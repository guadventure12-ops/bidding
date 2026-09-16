"""Read-only source extraction and faithful DOCX -> PDF bid exports.

The text and metadata returned by this module are untrusted source material,
never agent instructions. DOCX locations are structural; only PDFs have real
page numbers. Original source files are never modified.
"""
from __future__ import annotations

import base64
import hashlib
import importlib.util
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import zipfile
from contextlib import contextmanager
from datetime import date
from pathlib import Path
from typing import Any
from . import bid_body, review_rules
from xml.etree import ElementTree as ET

MAX_FILE_BYTES = 512 * 1024 * 1024
MAX_ZIP_BYTES = 1024 * 1024 * 1024
MAX_XML_BYTES = 64 * 1024 * 1024
MAX_BLOCK_CHARS = 4000
MAX_TEXT_CHARS = 12_000_000
MAX_PDF_PAGES = 3000
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_OCR_ENGINE = None
_OCR_LOCK = threading.Lock()
_RENDER_LOCK = threading.Lock()
WORD_RENDER_TIMEOUT_SECONDS = 600


def _warn(result: dict, text: str) -> None:
    if text not in result["warnings"]:
        result["warnings"].append(text)


def _clean(text: Any) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(text or "")).strip()


def _split_text(text: str, size: int = MAX_BLOCK_CHARS):
    while len(text) > size:
        cut = max(text.rfind("\n", size // 2, size), text.rfind("。", size // 2, size))
        cut = cut + 1 if cut >= size // 2 else size
        yield text[:cut].strip()
        text = text[cut:].strip()
    if text:
        yield text


def _add_block(result: dict, text: str, locator: str, kind: str = "paragraph", page=None, **extra):
    text = _clean(text)
    if not text:
        return
    remaining = MAX_TEXT_CHARS - result["text_chars"]
    if review_rules.active('file_parse_safety') and len(text) > remaining:
        text = text[:max(0, remaining)]
        _warn(result, "文本超过 1200 万字符安全上限，尾部已截断；请拆分文件后重新导入。")
        result["truncated"] = True
    parts = list(_split_text(text))
    for i, part in enumerate(parts, 1):
        block = {
            "id": f"b{len(result['blocks']) + 1:06d}", "text": part,
            "locator": locator + (f" / 分片 {i}/{len(parts)}" if len(parts) > 1 else ""),
            "kind": kind, "page": page,
        }
        # Large table cells must not bypass the per-block size bound through metadata.
        block.update(extra)
        if len(parts) > 1 and "table" in block:
            block["table"] = {k: v for k, v in block["table"].items() if k != "cells"}
            block["table"]["split"] = True
        result["blocks"].append(block)
        result["text_chars"] += len(part)


def _safe_xml(data: bytes) -> ET.Element:
    if review_rules.active('file_parse_safety') and len(data) > MAX_XML_BYTES:
        raise ValueError("DOCX 内部 XML 超过 64 MB 安全上限。")
    if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", data, re.I):
        raise ValueError("文件包含不支持的 XML 实体声明。")
    return ET.fromstring(data)


def _check_zip(zf: zipfile.ZipFile):
    infos = zf.infolist()
    bounded = review_rules.active('file_parse_safety')
    if bounded and (len(infos) > 20000 or sum(i.file_size for i in infos) > MAX_ZIP_BYTES):
        raise ValueError("DOCX 解压规模超过安全上限。")
    for info in infos:
        if info.flag_bits & 1:
            raise ValueError("不支持加密的 DOCX 压缩包。")
        if bounded and info.file_size > 10 * 1024 * 1024 and info.file_size / max(1, info.compress_size) > 1000:
            raise ValueError("DOCX 内部文件压缩比异常，请检查来源。")
        if bounded and info.filename.endswith(".xml") and info.file_size > MAX_XML_BYTES:
            raise ValueError("DOCX 内部 XML 超过 64 MB 安全上限。")


def _paragraph_text(element: ET.Element) -> str:
    fragments = []
    def walk(node):
        if node.tag in {W + "del", W + "txbxContent"}:
            return
        if node.tag == W + "t":
            fragments.append(node.text or "")
        elif node.tag in {W + "br", W + "cr"}:
            fragments.append("\n")
        elif node.tag == W + "tab":
            fragments.append("\t")
        else:
            for child in node:
                walk(child)
    walk(element)
    return _clean("".join(fragments))


def _ocr_bytes(image_data: bytes) -> tuple[str, float | None]:
    """Local OCR; no remote calls. Serialize the shared ONNX engine."""
    global _OCR_ENGINE
    from PIL import Image
    from io import BytesIO
    import numpy as np
    if review_rules.active('file_parse_safety') and len(image_data) > 80 * 1024 * 1024:
        raise ValueError("图片超过 80 MB OCR 上限")
    with Image.open(BytesIO(image_data)) as source:
        if review_rules.active('file_parse_safety') and source.width * source.height > 60_000_000:
            raise ValueError("图片像素超过 OCR 安全上限")
        image = source.convert("RGB")
        image.thumbnail((3200, 4800))
        arr = np.asarray(image)
    with _OCR_LOCK:
        if _OCR_ENGINE is None:
            from rapidocr_onnxruntime import RapidOCR
            _OCR_ENGINE = RapidOCR(intra_op_num_threads=2, inter_op_num_threads=2)
        rows, _ = _OCR_ENGINE(arr)
    if not rows:
        return "", None
    return "\n".join(str(row[1]) for row in rows), sum(float(row[2]) for row in rows) / len(rows)


def _parse_docx(path: Path, result: dict, ocr: bool):
    import posixpath
    with zipfile.ZipFile(path) as zf:
        _check_zip(zf)
        names = set(zf.namelist())
        if "word/document.xml" not in names:
            raise ValueError("不是有效的 DOCX 文档。")
        root = _safe_xml(zf.read("word/document.xml"))
        body = root.find(W + "body")
        if body is None:
            raise ValueError("DOCX 缺少正文。")
        media = [n for n in names if n.startswith("word/media/") and not n.endswith("/")]
        result["image_count"] = len(media)
        result["table_count"] = len(list(body.iter(W + "tbl")))
        result["heading_outline"] = []
        style_names = {}
        if "word/styles.xml" in names:
            for style in _safe_xml(zf.read("word/styles.xml")).findall(W + "style"):
                name = style.find(W + "name")
                if name is not None:
                    style_names[style.attrib.get(W + "styleId", "")] = name.attrib.get(W + "val", "")
        _warn(result, "DOCX 页码随 Word 排版变化；引用使用段落、表格和行定位，不代表原文页码。")
        if "docProps/app.xml" in names:
            try:
                app = _safe_xml(zf.read("docProps/app.xml"))
                value = next((x.text for x in app.iter() if x.tag.endswith("}Pages")), None)
                result["reported_page_count"] = int(value) if value else None
            except (ValueError, ET.ParseError):
                pass
        if media:
            _warn(result, f"文档含 {len(media)} 个嵌入图像文件。图示、印章、签名及截图版式需核对原件；OCR 不能验证其真实性。")
            if not ocr:
                _warn(result, "DOCX 嵌入图片未启用 OCR，图片中的文字未计入提取结果。")
        if any(n.startswith("word/embeddings/") and not n.endswith("/") for n in names):
            _warn(result, "文档包含嵌入对象/附件，未读取其内部内容，请另行导入附件。")
        if list(root.iter(W + "altChunk")):
            _warn(result, "文档包含外部格式片段 altChunk，未展开，请核对原件。")
        if list(root.iter(W + "ins")) or list(root.iter(W + "del")):
            _warn(result, "文档存在修订记录，提取采用保留插入、排除删除的正文视图；请人工确认有效版本。")
        if list(root.iter(W + "numPr")):
            _warn(result, "Word 自动列表编号可能未显示在提取文本中；条款请按原段落位置复核。")
        if list(root.iter(W + "txbxContent")):
            _warn(result, "文本框文字已单独提取，浮动对象的视觉阅读顺序需按原件复核。")
        if any(x.tag.endswith("}anchor") for x in root.iter()):
            _warn(result, "存在浮动图形，图形位置和关联关系未保留，请核对原件。")
        para_index = 0
        table_index = 0
        ocr_cache = {}
        ocr_attempted = set()

        def rels_for(part):
            rel_name = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
            if rel_name not in names:
                return {}
            relroot = _safe_xml(zf.read(rel_name))
            return {r.attrib.get("Id"): posixpath.normpath(posixpath.join(posixpath.dirname(part), r.attrib.get("Target", "")))
                    for r in relroot if r.attrib.get("TargetMode") != "External"}

        def images_in(node, locator, rels):
            if not ocr:
                return
            image_refs = [n.attrib.get(R + "embed") for n in node.iter(A + "blip")]
            image_refs += [n.attrib.get(R + "id") for n in node.iter() if n.tag.endswith("}imagedata")]
            for image_index, rel_id in enumerate(image_refs, 1):
                media_name = rels.get(rel_id)
                if not media_name or media_name not in names:
                    continue
                try:
                    ocr_attempted.add(media_name)
                    if media_name not in ocr_cache:
                        ocr_cache[media_name] = _ocr_bytes(zf.read(media_name))
                    text, confidence = ocr_cache[media_name]
                    if text:
                        _add_block(result, text, f"{locator} / 图片 {image_index} OCR", "image_ocr", confidence=confidence)
                        if confidence is not None and confidence < .8:
                            _warn(result, f"图片 OCR 置信度偏低：{locator}，请核对数字、参数及名称。")
                    else:
                        _warn(result, f"图片未识别出文字：{locator} / 图片 {image_index}，请查看原件。")
                except Exception as exc:
                    _warn(result, f"图片 OCR 未完成：{locator}，{type(exc).__name__}：{str(exc)[:160]}。")

        def paragraph(node, label, rels):
            nonlocal para_index
            para_index += 1
            locator = f"{label} / 段落 {para_index}"
            text = _paragraph_text(node)
            style = node.find("./" + W + "pPr/" + W + "pStyle")
            style_id = style.attrib.get(W + "val", "") if style is not None else ""
            style_name = style_names.get(style_id, style_id)
            kind = "heading" if re.search(r"heading|标题|^title$", style_name, re.I) else "paragraph"
            if kind == "heading" and text:
                result["heading_outline"].append({"text": text[:300], "locator": locator, "style": style_name})
            _add_block(result, text, locator, kind, style=style_name, style_id=style_id)
            images_in(node, locator, rels)
            for box_index, box in enumerate(node.iter(W + "txbxContent"), 1):
                texts = [_paragraph_text(p) for p in box.iter(W + "p")]
                _add_block(result, "\n".join(texts), f"{locator} / 文本框 {box_index}", "textbox")

        def table(node, label, rels):
            nonlocal table_index
            table_index += 1
            tid = table_index
            for row_no, row in enumerate(node.findall(W + "tr"), 1):
                cells = []
                spans = []
                for cell in row.findall(W + "tc"):
                    merge = cell.find("./" + W + "tcPr/" + W + "vMerge")
                    span = cell.find("./" + W + "tcPr/" + W + "gridSpan")
                    spans.append(int(span.attrib.get(W + "val", 1)) if span is not None else 1)
                    if merge is not None and merge.attrib.get(W + "val", "continue") == "continue":
                        cells.append("")
                    else:
                        cells.append("\n".join(filter(None, (_paragraph_text(p) for p in cell.iter(W + "p")))))
                locator = f"{label} / 表格 {tid} / 第 {row_no} 行"
                _add_block(result, " | ".join(cells), locator, "table_row", table={"index": tid, "row": row_no, "cells": cells, "spans": spans})
                images_in(row, locator, rels)

        def visit(container, label, rels):
            for child in container:
                if child.tag == W + "p":
                    paragraph(child, label, rels)
                elif child.tag == W + "tbl":
                    table(child, label, rels)
                elif child.tag not in {W + "sectPr", W + "del"}:
                    visit(child, label, rels)

        visit(body, "正文", rels_for("word/document.xml"))
        for part in sorted(n for n in names if re.fullmatch(r"word/(header|footer)\d+\.xml", n)):
            partroot = _safe_xml(zf.read(part))
            visit(partroot, "页眉 " + Path(part).stem if "/header" in part else "页脚 " + Path(part).stem, rels_for(part))
        for part, label, item in [("word/footnotes.xml", "脚注", "footnote"), ("word/endnotes.xml", "尾注", "endnote")]:
            if part in names:
                for note in _safe_xml(zf.read(part)).findall(W + item):
                    if int(note.attrib.get(W + "id", 0)) > 0:
                        visit(note, f"{label} {note.attrib.get(W + 'id')}", rels_for(part))
        result["ocr_images_attempted"] = len(ocr_attempted)
        result["ocr_images_with_text"] = sum(bool(v[0]) for v in ocr_cache.values())
        if ocr and len(ocr_attempted) < len(media):
            _warn(result, f"{len(media) - len(ocr_attempted)} 个嵌入图像未在可识别正文位置展开；请检查原文图形和附件。")


def _parse_pdf(path: Path, result: dict, ocr: bool):
    scanned = []
    if importlib.util.find_spec("fitz"):
        import fitz
        with fitz.open(path) as pdf:
            if pdf.needs_pass:
                raise ValueError("PDF 已加密，请先提供无密码的可读取副本。")
            result["page_count"] = len(pdf)
            if review_rules.active('file_parse_safety') and len(pdf) > MAX_PDF_PAGES:
                raise ValueError(f"PDF 超过 {MAX_PDF_PAGES} 页，请分卷导入。")
            for index, page in enumerate(pdf, 1):
                text = page.get_text("text", sort=True)
                sparse = len(re.sub(r"\s", "", text)) < 30
                images = page.get_images(full=True)
                image_coverage = 0.0
                if images:
                    try:
                        image_coverage = max((r.get_area() / max(1, page.rect.get_area()) for im in images for r in page.get_image_rects(im[0])), default=0)
                    except Exception:
                        pass
                # Text layers may contain only page numbers or an incomplete overlay.
                needs_ocr = sparse or (image_coverage > .6 and len(re.sub(r"\s", "", text)) < 250)
                if needs_ocr:
                    scanned.append(index)
                _add_block(result, text, f"第 {index} 页", "pdf_text", page=index)
                if ocr and needs_ocr:
                    try:
                        scale = min(2.5, math.sqrt(16_000_000 / max(1, page.rect.width * page.rect.height)))
                        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                        ocr_text, confidence = _ocr_bytes(pix.tobytes("png"))
                        if ocr_text:
                            _add_block(result, ocr_text, f"第 {index} 页 / OCR", "ocr", page=index, confidence=confidence)
                            if confidence is not None and confidence < .8:
                                _warn(result, f"第 {index} 页 OCR 置信度偏低，请复核金额、参数和名称。")
                        else:
                            _warn(result, f"第 {index} 页 OCR 未识别出文字，请查看原件。")
                    except Exception as exc:
                        _warn(result, f"第 {index} 页 OCR 未完成：{type(exc).__name__}：{str(exc)[:160]}。")
                if images and not needs_ocr:
                    result.setdefault("pages_with_non_ocr_images", []).append(index)
    else:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ValueError("PDF 已加密，请先提供无密码的可读取副本。")
        result["page_count"] = len(reader.pages)
        if review_rules.active('file_parse_safety') and len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError(f"PDF 超过 {MAX_PDF_PAGES} 页，请分卷导入。")
        for index, page in enumerate(reader.pages, 1):
            try:
                text = page.extract_text() or ""
                _add_block(result, text, f"第 {index} 页", "pdf_text", page=index)
                if len(re.sub(r"\s", "", text)) < 30:
                    scanned.append(index)
            except Exception as exc:
                _warn(result, f"第 {index} 页文本提取失败：{type(exc).__name__}，请检查原件。")
        _warn(result, "当前采用 pypdf 提取；未检测图像覆盖，复杂多栏表格和截图请核对原件。")
        if ocr:
            _warn(result, "OCR 未执行：缺少 PyMuPDF 页面渲染组件。")
    result["possible_scanned_pages"] = scanned
    if scanned:
        label = "、".join(map(str, scanned[:50])) + (" 等" if len(scanned) > 50 else "")
        _warn(result, f"发现 {len(scanned)} 个少文字或疑似扫描页：{label}；" + ("已尝试 OCR，空白页、图示及识别结果仍需人工核对。" if ocr else "未启用 OCR，这些页可能存在未提取的要求。"))
    if result.get("pages_with_non_ocr_images"):
        _warn(result, "部分文字页含图片，未对每张图片单独 OCR；请复核截图、图示、签章及图中要求。")
    _warn(result, "PDF 提取保留页码，表格的合并单元格和多栏阅读顺序需按原页复核。")


def parse_document(path: str, ocr: bool = False) -> dict:
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"文件不存在：{source.name}")
    if review_rules.active('file_parse_safety') and source.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("文件超过 512 MB，请拆分后导入。")
    ext = source.suffix.lower().lstrip(".")
    if ext not in {"docx", "pdf", "md", "txt", "markdown", "xlsx", "csv", "pptx", "html", "htm", "png", "jpg", "jpeg", "tif", "tiff", "webp", "bmp"}:
        raise ValueError("不支持该文件格式，请使用 DOCX、PDF、TXT、MD、XLSX、CSV、PPTX、HTML 或常见图片；旧版 DOC 请先另存为 DOCX。")
    result = {"name": source.name, "format": ext, "blocks": [], "warnings": [], "page_count": None, "text_chars": 0, "truncated": False}
    digest = hashlib.sha256()
    with source.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    result["sha256"] = digest.hexdigest()
    if ext == "docx":
        try:
            _parse_docx(source, result, bool(ocr))
        except zipfile.BadZipFile:
            raise ValueError("Office 文件可能已加密或损坏，无法直接提取；请在有权访问的 Office 中打开后另存为可读取的 DOCX 副本。") from None
    elif ext == "pdf":
        _parse_pdf(source, result, bool(ocr))
    elif ext in {"xlsx", "csv", "pptx", "html", "htm"}:
        from .extra_documents import parse_extra
        parse_extra(source, result, bool(ocr))
    elif ext in {"txt", "md", "markdown"}:
        raw = source.read_bytes()
        encoding = None
        for candidate in ("utf-8-sig", "utf-16" if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else "gb18030"):
            try:
                text = raw.decode(candidate)
                encoding = candidate
                break
            except UnicodeError:
                continue
        if encoding is None:
            text = raw.decode("utf-8", errors="replace")
            encoding = "utf-8-with-replacement"
            _warn(result, "文本编码无法完整识别，部分字符已替换，请核对原件。")
        result["encoding"] = encoding
        group = []
        start = 1
        for line_no, line in enumerate(text.splitlines(), 1):
            if not group:
                start = line_no
            group.append(line)
            if not line.strip() or sum(len(x) for x in group) >= MAX_BLOCK_CHARS:
                _add_block(result, "\n".join(group), f"第 {start}-{line_no} 行", "text")
                group = []
        if group:
            _add_block(result, "\n".join(group), f"第 {start}-{len(text.splitlines())} 行", "text")
    else:
        result["image_count"] = 1
        if ocr:
            try:
                text, confidence = _ocr_bytes(source.read_bytes())
                _add_block(result, text, "图像 / OCR", "image_ocr", confidence=confidence)
                _warn(result, "图片文字由本地 OCR 识别，需核对原图中的金额、参数、签章及细小文字。")
            except Exception as exc:
                _warn(result, f"图片 OCR 未完成：{type(exc).__name__}：{str(exc)[:160]}。")
        else:
            _warn(result, "图片未启用 OCR，未提取文字。")
    if not result["blocks"]:
        _warn(result, "未提取到可用文字，不能据此判定文件没有招标要求；请检查扫描件、密码或格式。")
    return result


def _font(run, size=None, bold=None, name="宋体"):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    run.font.name = name
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:cs"), name)
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        fonts.attrib.pop(qn('w:'+key),None)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _field(paragraph, instruction: str, hint: str = ""):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    for tag, value in (("begin", None), ("instruction", instruction), ("separate", None)):
        el = OxmlElement("w:instrText" if tag == "instruction" else "w:fldChar")
        if value:
            el.set(qn("xml:space"), "preserve")
            el.text = value
        else:
            el.set(qn("w:fldCharType"), tag)
        run._r.append(el)
    if hint:
        paragraph.add_run(hint)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def _style_table(table, widths):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            edge.set(qn("w:" + key), value)
        borders.append(edge)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, width in (("top", "90"), ("bottom", "90"), ("left", "90"), ("right", "90")):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:w"), width)
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    props.append(margins)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    header_repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header_repeat)
    for i, row in enumerate(table.rows):
        estimated_lines = max((sum(max(1, math.ceil(sum(1 if ord(c) > 127 else .55 for c in line) / max(1, widths[j] * 28.35 / 12))) for line in cell.text.splitlines()) for j, cell in enumerate(row.cells)), default=1)
        if i == 0 or estimated_lines <= 22:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E8EEF5")
                cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.left_indent = Pt(0)
                indent = p._p.get_or_add_pPr().get_or_add_ind()
                indent.set(qn("w:firstLineChars"), "0")
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.keep_with_next = i == 0
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or (j == 0 and len(cell.text) < 15) else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    _font(run, 12, False)


def _inline(paragraph, text: str):
    # Markdown links remain visible labels + URLs; never fetch external content.
    text = re.sub(r"!?\[([^\]]*)\]\(([^)]+)\)", r"\1（\2）", text)
    for index, part in enumerate(re.split(r"\*\*(.*?)\*\*", text)):
        if part:
            _font(paragraph.add_run(part.strip("`") if part.startswith("`") else part), size=12, bold=False)


def clean_export_locators(text, requirement_ids=()):
    """Remove unambiguous internal trace annotations from export copies only."""
    ids={str(value) for value in requirement_ids}
    prefixes={value[:8] for value in ids if len(value)==32}
    known=ids|prefixes
    part=r'(?:段落\s*\d+(?:[-–—]\d+)?|表格\s*\d+|第\s*\d+\s*(?:行|页|段)|页\s*\d+|行\s*\d+)'
    location=r'(?:(?:正文|附件(?:\s*\d+)?|第\s*\d+\s*页)(?:\s*/\s*'+part+r')+|第\s*\d+\s*页)'
    label=re.compile(r'^(?:采购条款定位|采购文件定位|招标原文位置|原文位置|来源定位)\s*[:：]\s*'+location+r'\s*[。；;]?\s*$')
    annotation=re.compile(r'[（(]\s*(?:对应要求|关联要求|要求\s*ID|需求\s*ID)\s*[:：]?\s*([0-9a-fA-F]{8,32}(?:\s*[、,，/]\s*[0-9a-fA-F]{8,32})*)\s*[）)]')
    output,removed=[],[]
    fence=None
    for number,line in enumerate((text or '').splitlines(keepends=True),1):
        marker=re.match(r'^\s*(`{3,}|~{3,})',line)
        if marker:
            if fence is None:fence=marker[1][0]
            elif marker[1][0]==fence:fence=None
            output.append(line);continue
        if fence:
            output.append(line);continue
        def remove_reference(match):
            values=re.split(r'\s*[、,，/]\s*',match[1])
            if all(value in known for value in values) or re.match(r'^\s*#{1,6}\s+',line):
                removed.append({'line':number,'raw':match[0],'kind':'requirement_annotation'})
                return ''
            return match[0]
        current=annotation.sub(remove_reference,line)
        plain=re.sub(r'^(?:#{1,6}\s+|>\s*|[-*+]\s+)', '',current.strip()).replace('**','')
        if label.fullmatch(plain):
            removed.append({'line':number,'raw':current.rstrip('\r\n'),'kind':'source_locator'})
            continue
        if current.lstrip().startswith('|'):
            cells=re.split(r'(?<!\\)\|',current)
            for index,cell in enumerate(cells):
                if label.fullmatch(cell.strip().replace('**','')):
                    removed.append({'line':number,'raw':cell,'kind':'source_locator'})
                    ending=re.search(r'(\r?\n)$',cell)
                    cells[index]=' '+(ending[0] if ending else '')
            current='|'.join(cells)
        output.append(current)
    return {'content':''.join(output),'removed':removed}


_ASSET_IMAGE_LINE = re.compile(r"!\[((?:\\.|[^\]\\])*)\]\(asset:(asset-[0-9a-f]{64})\)")


def _drop_internal_evidence_columns(text):
    """Drop only pure internal-reference columns from a Markdown export copy."""
    lines = text.splitlines(keepends=True)
    removed, fence, i = [], None, 0
    def cells(line):
        line = line.strip()
        if not line.startswith('|') or not line.endswith('|'):
            return None
        return re.split(r'(?<!\\)\|', line)[1:-1]
    while i < len(lines):
        marker = re.match(r'^\s*(`{3,}|~{3,})', lines[i])
        if marker:
            fence = marker[1][0] if fence is None else (None if marker[1][0] == fence else fence)
        headers = cells(lines[i]) if not fence and not marker else None
        separator = cells(lines[i+1]) if headers and i+1 < len(lines) else None
        if not separator or len(separator) != len(headers) or not all(re.fullmatch(r'\s*:?-{3,}:?\s*', v) for v in separator):
            i += 1
            continue
        end, rows = i+2, [headers, separator]
        while end < len(lines) and cells(lines[end]) is not None:
            rows.append(cells(lines[end])); end += 1
        drop = []
        if len(rows) > 2 and all(len(row) == len(headers) for row in rows):
            for column, header in enumerate(headers):
                values = [row[column].strip() for row in rows[2:]]
                if header.strip() in {'证据', '来源'} and any(values) and all(
                    re.fullmatch(r'(?:\s*\[E:[A-Za-z0-9_-]+\]\s*)*', value) for value in values
                ):
                    drop.append(column)
        # Never erase a table which has no remaining non-reference column.
        if drop and len(drop) < len(headers):
            for column in drop:
                removed.append({'kind':'internal_evidence_column', 'line':i+1,
                                'column':column+1, 'header':headers[column].strip(),
                                'raw_cells':[row[column] for row in rows[2:]]})
            for offset, row in enumerate(rows):
                original = lines[i+offset]
                ending = '\r\n' if original.endswith('\r\n') else ('\n' if original.endswith('\n') else '')
                indent = re.match(r'^\s*', original)[0]
                lines[i+offset] = indent+'|'+'|'.join(value for c,value in enumerate(row) if c not in drop)+'|'+ending
        i = end
    return {'content':''.join(lines), 'removed':removed}


def _asset_image_paragraphs(doc, caption, asset_id, assets, asset_root):
    """Embed an explicitly selected immutable asset, never a model-supplied path."""
    import copy
    import io
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from .document_assets import verified_asset_path

    if not isinstance(assets, dict) or asset_id not in assets:
        raise ValueError("图片资源未在本次获准范围内：" + asset_id)
    asset = assets[asset_id]
    if not isinstance(asset, dict) or asset.get('asset_id') != asset_id:
        raise ValueError("图片资源ID与本次选择记录不一致")
    if not asset_root:
        raise ValueError("Word图片缺少明确的本地资源目录")
    source = asset.get('source') or {}
    if not isinstance(source, dict):
        raise ValueError("图片来源定位记录无效")
    if source.get('display_transform'):
        raise ValueError("图片在原件中的裁剪、旋转或翻转尚未复现，不能直接导出：" + asset_id)
    caption = _clean(re.sub(r'\\([\\\[\]])', r'\1', caption))
    if not caption or len(caption) > 300:
        raise ValueError("图片须提供1至300字的明确图题")
    path = verified_asset_path(asset, asset_root=asset_root)
    data = path.read_bytes()
    if hashlib.sha256(data).hexdigest() != asset.get('sha256'):
        raise ValueError("图片在校验后发生变化，未写入Word：" + asset_id)

    page = doc.sections[-1]
    available_width = int(page.page_width - page.left_margin - page.right_margin)
    available_height = int(page.page_height - page.top_margin - page.bottom_margin)
    # Reserve caption lines and paragraph spacing inside the usable page. Small
    # rasters retain their 96-dpi size rather than being enlarged into blurry art.
    estimated_caption_width = sum(12 if ord(char) > 127 else 6 for char in caption)
    caption_lines = max(1, math.ceil(estimated_caption_width / max(1, available_width / 12700)))
    available_height -= int(Pt(caption_lines * 15 + 30))
    if available_width <= 0 or available_height <= 0:
        raise ValueError("当前页面没有足够空间放置图片及图题")
    native_width, native_height = asset['width'] * 9525, asset['height'] * 9525
    scale = min(1.0, available_width / native_width, available_height / native_height)
    width, height = max(1, int(native_width * scale)), max(1, int(native_height * scale))

    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.first_line_indent = Pt(0)
    picture_paragraph.paragraph_format.left_indent = Pt(0)
    picture_paragraph._p.get_or_add_pPr().get_or_add_ind().set(qn('w:firstLineChars'), '0')
    picture_paragraph.paragraph_format.space_before = Pt(6)
    picture_paragraph.paragraph_format.space_after = Pt(3)
    picture_paragraph.paragraph_format.line_spacing = 1
    picture_paragraph.paragraph_format.keep_together = True
    picture_paragraph.paragraph_format.keep_with_next = True
    picture = picture_paragraph.add_run().add_picture(io.BytesIO(data), width=width, height=height)
    picture._inline.docPr.set('descr', caption)
    picture._inline.docPr.set('title', caption)

    label = doc.add_paragraph(style='Caption')
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.first_line_indent = Pt(0)
    label.paragraph_format.left_indent = Pt(0)
    label._p.get_or_add_pPr().get_or_add_ind().set(qn('w:firstLineChars'), '0')
    label.paragraph_format.space_before = Pt(0)
    label.paragraph_format.space_after = Pt(6)
    label.paragraph_format.line_spacing = 1.15
    label.paragraph_format.keep_together = True
    label.paragraph_format.keep_with_next = False
    run = label.add_run(caption)
    _font(run, 12, False)
    run.italic = False
    return {"asset_id": asset_id, "sha256": asset['sha256'],
            "document_id": asset.get('document_id'), "source_sha256": asset.get('source_sha256'),
            "source": copy.deepcopy(source), "caption": caption}


def _reject_unresolved_image_syntax(line, assets):
    # Legacy profiles retain their original visible link text and never fetch
    # images. When an explicit manifest is supplied, all pictures must use its
    # asset IDs; inline/table references cannot silently print internal IDs.
    visible = re.sub(r'`+[^`]*`+', '', line)
    if re.search(r'!\[[^\n]*\]\(\s*asset:', visible):
        raise ValueError("资源图片必须独占一行，格式为 ![图题](asset:asset-ID)")
    if assets and re.search(r'!\[[^\n]*\]\(', visible):
        raise ValueError("Word图片只接受本次已选asset ID，不接受模型给出的文件路径或URL")


def _proposal_column_widths(data, total=16.0):
    """Reserve readable short columns; give prose columns space to explain."""
    headers=data[0];count=len(headers)
    short=r'序号|编号|条款号|(?:对应)?页码|页数|分值|数量'
    minimum=[1.2 if re.fullmatch(short,str(h)) else 1.5 for h in headers]
    if sum(minimum)>=total:return [total/count]*count
    weights=[]
    for index,header in enumerate(headers):
        lengths=sorted(len(str(row[index])) for row in data[1:] if index<len(row))
        typical=lengths[min(len(lengths)-1,len(lengths)*3//4)] if lengths else len(header)
        weight=math.sqrt(max(1,min(typical,240)))
        if re.search('内容|说明|措施|标准|条款|描述|成果',str(header)):weight*=1.7
        if re.fullmatch(short,str(header)):weight=.3
        weights.append(weight)
    remainder=total-sum(minimum)
    return [base+remainder*weight/sum(weights) for base,weight in zip(minimum,weights)]


def _write_markdown(doc, content: str, base_level=2, *, assets=None, asset_root=None, proposal_layout=False, keep_short_forms=False):
    lines = _clean(content).splitlines()
    i = 0
    heading_stack = []
    used_assets = []
    image_code_fence = None
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        marker = re.match(r'^(`{3,}|~{3,})(.*)$', line)
        if marker:
            if image_code_fence is None:
                image_code_fence = (marker[1][0], len(marker[1]))
            elif marker[1][0] == image_code_fence[0] and len(marker[1]) >= image_code_fence[1] and not marker[2].strip():
                image_code_fence = None
        picture = _ASSET_IMAGE_LINE.fullmatch(line) if image_code_fence is None and not marker else None
        if picture:
            used_assets.append(_asset_image_paragraphs(doc, picture[1], picture[2], assets, asset_root))
            i += 1
            continue
        if image_code_fence is None and not marker:
            _reject_unresolved_image_syntax(line, assets)
        heading = re.match(r"^(#{1,6})\s+(.+)$", line) if image_code_fence is None and not marker else None
        if heading:
            depth = len(heading.group(1))
            while heading_stack and heading_stack[-1] >= depth:
                heading_stack.pop()
            heading_stack.append(depth)
            from .chapter_outline import _heading_title
            doc.add_heading(_heading_title(heading.group(2)), level=min(9, base_level - 1 + len(heading_stack)))
        elif "|" in line and i + 1 < len(lines) and (re.fullmatch(r"\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*", lines[i + 1])
                or proposal_layout and re.fullmatch(r"\s*\|\s*:?-{3,}:?\s*\|\s*", lines[i + 1])):
            headers = [v.strip() for v in line.strip("|").split("|")]
            data = [headers]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                if image_code_fence is None:
                    _reject_unresolved_image_syntax(lines[i], assets)
                row = [v.strip() for v in lines[i].strip().strip("|").split("|")]
                data.append((row + [""] * len(headers))[:len(headers)])
                i += 1
            table = doc.add_table(rows=len(data), cols=len(headers))
            for ri, row in enumerate(data):
                for ci, value in enumerate(row):
                    table.cell(ri, ci).text = value.replace("**", "")
            _style_table(table, _proposal_column_widths(data) if proposal_layout else [16 / len(headers)] * len(headers))
            # Compact qualification identity/relationship forms should not
            # strand their final fields on a second page. Long condition tables
            # retain normal pagination; no body text or grid is changed.
            if keep_short_forms and 2 <= len(data) <= 8 and sum(len(v) for row in data for v in row) <= 500:
                for ri,row in enumerate(table.rows):
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.keep_with_next = ri < len(data)-1
            doc.add_paragraph()
            continue
        elif line.startswith("```"):
            pass
        else:
            bullet = re.match(r"^[-*+]\s+(.+)", line)
            numbered = re.match(r"^\d+[.)]\s+(.+)", line)
            quote = re.match(r"^>\s+(.*)", line)
            p = doc.add_paragraph(style="List Bullet" if bullet else "Normal")
            _inline(p, bullet.group(1) if bullet else quote.group(1) if quote else line)
            # Plain stage labels are subheadings even without Markdown hashes.
            if proposal_layout and re.fullmatch(r'（[一二三四五六七八九十]+）[^。；！？\n]{1,24}阶段', p.text):
                p.paragraph_format.keep_with_next = True
        i += 1
    return used_assets


def _deviation_value(requirement: dict) -> str:
    explicit = _clean(requirement.get("deviation"))
    if explicit:
        return explicit
    match = re.search(r"(?<!非)(无偏离|正偏离|负偏离)", _clean(requirement.get("response")))
    return match.group(1) if match else "________________"


def _requirement_display_labels(requirements: list) -> dict:
    labels = {}
    for index, requirement in enumerate(requirements, 1):
        id = str(requirement.get("id") or "")
        if id:
            labels.setdefault(id, f"要求{index}")
    return labels


def _requirement_reference_replacer(labels: dict):
    """Display known IDs, with short prefixes allowed only in requirement context."""
    full = {id: label for id, label in labels.items() if re.fullmatch(r"[0-9a-f]{32}", id)}
    prefix_ids = {}
    for id in full:
        prefix_ids.setdefault(id[:8], []).append(id)
    short = {prefix: full[ids[0]] for prefix, ids in prefix_ids.items() if len(ids) == 1}
    token = re.compile(r"(?<![A-Za-z0-9_./-])(?:[0-9a-f]{32}|[0-9a-f]{8})(?![A-Za-z0-9_./-])")
    reference = re.compile(r"(?P<lead>对应|关联|相关|涉及|参见|见)?(?:要求|需求|条款)(?:\s*(?:ID|编号|标识))?\s*[:：]?\s*`?(?P<id>[0-9a-f]{32}|[0-9a-f]{8})`?(?![A-Za-z0-9_./-])", re.I)
    separator = re.compile(r"\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*")
    requirement_column = re.compile(r"(?:对应|关联|相关|涉及)?(?:要求|需求|条款)(?:ID|编号|标识)?", re.I)
    checksum_label = re.compile(r"(?:hash|sha(?:-?\d+)?|md5|哈希|校验(?:码|值)|摘要)", re.I)

    def replace_cell(value, allow_short=False, checksum=False):
        if checksum:
            return value
        value = reference.sub(lambda match: (match.group("lead") or "") + (full.get(match.group("id")) or short.get(match.group("id")))
                              if match.group("id") in full or match.group("id") in short else match.group(), value)
        def replace_token(match):
            # A checksum explicitly described as such remains source data, even
            # in the unlikely case that its value equals an internal ID.
            preceding = value[max(0, match.start() - 24):match.start()]
            if checksum_label.search(preceding) and not re.search(r"[。；;\n|]", preceding):
                return match.group()
            return full.get(match.group()) or (short.get(match.group()) if allow_short else None) or match.group()
        return token.sub(replace_token, value)

    def replace(text):
        lines = _clean(text).splitlines()
        requirement_columns, checksum_columns = set(), set()
        in_table = False
        for index, line in enumerate(lines):
            if "|" in line and index + 1 < len(lines) and separator.fullmatch(lines[index + 1]):
                headers = line.strip().strip("|").split("|")
                headers = [re.sub(r"[\s*`]", "", value) for value in headers]
                requirement_columns = {i for i, value in enumerate(headers) if requirement_column.fullmatch(value)}
                checksum_columns = {i for i, value in enumerate(headers) if checksum_label.search(value)}
                in_table = True
            elif in_table and "|" in line and line.strip():
                if separator.fullmatch(line):
                    continue
                cells = line.split("|")
                offset = 1 if line.lstrip().startswith("|") else 0
                lines[index] = "|".join(replace_cell(cell, i - offset in requirement_columns, i - offset in checksum_columns) for i, cell in enumerate(cells))
                continue
            else:
                in_table = False
            lines[index] = replace_cell(line)
        return "\n".join(lines)
    return replace


def _append_deviation_tables(doc, requirements: list, project: dict, company: dict):
    """Column schemas observed in the supplied tender, not template fidelity."""
    technical_categories = {"technical", "implementation", "security", "service", "other", "技术", "实施", "安全", "服务"}
    technical = [r for r in requirements if r.get("category", "technical") in technical_categories]
    commercial = [r for r in requirements if r not in technical]
    for is_technical, items in ((False, commercial), (True, technical)):
        title = "附件12 技术偏离表" if is_technical else "附件11 商务条款偏离表"
        doc.add_heading(title, 1)
        if any(r.get("_display_requirement_label") for r in items):
            doc.add_paragraph("说明中的“要求N”为本稿响应台账标识，各分册使用同一编号；采购文件条目号和原文位置另列。").paragraph_format.keep_with_next = True
        doc.add_paragraph("采购项目：" + _clean((project.get('_cover_identity') or {}).get('project_name') or project.get("name") or project.get("title"))).paragraph_format.keep_with_next = True
        doc.add_paragraph("项目编号：" + _clean(project.get("project_number") or project.get("number") or "________________")).paragraph_format.keep_with_next = True
        kind = "技术" if is_technical else "商务"
        headers = ["序号", "采购文件条目号", f"采购文件的{kind}条款", f"响应文件的{kind}条款响应"]
        if is_technical:
            headers.append("偏离")
        headers.append("说明")
        table = doc.add_table(rows=1, cols=len(headers))
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        for index, requirement in enumerate(items, 1):
            row = table.add_row().cells
            source = _clean(requirement.get("source_locator") or requirement.get("locator"))
            row[0].text = str(index)
            row[1].text = _clean(requirement.get("number")) or source or "________________"
            row[2].text = _clean(requirement.get("text"))
            row[3].text = _clean(requirement.get("response")) or "________________"
            explanation = []
            if requirement.get("_display_requirement_label"):
                explanation.append("本稿要求标识：" + requirement["_display_requirement_label"])
            if is_technical:
                row[4].text = _deviation_value(requirement)
            else:
                explanation.append("偏离：" + _deviation_value(requirement))
            if source:
                explanation.append("原文位置：" + source)
            refs = requirement.get("_evidence_labels") or []
            if refs:
                explanation.append("内容来源：" + "、".join(refs))
            row[-1].text = "\n".join(explanation)
        if not items:
            row = table.add_row().cells
            row[0].text = "1"
            row[2].text = "________________"
            row[3].text = "________________"
        _style_table(table, [0.8, 2.2, 4.0, 4.1, 1.6, 3.3] if is_technical else [0.8, 2.5, 4, 5, 3.7])
        doc.add_paragraph("供应商：" + _clean(company.get("name") or company.get("company_name") or "________________") + "（盖章）")
        doc.add_paragraph("法定代表人/负责人或其授权代表：" + _clean(company.get("representative") or "________________"))
        doc.add_paragraph("日期：________________")


def _append_evidence_index(doc, company: dict):
    evidence = company.get("evidence") or []
    if not evidence:
        return
    doc.add_heading("企业资料来源索引", 1)
    doc.add_paragraph("正文引用与原始企业材料对应如下。原件及所需签章应按招标文件要求另行编排。")
    for item in evidence:
        if not isinstance(item, dict):
            continue
        label = _clean(item.get("_display_label") or "企业资料") + "　" + _clean(item.get("document_name") or item.get("name") or "企业材料")
        locator = _clean(item.get("locator"))
        doc.add_paragraph(label + (" / " + locator if locator else ""))


def cover_identity(project, source_chunks=()):
    """Use explicit tender cover text without renaming the workspace project."""
    buyer = _clean(project.get('buyer'))
    candidates, sources = set(), []
    chunks = list(source_chunks)
    for index, chunk in enumerate(chunks):
        text = _clean(chunk.get('text'))
        match = re.fullmatch(r'项目名称\s*[:：]\s*([^\n|。；]{2,120})', text)
        value = match.group(1).strip() if match else ''
        if not value and buyer and text == buyer and index + 1 < len(chunks):
            following = chunks[index + 1]
            next_text = _clean(following.get('text'))
            if (following.get('document_id') == chunk.get('document_id') and
                    re.fullmatch(r'[^\n|。；:：]{2,100}项目', next_text)):
                value, chunk = next_text, following
        if value:
            candidates.add(value)
            sources.append({k:chunk.get(k) for k in ('document_id','id','locator','text')})
    name = next(iter(candidates)) if len(candidates) == 1 else _clean(project.get('name') or project.get('title') or '项目投标文件')
    if buyer and name.startswith(buyer) and name != buyer:
        name = name[len(buyer):].lstrip(' \n')
    return {'buyer':buyer, 'project_name':name, 'status':'source_verified' if len(candidates)==1 else 'conflict' if candidates else 'project_field', 'sources':sources}


def _export_numbering(doc):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    root = doc.part.numbering_part.element
    abstract_id = max([int(n.get(qn('w:abstractNumId'))) for n in root.findall(qn('w:abstractNum'))]+[-1]) + 1
    num_id = max([int(n.get(qn('w:numId'))) for n in root.findall(qn('w:num'))]+[0]) + 1
    definition = OxmlElement('w:abstractNum');definition.set(qn('w:abstractNumId'),str(abstract_id))
    nsid=OxmlElement('w:nsid');nsid.set(qn('w:val'),f'{0x4D580000+abstract_id:08X}');definition.append(nsid)
    kind = OxmlElement('w:multiLevelType');kind.set(qn('w:val'),'multilevel');definition.append(kind)
    for level in range(9):
        item = OxmlElement('w:lvl');item.set(qn('w:ilvl'),str(level))
        values = [('start','1'),('numFmt','chineseCounting' if level==0 else 'decimal')]
        # Default multilevel semantics restart children after a parent. Each
        # heading paragraph carries its explicit numPr/ilvl binding.
        values += [('suff','space'),('lvlText','%1、' if level==0 else '.'.join('%'+str(i) for i in range(2,level+2))+'.'),('lvlJc','left')]
        for tag,value in values:
            el=OxmlElement('w:'+tag);el.set(qn('w:val'),value);item.append(el)
        props=OxmlElement('w:pPr');indent=OxmlElement('w:ind');indent.set(qn('w:left'),'0');indent.set(qn('w:firstLine'),'0');props.append(indent);item.append(props)
        definition.append(item)
    existing_nums=root.findall(qn('w:num'))
    root.insert(root.index(existing_nums[0]) if existing_nums else len(root),definition)
    num=OxmlElement('w:num');num.set(qn('w:numId'),str(num_id));ref=OxmlElement('w:abstractNumId');ref.set(qn('w:val'),str(abstract_id));num.append(ref);root.append(num)
    return num_id


def _format_body(doc, start, num_id, directory_labels=()):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for p in doc.paragraphs[start:]:
        heading = re.fullmatch(r'Heading ([1-9])',p.style.name)
        if heading:
            level=int(heading.group(1))-1
            # Remove an explicit heading label only, never numbers in prose,
            # tables, versions, product names or attachment identifiers.
            text=re.sub(r'^(?:第[一二三四五六七八九十百\d]+[章节篇]\s+|[一二三四五六七八九十百]+[、．]\s*|\d+(?:\.\d+)*(?:[.．)）]\s+|、\s*|\s+))(?=\S)', '',p.text)
            if p._p not in directory_labels and text != p.text and text:p.text=text
            p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.left_indent=Pt(0)
            props=p._p.get_or_add_pPr();props.get_or_add_ind().set(qn('w:firstLineChars'),'0')
            numbering=props.get_or_add_numPr()
            numbering.get_or_add_ilvl().val=level;numbering.get_or_add_numId().val=num_id
        elif p.style.name == 'Caption':
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.left_indent=Pt(0)
            p.paragraph_format.first_line_indent=Pt(0)
            p._p.get_or_add_pPr().get_or_add_ind().set(qn('w:firstLineChars'),'0')
            p.paragraph_format.keep_together=True
            p.paragraph_format.keep_with_next=False
            for run in p.runs:
                _font(run,12,False)
                run.italic=False
        elif p.text and p.text != '目录' and not p._p.xpath('.//w:instrText'):
            p.paragraph_format.left_indent=Pt(0)
            p.paragraph_format.first_line_indent=Pt(24)
            p._p.get_or_add_pPr().get_or_add_ind().set(qn('w:firstLineChars'),'200')
            for run in p.runs:_font(run,12,False)


def _write_cover(doc, project, company, cover, bidder, final):
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    def paragraph(text='',size=12,bold=False,center=True,style=None):
        p=doc.add_paragraph(style=style)
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1.5;p.paragraph_format.keep_with_next=False
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Pt(0)
        p._p.get_or_add_pPr().get_or_add_ind().set(qn('w:firstLineChars'),'0')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        _font(p.add_run(text),size,bold)
        # Empty spacer lines inherit the paragraph mark, not an empty run.
        mark=OxmlElement('w:rPr')
        for tag in ('w:sz','w:szCs'):
            el=OxmlElement(tag);el.set(qn('w:val'),str(size*2));mark.append(el)
        p._p.get_or_add_pPr().append(mark)
        return p
    paragraph().paragraph_format.line_spacing=1.8
    paragraph(size=16)
    paragraph('项目编号：'+_clean(project.get('project_number') or project.get('number') or '________________'),16,True,False)
    for _ in range(3):paragraph(size=16)
    paragraph(cover['buyer'],26,True,style='Title')
    paragraph(cover['project_name'],26,True,style='Title')
    paragraph();paragraph(size=24)
    paragraph(_clean(project.get('document_title') or '商务、技术文件'),28,True,style='Subtitle')
    paragraph('草稿 待审核' if not final else '')
    paragraph()
    paragraph('供应商名称：'+bidder+'（盖章）')
    paragraph('法定代表人/负责人或者其委托代理人：'+_clean(company.get('representative') or '________________')+'（签字）')
    paragraph(_clean(project.get('document_date')) or '____年____月____日')
    paragraph()
    if not final and any(c.get('code')=='model_review_missing' for c in company.get('checks',[]) if isinstance(c,dict)):
        paragraph('当前版本独立 AI 复核未完成\n本稿仅供内部校核，不可递交',10)


def compose_docx(project: dict, requirements: list, sections: list, output_path: str, company: dict | None = None, final: bool = False) -> dict:
    import copy
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    company = copy.deepcopy(company or {})
    selected_assets = company.get('assets') or []
    if not isinstance(selected_assets, list):
        raise ValueError("本次图片资源必须是明确选择的manifest列表")
    asset_map = {}
    for asset in selected_assets:
        if not isinstance(asset, dict) or not re.fullmatch(r'asset-[0-9a-f]{64}', str(asset.get('asset_id', ''))):
            raise ValueError("本次图片资源记录缺少有效asset ID")
        key = asset['asset_id']
        if key in asset_map and asset_map[key] != asset:
            raise ValueError("同一asset ID存在冲突的图片记录")
        asset_map[key] = asset
    used_assets = []
    if final and review_rules.active('export_basic_fields'):
        required_fields = (("项目名称", project.get("name") or project.get("title")), ("项目编号", project.get("project_number") or project.get("number")), ("采购人", project.get("buyer")), ("投标人法定名称", company.get("name") or company.get("company_name") or project.get("company_name")))
        missing = [label for label, value in required_fields if not _clean(value) or re.search(r"【待|待(?:补充|填写|确认)", _clean(value))]
        if missing:
            raise ValueError("正式导出被阻止：请补全并核对" + "、".join(missing) + "。")
    requirements = copy.deepcopy(requirements)
    sections = copy.deepcopy(sections)
    internal_notes = []
    source_notes = []
    def clean_internal_columns(text, owner):
        if project.get('_hide_inline_evidence') and (project.get('metadata') or {}).get('generation_profile') == 'technical_proposal':
            cleaned = _drop_internal_evidence_columns(text)
            source_notes.extend({**owner, **item} for item in cleaned['removed'])
            return cleaned['content']
        return text
    requirement_ids = [r.get('id') for r in requirements if r.get('id')]
    for section in sections:
        cleaned=clean_export_locators(clean_internal_columns(section.get('content') or section.get('text') or '',
                                    {'section_id':section.get('id')}),requirement_ids)
        section['content']=cleaned['content']
        source_notes.extend({'section_id':section.get('id'),'title':section.get('title'),**item} for item in cleaned['removed'])
        if section.get('title') and not section.get('outline_group_id'):
            section['title']=clean_export_locators(section['title'],requirement_ids)['content']
    for requirement in requirements:
        cleaned=clean_export_locators(clean_internal_columns(requirement.get('response') or '',
                                    {'requirement_id':requirement.get('id')}),requirement_ids)
        requirement['response']=cleaned['content']
        source_notes.extend({'requirement_id':requirement.get('id'),**item} for item in cleaned['removed'])
    if review_rules.active('body_separation'):
        for section in sections:
            split = bid_body.separate(section.get('content') or section.get('text') or '')
            section['content'] = split['content']
            internal_notes.extend({'title':section.get('title',''), **i} for i in split['items'])
        for requirement in requirements:
            split = bid_body.separate(requirement.get('response') or '')
            requirement['response'] = split['content']
            internal_notes.extend({'title':requirement.get('title',''), **i} for i in split['items'])
    requirement_labels = project.get("_requirement_display_labels") or _requirement_display_labels(requirements)
    replace_requirement_references = _requirement_reference_replacer(requirement_labels)
    evidence_labels = {}
    for item in company.get("evidence") or []:
        if isinstance(item, dict) and item.get("id"):
            source_id = str(item["id"])
            if source_id not in evidence_labels:
                evidence_labels[source_id] = f"资料{len(evidence_labels) + 1}"
            item["_display_label"] = evidence_labels[source_id]
    unknown_refs = set()
    def display_reference(source_id):
        source_id = str(source_id)
        if source_id not in evidence_labels:
            unknown_refs.add(source_id)
            return ""
        return evidence_labels[source_id]
    # Removed columns still participate in the normal unknown-reference check.
    for note in source_notes:
        if note.get('kind') == 'internal_evidence_column':
            for cell in note['raw_cells']:
                for source_id in re.findall(r"\[E:([A-Za-z0-9_-]+)\]", cell):
                    display_reference(source_id)
    def replace_references(text):
        if project.get('_hide_inline_evidence'):
            def consume(match):
                display_reference(match.group(1))
                return ''
            return replace_requirement_references(re.sub(r"\[E:([A-Za-z0-9_-]+)\]",consume,_clean(text)))
        return replace_requirement_references(re.sub(r"\[E:([A-Za-z0-9_-]+)\]", lambda m: "（" + display_reference(m.group(1)) + "）", _clean(text)))
    for section in sections:
        section["content"] = replace_references(section.get("content") or section.get("text"))
        if (section.get("title") or section.get("name")) and not section.get('outline_group_id'):
            section["title"] = replace_references(section.get("title") or section.get("name"))
    for requirement in requirements:
        requirement["response"] = replace_references(requirement.get("response"))
        requirement["_display_requirement_label"] = requirement_labels.get(str(requirement.get("id") or ""), "")
        requirement["_evidence_labels"] = [display_reference(item) for item in requirement.get("evidence_ids", [])]
    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.suffix.lower() != ".docx":
        raise ValueError("Word 导出路径必须以 .docx 结尾。")
    doc = Document()
    for style in doc.styles:
        if style.type == 1:
            style.font.name = "宋体"
            style.font.color.rgb = RGBColor(0, 0, 0)
            style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
            for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
                style._element.rPr.rFonts.attrib.pop(qn('w:'+key),None)
            # The bundled base template contains a theme-blue Title border.
            # A formal cover uses black typography and whitespace only.
            for border in style._element.xpath(".//w:pBdr"):
                border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.font.bold = False
    normal_cs = OxmlElement('w:szCs');normal_cs.set(qn('w:val'),'24');normal._element.get_or_add_rPr().append(normal_cs)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Pt(24)
    for name, size in [("Title", 26), ("Subtitle", 28)] + [(f"Heading {level}",18 if level == 1 else 15) for level in range(1,10)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        rpr=style._element.get_or_add_rPr()
        for el in rpr.findall(qn('w:szCs')):rpr.remove(el)
        complex_size=OxmlElement('w:szCs');complex_size.set(qn('w:val'),str(size*2));rpr.append(complex_size)
        family = "宋体" if name in ("Title","Subtitle") else "黑体"
        style.font.name = family
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), family)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.left_indent = Pt(0)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
    numbering_id = _export_numbering(doc)
    page = doc.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin, page.bottom_margin, page.left_margin, page.right_margin = Cm(2.3), Cm(2.1), Cm(2.7), Cm(2.3)
    page.header_distance, page.footer_distance = Cm(1.1), Cm(1.0)
    page.different_first_page_header_footer = True
    cover = project.get("_cover_identity") or cover_identity(project)
    title = cover["project_name"]
    bidder = _clean(company.get("name") or company.get("company_name") or project.get("company_name") or "________________")
    header = page.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.first_line_indent = Pt(0)
    _font(header.add_run(title[:70]), 9)
    footer = page.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    _font(footer.add_run("第 "), 9)
    _field(footer, " PAGE ", "1")
    _font(footer.add_run(" 页 / 共 "), 9)
    _field(footer, " NUMPAGES ", "1")
    _font(footer.add_run(" 页"), 9)
    _write_cover(doc, project, company, cover, bidder, final)
    from docx.enum.section import WD_SECTION_START
    # The reference cover has its own margins; wide body tables retain their
    # existing 16 cm text area rather than spilling outside the new cover width.
    page.top_margin, page.bottom_margin = Cm(2.54), Cm(2.54)
    page.left_margin, page.right_margin = Cm(3.17), Cm(3.17)
    body_page = doc.add_section(WD_SECTION_START.NEW_PAGE)
    body_page.top_margin, body_page.bottom_margin = Cm(2.3), Cm(2.1)
    body_page.left_margin, body_page.right_margin = Cm(2.7), Cm(2.3)
    body_page.different_first_page_header_footer = False
    body_start = len(doc.paragraphs)
    toc_title = doc.add_paragraph("目录")
    toc_title.paragraph_format.first_line_indent = Pt(0)
    _font(toc_title.runs[0], 18, True, "黑体")
    toc_paragraph = doc.add_paragraph()
    toc_paragraph.paragraph_format.space_before = Pt(0)
    toc_paragraph.paragraph_format.space_after = Pt(0)
    # Word leaves an empty field-ending paragraph after expanding the TOC.
    # Shrink only its paragraph mark so it cannot overflow onto a blank page;
    # the visible TOC entries retain their own normal text styles.
    toc_mark = OxmlElement("w:rPr")
    for property_name in ("w:sz", "w:szCs"):
        size = OxmlElement(property_name)
        size.set(qn("w:val"), "2")  # half-points: one-point paragraph mark
        toc_mark.append(size)
    toc_paragraph._p.get_or_add_pPr().append(toc_mark)
    toc_levels = '1-4' if project.get('metadata', {}).get('outline_selection') else '1-3'
    _field(toc_paragraph, ' TOC \\o "' + toc_levels + '" \\h \\z \\u ', "在 Word 中更新域以生成目录与页码")
    if not sections:
        first_heading = doc.add_heading("投标文件正文", 1)
        first_heading.paragraph_format.page_break_before = True
        doc.add_paragraph("")
    from .export_outline import outline_sections, part_body
    outline = outline_sections(sections, project.get('_split_section_families',()))
    directory_labels=[]
    for index, group in enumerate(outline, 1):
        chapter_heading = doc.add_heading(group['title'] or f'第 {index} 章', 1)
        if group.get('id'):directory_labels.append(chapter_heading._p)
        if index == 1:
            chapter_heading.paragraph_format.page_break_before = True
        for part in group['parts']:
            if group['grouped']:
                topic=doc.add_heading(part['title'], 2)
                if group.get('id'):directory_labels.append(topic._p)
            from .chapter_outline import structure_body
            body = structure_body(part['section'], part_body(part)) if group['grouped'] else part_body(part)
            used = _write_markdown(doc, body, base_level=3 if group['grouped'] else 2,
                                   assets=asset_map, asset_root=company.get('asset_root'),proposal_layout=project.get('metadata',{}).get('generation_profile')=='technical_proposal',
                                   keep_short_forms=project.get('_proposal_volume')=='qualification')
            used_assets.extend({**item, 'section_id':part['section'].get('id')} for item in used)
    if project.get("_deviation_tables"):
        _append_deviation_tables(doc, requirements, project, company)
    elif not project.get("_omit_response_table"):
        doc.add_heading("逐条招标要求响应表", 1)
        doc.add_paragraph("下表按已提取招标条款逐项编排。“要求N”为本稿响应台账标识，各分册使用同一编号；采购文件条目号及原文位置另列。响应意见、偏离情况及证明材料应由投标负责人逐条复核。").paragraph_format.keep_with_next = True
        table = doc.add_table(rows=1, cols=4)
        for cell, text in zip(table.rows[0].cells, ("序号", "招标要求及来源", "投标响应及偏离", "内容来源")):
            cell.text = text
        for index, requirement in enumerate(requirements, 1):
            row = table.add_row().cells
            row[0].text = str(index)
            raw = _clean(requirement.get("text"))
            source = _clean(requirement.get("source_locator") or requirement.get("locator"))
            label = requirement.get("_display_requirement_label")
            row[1].text = (label + "\n" if label else "") + raw + (f"\n原文位置：{source}" if source else "")
            response = _clean(requirement.get("response")) or "________________"
            row[2].text = response + "\n偏离：" + _deviation_value(requirement)
            refs = requirement.get("_evidence_labels") or []
            row[3].text = "\n".join(refs)
        if not requirements:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text, row[3].text = "1", "________________", "________________", "________________"
        _style_table(table, [1.0, 6.0, 6.0, 3.0])
        doc.add_paragraph()
    if not project.get('_omit_evidence_index'):_append_evidence_index(doc, company)
    if not project.get("_omit_attachments"):
        doc.add_heading("附件清单", 1)
        attachments = company.get("attachments") or project.get("attachments") or []
        if attachments:
            for index, attachment in enumerate(attachments, 1):
                if isinstance(attachment, dict):
                    name = _clean(attachment.get("name") or attachment.get("title") or "附件")
                    note = _clean(attachment.get("note") or "")
                else:
                    name, note = _clean(attachment), ""
                doc.add_paragraph(f"{index}. {name}" + (f"（{note}）" if note else ""))
        else:
            doc.add_paragraph("________________")
    _format_body(doc, body_start, numbering_id,directory_labels)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings._element.append(update)
    doc.core_properties.title = title + " 投标文件"
    doc.core_properties.author = bidder if not bidder.startswith("【") else ""
    doc.core_properties.subject = "投标文件"
    doc.core_properties.keywords = "招投标"
    temp = output.with_name(output.name + ".tmp")
    try:
        doc.save(temp)
        os.replace(temp, output)
    finally:
        temp.unlink(missing_ok=True)
    structure_notice = ("按本项目采购目录和已有表单结构导出；实际填写内容、签章及附件完整性仍需核对。"
                        if (project.get('metadata') or {}).get('generation_profile') == 'technical_proposal'
                        else "采用通用投标章节结构，未自动回填招标方固定表单；请核对指定格式、签章和附件装订。")
    warnings = [structure_notice, "DOCX 目录页码需在 Word 中更新域；PDF 导出会对同份文档更新目录后渲染。"]
    disabled = [rule for rule in ('export_basic_fields', 'body_separation') if not review_rules.active(rule)]
    if disabled:
        warnings.append("本次按当前审核设置导出，以下检查或处理已停用，未据此认定其已通过：" + "、".join(disabled))
    if not final:
        warnings.append("当前为草稿，含待补充或待审核内容，不可直接提交。")
    if unknown_refs:
        warnings.append(f"有 {len(unknown_refs)} 个引用未找到对应企业来源，已记录为导出警告，不能据此确认响应事实。")
    if final and any("【待" in p.text for p in doc.paragraphs):
        warnings.append("正式版仍存在待补充内容；导出并不代表满足全部招标条件。")
    return {"path": str(output), "warnings": warnings, "internal_notes": internal_notes, "source_notes": source_notes, "outline": [{"title":g["title"],"grouped":g["grouped"],"parts":[{"section_id":p["section"].get("id"),"title":p["title"]} for p in g["parts"]]} for g in outline], "disabled_rules": disabled, "cover_identity": cover, "used_assets": used_assets}


_BID_PRICE = re.compile(r"(?:含税(?:投标)?报价|投标(?:含税)?总价|投标报价|报价总额|报价金额|项目报价)\s*[:：为]?\s*[￥¥]?\s*([0-9][0-9,]*(?:\.[0-9]{1,2})?)\s*(万元|元)?")


def amount_in_chinese(value) -> str:
    """Deterministic uppercase CNY amount, without binary float rounding."""
    from decimal import Decimal, InvalidOperation
    try:
        amount = Decimal(str(value).replace(",", ""))
    except (InvalidOperation, TypeError):
        raise ValueError("报价金额无法转换为人民币大写。") from None
    if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -2:
        raise ValueError("报价金额须为非负数，最多两位小数。")
    if amount >= Decimal("10000000000000000"):
        raise ValueError("报价金额超过人民币大写转换范围。")
    digits = "零壹贰叁肆伍陆柒捌玖"
    integer = int(amount)

    def group_text(number):
        text = ""
        pending_zero = False
        for power, unit in ((1000, "仟"), (100, "佰"), (10, "拾"), (1, "")):
            digit, number = divmod(number, power)
            if digit:
                text += ("零" if pending_zero else "") + digits[digit] + unit
                pending_zero = False
            elif text:
                pending_zero = True
        return text

    groups = []
    remaining = integer
    while remaining:
        remaining, group = divmod(remaining, 10000)
        groups.append(group)
    text = ""
    skipped_group = False
    for index in range(len(groups) - 1, -1, -1):
        group = groups[index]
        if not group:
            if text:
                skipped_group = True
            continue
        if text and (skipped_group or group < 1000):
            text += "零"
        text += group_text(group) + ("", "万", "亿", "兆")[index]
        skipped_group = False
    text = (text or "零") + "元"
    fraction = int((amount - integer) * 100)
    jiao, fen = divmod(fraction, 10)
    if not fraction:
        return text + "整"
    if jiao:
        text += digits[jiao] + "角"
    elif integer and fen:
        text += "零"
    if fen:
        text += digits[fen] + "分"
    return text


def _confirmed_quotation(project, company, pricing_sections):
    from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
    quotation = company.get("quotation") or project.get("quotation") or {}
    check_decimal = review_rules.active('quote_decimal')
    check_arithmetic = review_rules.active('quote_arithmetic')
    check_uppercase = review_rules.active('quote_uppercase')

    def uppercase_amount(value, supplied=''):
        # Turning a check off never authorizes rounding a price or replacing
        # the user's uppercase text. Unsupported conversion stays blank.
        representable = (value >= 0 and value.as_tuple().exponent >= -2
                         and value < Decimal('10000000000000000'))
        computed = amount_in_chinese(value) if check_decimal or representable else ''
        if check_uppercase and supplied and computed and supplied != computed:
            raise ValueError("报价人民币大写与含税总金额不一致，请复核。")
        return supplied if supplied and not check_uppercase else computed or supplied

    if isinstance(quotation, dict) and quotation.get("confirmed") is True:
        raw = quotation.get("total_including_tax", quotation.get("total"))
        if raw in (None, '') and not review_rules.active('package_quotation'):
            return None
        try:
            value = Decimal(str(raw).replace(",", ""))
            if not value.is_finite():
                raise ValueError("报价无法执行金额处理：含税总金额不是有限数字。")
            if check_decimal and (value < 0 or value.as_tuple().exponent < -2):
                raise ValueError("含税报价须为非负金额，最多两位小数。")
            if check_decimal and value >= Decimal('10000000000000000'):
                raise ValueError("报价金额超过人民币大写转换范围。")
            items = quotation.get("items") or []
            if items:
                total = Decimal("0")
                for item in items:
                    if not isinstance(item, dict):
                        raise ValueError("报价明细格式无效。")
                    quantity = Decimal(str(item.get("quantity", "")).replace(",", ""))
                    unit_price = Decimal(str(item.get("unit_price", "")).replace(",", ""))
                    subtotal = Decimal(str(item.get("subtotal", "")).replace(",", ""))
                    if any(not amount.is_finite() for amount in (quantity, unit_price, subtotal)):
                        raise ValueError("报价无法执行金额处理：明细数量、单价或小计不是有限数字。")
                    if check_decimal and any(amount < 0 for amount in (quantity, unit_price, subtotal)):
                        raise ValueError("报价明细的数量、单价和小计须为非负数。")
                    if check_decimal and any(amount.as_tuple().exponent < -2 for amount in (unit_price, subtotal)):
                        raise ValueError("报价明细的单价和小计最多保留两位小数。")
                    if check_arithmetic and (quantity * unit_price).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) != subtotal:
                        raise ValueError("报价明细数量乘单价与小计不一致，请复核。")
                    total += subtotal
                if check_arithmetic and total != value:
                    raise ValueError("报价明细合计与含税总报价不一致，请复核。")
            original_uppercase = _clean(quotation.get("total_uppercase"))
            supplied_uppercase = original_uppercase.removeprefix("人民币").replace(" ", "")
            uppercase = uppercase_amount(value, supplied_uppercase)
            if not check_uppercase and original_uppercase:
                uppercase = original_uppercase
            return {**quotation, "total_including_tax": format(value, "f"), "total_uppercase": uppercase, "confirmation": "企业填写并确认的报价数据"}
        except (InvalidOperation, TypeError):
            raise ValueError("已确认报价缺少有效的含税总金额。")
    for section in pricing_sections:
        text = _clean(section.get("content") or section.get("text"))
        if section.get("status") != "approved" or not section.get("user_edited") or re.search(r"待(?:补充|填写|确认)|【待", text):
            continue
        match = _BID_PRICE.search(text)
        # Explicit currency unit is required for free-text confirmation.
        if match and match.group(2):
            amount = Decimal(match.group(1).replace(",", "")) * (10000 if match.group(2) == "万元" else 1)
            return {"total_including_tax": format(amount, "f"), "total_uppercase": uppercase_amount(amount), "confirmed": True, "confirmation": "人工编辑并审核通过的报价章节"}
    return None


def _quotation_item_name(requirements: list) -> dict:
    """Use only an explicit, source-verified tender value for the price form."""
    candidates = {}
    field = r"(?:报价项目|报价标的(?:名称)?|报价名称)"
    quoted_value = re.compile(field + r"\s*(?:应(?:当)?填写|须填写|填写为|名称为|为)\s*[：:]?\s*[“\"‘']([^”\"’'\n]{2,150})[”\"’']")
    for requirement in requirements:
        if (requirement.get("category") not in {"pricing", "报价"} or requirement.get("verified") not in (1, True)
                or not requirement.get("chunk_id") or not (requirement.get("locator") or requirement.get("source_locator"))):
            continue
        quote = _clean(requirement.get("quote"))
        if re.search(r"例如|示例|样例|仅供参考", quote):
            continue
        context = _clean(requirement.get("title")) + "\n" + _clean(requirement.get("text"))
        values = quoted_value.findall(quote)
        # The supplied price form contains a fixed first cell and two empty
        # amount/note cells. Context identifies the field, never supplies its value.
        if not values and re.search(field, context) and "表格" in str(requirement.get("locator") or requirement.get("source_locator")) and "\n" not in quote:
            cells = [cell.strip() for cell in quote.split("|")]
            if quote.startswith("|") and quote.endswith("|"):
                cells = cells[1:-1]
            if len(cells) == 3 and cells[0] and not any(cells[1:]):
                values = [cells[0]]
        for value in values:
            if (not 2 <= len(value) <= 150 or re.search(r"[\r\n|]|【待|待填写|待确认|待补充", value)
                    or value in {"报价项目", "报价标的", "报价名称"}):
                continue
            source = {key: requirement.get(key) for key in ("id", "chunk_id", "quote", "locator")}
            source["locator"] = requirement.get("locator") or requirement.get("source_locator")
            candidates.setdefault(value, []).append(source)
    if len(candidates) == 1:
        name = next(iter(candidates))
        return {"status": "source_verified", "name": name, "sources": candidates[name]}
    if candidates:
        return {"status": "conflict", "name": "【待填写 核对招标文件中的报价项目名称】",
                "candidates": [{"name": name, "sources": sources} for name, sources in candidates.items()]}
    return {"status": "not_found", "name": None, "sources": []}


def compose_bid_package(project: dict, requirements: list, sections: list, output_path: str, company: dict | None = None, final: bool = False) -> dict:
    """Generate separate qualification, commercial/technical and price volumes.

    This implements the supplied tender's volume separation and observed
    deviation-table schemas. It is not an arbitrary DOCX template round-trip.
    No credentials, platform submission, signature, or corporate seal is made.
    """
    import copy
    from .export_outline import split_families
    families=split_families(sections)
    company = copy.deepcopy(company or {})
    requirement_labels = _requirement_display_labels(requirements)
    output = Path(output_path).expanduser().resolve()
    if output.suffix.lower() != ".zip":
        raise ValueError("整套文件输出路径必须以 .zip 结尾。")
    output.parent.mkdir(parents=True, exist_ok=True)
    buckets = {"qualification": [], "technical": [], "pricing": []}
    req_buckets = {"qualification": [], "technical": [], "pricing": []}
    req_categories = {}
    for requirement in requirements:
        category = requirement.get("category", "technical")
        bucket = "pricing" if category in {"pricing", "报价"} else "qualification" if category in {"qualification", "资格"} else "technical"
        if _BID_PRICE.search(_clean(requirement.get("response"))):
            bucket = "pricing"
        req_buckets[bucket].append(requirement)
        req_categories[str(requirement.get("id", ""))] = bucket
    warnings = ["整套文件按资格、商务技术、报价分卷；正式提交前需按采购文件及电子交易平台要求编排、签章和加密。", "商务/技术偏离表使用内置结构和表头；并非任意采购模板的原样回填，请按本项目采购格式复核。"]
    moved_price_sections = []
    package_internal_notes = []
    for section in sections:
        title = _clean(section.get("title") or section.get("name"))
        content = _clean(section.get("content") or section.get("text"))
        ids = section.get("requirement_ids") or []
        categories = {req_categories.get(str(item)) for item in ids} if isinstance(ids, list) else set()
        if "pricing" in categories or re.search(r"报价|价格文件|费用说明", title) or _BID_PRICE.search(content):
            bucket = "pricing"
            if not re.search(r"报价|价格文件|费用说明", title) and "pricing" not in categories:
                moved_price_sections.append(title)
        elif categories == {"qualification"} or re.search(r"^企业资格|^资格|响应函|法定代表人.*授权", title):
            bucket = "qualification"
        else:
            bucket = "technical"
        buckets[bucket].append(copy.deepcopy(section))
    if moved_price_sections:
        warnings.append("检测到明确报价金额，相关章节已移入报价分册以避免进入商务技术文件：" + "、".join(moved_price_sections))
    quotation = _confirmed_quotation(project, company, buckets["pricing"])
    if final and review_rules.active('package_quotation') and quotation is None:
        raise ValueError("正式整套导出被阻止：缺少企业确认的报价。请人工填写并审核报价章节（明确含税报价金额和元/万元单位），不得使用招标预算或模型推断金额。")
    if quotation is None:
        warnings.append("报价分册尚无企业确认的金额，报价表保留待填写，不能作为正式报价提交。")
    if not buckets["qualification"]:
        buckets["qualification"] = [{"title": "资格响应文件", "content": "【待补充 按招标文件编制响应函、法定代表人或负责人授权书、营业执照及资格证明材料；涉及声明、授权、信用记录和签章必须逐项核验】"}]
    if not buckets["technical"]:
        buckets["technical"] = [{"title": "商务技术方案", "content": "【待补充 请完成商务条款、技术方案、实施交付和运维服务章节编制】"}]
    # A disabled quotation gate permits an empty form; it never derives a price
    # from the buyer's budget or an unconfirmed model-generated paragraph.
    empty_amount = "【待填写 经企业负责人确认后填写】" if review_rules.active('package_quotation') else "________________"
    amount = quotation["total_including_tax"] if quotation else empty_amount
    item_source = _quotation_item_name(requirements)
    if item_source["status"] == "conflict":
        warnings.append("报价项目名称在已核招标来源中存在冲突，报价一览表保留待填写，请核对原文及澄清文件。")
    elif item_source["status"] == "not_found":
        warnings.append("未找到经原文校验的指定报价项目名称，报价一览表暂用项目名称；请按招标格式核对。")
    item_name = item_source["name"] or _clean(project.get("name") or "项目")
    price_table = "| 报价项目 | 含税报价（元） | 备注 |\n| --- | --- | --- |\n| " + item_name.replace("|", "／") + " | " + amount + " | " + ("企业已确认金额" if quotation else "待企业确认") + " |"
    uppercase_line = "人民币大写：" + (quotation["total_uppercase"] if quotation else empty_amount)
    overview = {"title": "附件16 初次报价一览表", "content": price_table + "\n\n" + uppercase_line + "\n\n供应商：" + _clean(company.get("name") or company.get("company_name") or "________________") + "（盖章）\n\n法定代表人/负责人或其授权委托人：________________（签字）\n\n日期：________________"}
    buckets["pricing"].insert(0, overview)
    if quotation and quotation.get("items"):
        rows = ["| 费用项目 | 数量 | 单价（元） | 含税小计（元） | 说明 |", "| --- | --- | --- | --- | --- |"]
        for item in quotation["items"]:
            if not isinstance(item, dict):
                continue
            values = [item.get("name"), item.get("quantity"), item.get("unit_price"), item.get("subtotal"), item.get("note", "")]
            rows.append("| " + " | ".join(_clean(value if value is not None else "【待填写】").replace("|", "／") for value in values) + " |")
        buckets["pricing"].append({"title": "报价明细表", "content": "\n".join(rows)})
    elif not any("报价明细" in (_clean(s.get("title")) + _clean(s.get("content"))) for s in buckets["pricing"]):
        empty_item = "【待填写】" if review_rules.active('package_quotation') else "________________"
        buckets["pricing"].append({"title": "报价明细表", "content": "| 费用项目 | 数量 | 单价（元） | 含税小计（元） | 说明 |\n| --- | --- | --- | --- | --- |\n| " + " | ".join([empty_item] * 4) + " | 请依据经企业确认的报价方案填写 |"})
    if req_buckets["pricing"]:
        buckets["pricing"].append({"title": "报价要求核对", "content": "\n\n".join("### " + requirement_labels.get(str(r.get("id") or ""), "采购条款") + "\n\n招标要求：" + _clean(r.get("text")) + ("\n\n原文位置：" + _clean(r.get("source_locator") or r.get("locator")) if r.get("source_locator") or r.get("locator") else "") + "\n\n响应：" + (_clean(r.get("response")) or "________________") for r in req_buckets["pricing"])})
    if final and review_rules.active('package_quotation') and any(re.search(r"【待|待(?:补充|填写|确认)", _clean(s.get("content"))) for s in buckets["pricing"]):
        raise ValueError("正式整套导出被阻止：报价文件仍含待填写内容，请完成报价明细与报价要求响应后再导出。")
    records = []
    disabled = [rule for rule in ('export_basic_fields', 'body_separation', 'quote_decimal', 'quote_arithmetic',
                                 'quote_uppercase', 'package_quotation') if not review_rules.active(rule)]
    if disabled:
        warnings.append("本次按当前审核设置导出，停用的检查不代表已通过：" + "、".join(disabled))
    with tempfile.TemporaryDirectory(prefix="bidding-package-") as tempdir:
        root = Path(tempdir)
        for bucket, file_name, volume in (("qualification", "01_资格文件.docx", "资格文件"), ("technical", "02_商务技术文件.docx", "商务、技术文件"), ("pricing", "03_报价文件.docx", "报价文件")):
            volume_project = {**project, "document_title": volume, "_deviation_tables": bucket == "technical", "_omit_response_table": bucket == "pricing", "_omit_attachments": bucket == "pricing", "_requirement_display_labels": requirement_labels, '_split_section_families':sorted(families)}
            volume_company = copy.deepcopy(company)
            refs = {str(ref) for r in req_buckets[bucket] for ref in (r.get("evidence_ids") or [])}
            refs.update(str(ref) for section in buckets[bucket] for ref in (section.get("evidence_ids") or []))
            volume_company["evidence"] = [item for item in company.get("evidence", []) if isinstance(item, dict) and str(item.get("id")) in refs]
            result = compose_docx(volume_project, req_buckets[bucket], buckets[bucket], str(root / file_name), volume_company, final=final)
            package_internal_notes.extend(result.get("internal_notes", []))
            for warning in result["warnings"]:
                if warning not in warnings:
                    warnings.append(warning)
            records.append({"file": file_name, "volume": volume, "requirement_count": len(req_buckets[bucket]), "section_count": len(buckets[bucket]), "sha256": hashlib.sha256((root / file_name).read_bytes()).hexdigest()})
        final_status = "按当前设置导出，部分检查已停用，仍需核对最终提交文件" if disabled else "已通过导出前检查，仍需核对最终提交文件"
        checklist = ["# 内部审阅清单", "此清单用于企业内部复核，不属于对外提交投标正文。", "状态：" + (final_status if final else "草稿 待审核 不可直接提交"), "## 分册文件", *[f"- {r['file']}：{r['requirement_count']} 条要求，{r['section_count']} 个章节。" for r in records], "## 提交前核对", "- 按当前有效招标文件及澄清文件核对分册组成、附件顺序、正副本、签章和有效期。", "- 资格声明、授权文件、营业执照、证书及人员业绩证明须来自本企业真实资料。", "- 商务与技术偏离逐条确认；未填偏离不能由软件自动视为无偏离。", "- 确认报价总额、明细、税率、费用范围和报价有效期一致，并复核最高限价。", "- 报价文件独立提交，复核商务技术文件中没有误入的报价金额。", "- 在 Word 中更新目录和域，逐页查看表格跨页、空页、签章及附件。", "- 本程序未执行交易平台加密、电子签章、盖章或递交；使用采购文件指定客户端完成。", "## 生成提醒", *["- " + w for w in warnings]]
        if package_internal_notes:
            checklist.extend(["## 从正文分离的内部待办", *["- " + i["title"] + "：" + i["message"] for i in package_internal_notes]])
        checks = company.get("checks") or []
        if company.get("evidence"):
            checklist.extend(["## 企业证据内部索引", *[f"- [E:{item.get('id', '')}] {_clean(item.get('document_name') or item.get('name'))} / {_clean(item.get('locator'))}" for item in company["evidence"] if isinstance(item, dict)]])
        if checks:
            checklist.extend(["## 当前审查发现", *["- " + _clean(item.get("message")) for item in checks if isinstance(item, dict)]])
        (root / "内部审阅清单.md").write_text("\n\n".join(checklist), encoding="utf-8")
        manifest = {"product": "招投标", "project": _clean(project.get("name")), "final": final, "files": records, "warnings": warnings, "quotation_confirmed": quotation is not None, "quotation_item_source": item_source, "template_basis": "内置响应文件格式；结构与表头适配，不是原样模板回填"}
        manifest['disabled_rules'] = disabled
        (root / "导出记录.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        staged = output.with_name(output.name + ".tmp")
        try:
            with zipfile.ZipFile(staged, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for item in sorted(root.iterdir()):
                    archive.write(item, item.name)
            os.replace(staged, output)
        finally:
            staged.unlink(missing_ok=True)
    return {"path": str(output), "warnings": warnings}


@contextmanager
def _render_workspace(result):
    workspace = tempfile.TemporaryDirectory(prefix="bidding-render-")
    try:
        yield workspace.name
    finally:
        try:
            workspace.cleanup()
        except OSError as exc:
            # Word may still be releasing file handles after its owned process
            # exits. Cleanup must never replace the actual render outcome.
            result["warnings"].append(f"转换临时目录暂未清理（{type(exc).__name__}），已保留供排查：{workspace.name}")


def _stop_owned_render_word(powershell, owned_word):
    """Stop only a recorded Word identity, never by image name or PID alone."""
    try:
        record = owned_word.read_text(encoding="ascii").strip()
    except (OSError, UnicodeError):
        return "owner_record_unavailable"
    match = re.fullmatch(r"([1-9][0-9]{0,9}),([1-9][0-9]{16,19})", record)
    if not match or int(match[1]) > 2147483647 or int(match[2]) > 9223372036854775807:
        return "owner_record_invalid"
    cleanup = r'''$ErrorActionPreference = 'Stop'
$record = $env:BIDDING_RENDER_IDENTITY.Split(',')
$owned = Get-Process -Id ([int]$record[0]) -ErrorAction SilentlyContinue
if ($null -eq $owned) { Write-Output 'not_running'; exit 0 }
if ($owned.ProcessName -ne 'WINWORD' -or $owned.StartTime.ToUniversalTime().Ticks.ToString() -ne $record[1]) {
    Write-Output 'identity_mismatch'; exit 0
}
Stop-Process -InputObject $owned -Force
if ($owned.WaitForExit(10000)) { Write-Output 'stopped' } else { Write-Output 'exit_pending' }
'''
    try:
        completed = subprocess.run([str(powershell), "-NoProfile", "-NonInteractive", "-EncodedCommand", base64.b64encode(cleanup.encode("utf-16le")).decode()],
                                   env=dict(os.environ, BIDDING_RENDER_IDENTITY=record), capture_output=True, timeout=15,
                                   creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        status = completed.stdout.decode("ascii", errors="replace").strip()
        return status if not completed.returncode and status in {"not_running", "identity_mismatch", "stopped", "exit_pending"} else "cleanup_unverified"
    except (OSError, subprocess.TimeoutExpired):
        return "cleanup_unverified"


def render_pdf(docx_path: str, output_path: str, updated_docx_path: str | None = None) -> dict:
    """Render the very same DOCX using native Word, never a text-based re-layout.

    A temporary copy is opened with macros disabled and never enters MRU.
    The original DOCX and any user's already-open documents stay untouched.
    """
    source = Path(docx_path).expanduser().resolve()
    target = Path(output_path).expanduser().resolve()
    updated_target=Path(updated_docx_path).expanduser().resolve() if updated_docx_path else None
    if not source.is_file() or source.suffix.lower() != ".docx":
        raise ValueError("请提供已生成的 DOCX 文件。")
    if target.suffix.lower() != ".pdf":
        raise ValueError("PDF 输出路径必须以 .pdf 结尾。")
    if updated_target and (updated_target.suffix.lower()!='.docx' or updated_target==source):
        raise ValueError('更新域后的Word必须保存为另一个DOCX文件，不能覆盖原件或旧导出')
    target.parent.mkdir(parents=True, exist_ok=True)
    if os.name != "nt":
        return {"path": None, "warnings": ["PDF 渲染未执行：当前环境未配置受支持的 Word 渲染器。请在 Windows Word 中将同份 DOCX 导出为 PDF。"], "engine": None}
    powershell = Path(os.environ.get("SystemRoot", "C:/Windows")) / "System32/WindowsPowerShell/v1.0/powershell.exe"
    if not powershell.is_file():
        return {"path": None, "warnings": ["PDF 渲染未执行：缺少 Windows PowerShell / Microsoft Word。"], "engine": None}
    result = {"path": None, "warnings": [], "engine": "word"}
    with _RENDER_LOCK, _render_workspace(result) as tempdir:
        tmp = Path(tempdir)
        copied = tmp / "input.docx"
        pdf = tmp / "output.pdf"
        owned_word = tmp / "owned-word.txt"
        updated_copy=tmp/'updated.docx'
        pagination_record=tmp/'pagination.json'
        shutil.copy2(source, copied)
        # Paths go through environment variables, never shell interpolation.
        script = r'''$ErrorActionPreference = 'Stop'
$ProgressPreference = 'SilentlyContinue'
$word = $null
$document = $null
$businessPageFields = New-Object 'System.Collections.Generic.List[object]'
$sectionBookmarkNames = New-Object 'System.Collections.Generic.List[string]'
$updateLinksOriginal = $null
[uint32]$ownedWordPid = 0
$ownedWordStartTicks = $null
$existingWordIds = @(Get-Process -Name WINWORD -ErrorAction SilentlyContinue | Select-Object -ExpandProperty Id)
try {
    Add-Type -TypeDefinition 'using System; using System.Runtime.InteropServices; public static class BiddingWordNative { [DllImport("user32.dll")] public static extern uint GetWindowThreadProcessId(IntPtr hwnd, out uint pid); }'
    $word = New-Object -ComObject Word.Application
    # Word.Application.Hwnd is unavailable on some installed Word versions.
    # A new blank document exposes the already-verified ActiveWindow.Hwnd API,
    # allowing ownership to be recorded before opening a large source file.
    $document = $word.Documents.Add()
    [void][BiddingWordNative]::GetWindowThreadProcessId([intptr]$document.ActiveWindow.Hwnd, [ref]$ownedWordPid)
    if ($ownedWordPid -eq 0 -or $existingWordIds -contains $ownedWordPid) { throw 'Unable to establish an isolated Word rendering process.' }
    $ownedProcess = Get-Process -Id $ownedWordPid
    if ($ownedProcess.ProcessName -ne 'WINWORD') { throw 'Unexpected Word application process identity.' }
    $ownedWordStartTicks = $ownedProcess.StartTime.ToUniversalTime().Ticks.ToString()
    Set-Content -LiteralPath $env:BIDDING_RENDER_OWNER -Value ($ownedWordPid.ToString() + ',' + $ownedWordStartTicks) -Encoding ASCII
    $document.Close([ref]0)
    [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($document)
    $document = $null
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $word.AutomationSecurity = 3
    $updateLinksOriginal = $word.Options.UpdateLinksAtOpen
    $word.Options.UpdateLinksAtOpen = $false
    $document = $word.Documents.Open([ref]$env:BIDDING_RENDER_INPUT, [ref]$false, [ref]$true, [ref]$false)
    foreach ($bookmark in $document.Bookmarks) {
        if ($bookmark.Name -match '^mxs_[0-9a-f]{32}$') {
            [void]$sectionBookmarkNames.Add($bookmark.Name)
        }
        [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($bookmark)
    }
    # Snapshot only original business page references before TOC expansion.
    # Keep live Field objects; numeric collection indices change as TOC grows.
    foreach ($field in $document.Fields) {
        if ($field.Code.Text -match '^\s*PAGEREF\s+(?:mxs_[0-9a-f]{32}|"mxs_[0-9a-f]{32}")(?=\s|$)') {
            [void]$businessPageFields.Add($field)
        }
    }
    [void]$document.Fields.Update()
    foreach ($toc in $document.TablesOfContents) { [void]$toc.Update() }
    [void]$document.Repaginate()
    foreach ($toc in $document.TablesOfContents) { [void]$toc.UpdatePageNumbers() }
    # The expanded TOC can move body pages. Refresh the captured business
    # references only. TOC-owned _Toc references were updated above and must
    # not be traversed/updated again through the now-expanded Fields collection.
    foreach ($businessField in $businessPageFields) {
        [void]$businessField.Update()
    }
    # Updating references can alter index-table height. Repaginate to a fixed
    # point, bounded to four passes; never publish stale reference results.
    $paginationStable = $false
    for ($pass = 0; $pass -lt 4; $pass++) {
        [void]$document.Repaginate()
        foreach ($toc in $document.TablesOfContents) { [void]$toc.UpdatePageNumbers() }
        foreach ($businessField in $businessPageFields) { [void]$businessField.Update() }
        [void]$document.Repaginate()
        $sectionPages = @{}
        foreach ($bookmarkName in $sectionBookmarkNames) {
            $bookmark = $document.Bookmarks.Item($bookmarkName)
            $bookmarkRange = $bookmark.Range
            # Collapse to the heading start: a heading spanning pages must
            # reference its first page. 3 is wdActiveEndPageNumber (physical).
            [void]$bookmarkRange.Collapse(1)
            $page = [int]$bookmarkRange.Information(3)
            if ($page -lt 1) { throw 'Word returned an invalid section page.' }
            $sectionPages[$bookmarkName.Substring(4).ToLowerInvariant()] = $page
            [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($bookmarkRange)
            [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($bookmark)
        }
        $paginationStable = $true
        foreach ($businessField in $businessPageFields) {
            if ($businessField.Code.Text -match '^\s*PAGEREF\s+"?(mxs_[0-9a-f]{32})"?(?=\s|$)') {
                $sectionId = $Matches[1].Substring(4).ToLowerInvariant()
                $actualResult = $businessField.Result.Text.Trim()
                if (-not $sectionPages.ContainsKey($sectionId) -or $actualResult -ne [string]$sectionPages[$sectionId]) {
                    $paginationStable = $false
                }
            }
        }
        if ($paginationStable) { break }
    }
    if (-not $paginationStable) { throw 'Business page references did not converge to actual section pages.' }
    $pageCount = [int]$document.ComputeStatistics(2)
    if ($pageCount -lt 1) { throw 'Word returned an invalid document page count.' }
    if ($env:BIDDING_RENDER_UPDATED_DOCX) {
        $document.SaveAs2($env:BIDDING_RENDER_UPDATED_DOCX, 12)
    }
    $document.ExportAsFixedFormat($env:BIDDING_RENDER_OUTPUT, 17)
    @{engine='word';page_count=$pageCount;section_pages=$sectionPages} | ConvertTo-Json -Depth 5 | Set-Content -LiteralPath $env:BIDDING_RENDER_PAGINATION -Encoding UTF8
    Write-Output 'BIDDING_RENDER_OK'
} catch {
    [Console]::Error.WriteLine($_.Exception.Message)
    exit 1
} finally {
    foreach ($businessField in $businessPageFields) {
        try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($businessField) } catch { }
    }
    $businessPageFields.Clear()
    if ($null -ne $document) {
        try { $document.Close([ref]0) } catch { }
        try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($document) } catch { }
    }
    if ($null -ne $word) {
        try { if ($null -ne $updateLinksOriginal) { $word.Options.UpdateLinksAtOpen = $updateLinksOriginal } } catch { }
        $owned = Get-Process -Id $ownedWordPid -ErrorAction SilentlyContinue
        if ($null -ne $owned -and $null -ne $ownedWordStartTicks -and $owned.ProcessName -eq 'WINWORD' -and $owned.StartTime.ToUniversalTime().Ticks.ToString() -eq $ownedWordStartTicks) {
            try { $word.Quit([ref]0) } catch { }
        }
        try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($word) } catch { }
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
'''
        encoded = base64.b64encode(script.encode("utf-16le")).decode("ascii")
        env = dict(os.environ, BIDDING_RENDER_INPUT=str(copied), BIDDING_RENDER_OUTPUT=str(pdf), BIDDING_RENDER_OWNER=str(owned_word))
        env['BIDDING_RENDER_UPDATED_DOCX']=str(updated_copy) if updated_target else ''
        env['BIDDING_RENDER_PAGINATION']=str(pagination_record)
        try:
            completed = subprocess.run([str(powershell), "-NoProfile", "-NonInteractive", "-OutputFormat", "Text", "-EncodedCommand", encoded], env=env, capture_output=True, timeout=WORD_RENDER_TIMEOUT_SECONDS, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            if completed.returncode or not pdf.is_file() or pdf.stat().st_size < 100:
                detail = completed.stderr.decode("utf-8", errors="replace").strip()[:400]
                result.update(warnings=["Word PDF 渲染失败；请确认 Microsoft Word 可正常激活并打开文档。" + (f" 错误：{detail}" if detail else "")], error_type="render_failed")
                return result
            if not pdf.read_bytes().startswith(b"%PDF-"):
                result.update(warnings=["Word 返回的输出不是有效 PDF，未保存。"], error_type="invalid_pdf")
                return result
            if pagination_record.is_file():
                # A successful process status alone is not pagination proof.
                pagination=json.loads(pagination_record.read_text(encoding='utf-8-sig'))
                if not isinstance(pagination,dict):
                    result.update(warnings=['Word分页记录无效，未保存输出。'],error_type='invalid_pagination')
                    return result
                page_count=pagination.get('page_count')
                pages=pagination.get('section_pages')
                if (pagination.get('engine')!='word' or type(page_count) is not int or page_count<1
                        or not isinstance(pages,dict) or any(
                            not re.fullmatch(r'[0-9a-f]{32}',str(key)) or type(value) is not int
                            or value<1 or value>page_count for key,value in pages.items())):
                    result.update(warnings=['Word分页记录无效，未保存输出。'],error_type='invalid_pagination')
                    return result
                result.update(section_pages=pages,page_count=page_count)
            staged = target.with_name(target.name + ".tmp")
            if updated_target:
                if not updated_copy.is_file() or not zipfile.is_zipfile(updated_copy):
                    result.update(warnings=['Word没有生成有效的已更新域DOCX，未宣称目录页码已保存'],error_type='updated_docx_missing')
                    return result
                updated_target.parent.mkdir(parents=True,exist_ok=True)
                updated_staged=updated_target.with_name(updated_target.name+'.tmp')
                shutil.copyfile(updated_copy,updated_staged);os.replace(updated_staged,updated_target)
                result['updated_docx_path']=str(updated_target)
            shutil.copyfile(pdf, staged)
            os.replace(staged, target)
            result.update(path=str(target))
            return result
        except subprocess.TimeoutExpired:
            # Kill only the exact Word process created by this export, checking
            # its PID *and* start time. Never terminate a pre-existing Word.
            cleanup_status = _stop_owned_render_word(powershell, owned_word)
            result.update(warnings=[f"Word 渲染超过 {WORD_RENDER_TIMEOUT_SECONDS} 秒，PDF 未保存；大册转换超时，请检查 Word 状态后重试或手工导出。"],
                          error_type="timeout", timeout_seconds=WORD_RENDER_TIMEOUT_SECONDS, owned_process_cleanup=cleanup_status)
            if cleanup_status not in {"stopped", "not_running"}:
                result["warnings"].append("未能确认转换进程已退出；未终止任何无法核验归属的进程。")
            return result
        except (OSError, ValueError) as exc:
            result.update(warnings=[f"PDF 渲染未完成：{type(exc).__name__}：{str(exc)[:250]}"], error_type=type(exc).__name__)
            return result
