"""Synthetic export checks; never read the live project or invoke PDF/Word."""
import copy
from io import BytesIO
from zipfile import ZipFile

from docx import Document

from app.documents import _requirement_display_labels, _requirement_reference_replacer, compose_bid_package, compose_docx


A = "3c21bb66" + "a" * 24
B = "62b52d84" + "b" * 24
C = "20260910" + "c" * 24


def text_from_docx(path_or_buffer):
    doc = Document(path_or_buffer)
    return "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])


def test_known_requirement_ids_and_contextual_prefixes_are_readable():
    replace = _requirement_reference_replacer(_requirement_display_labels([{"id": A}, {"id": B}, {"id": C}]))
    text = f"### 磋商组织（对应要求 {A[:8]}）\n要求ID：{B}\n引用（{A}）。\n普通文本 {A[:8]}，日期20260910，数字 20260910。"
    result = replace(text)
    assert "（对应要求1）" in result and "要求2" in result and "引用（要求1）" in result
    assert f"普通文本 {A[:8]}" in result and "日期20260910，数字 20260910" in result
    assert replace(f"对应要求 {C[:8]}") == "对应要求3"
    assert replace(f"对应要求 {A[:8]}、对应要求 {A}") == "对应要求1、对应要求1"


def test_only_requirement_columns_allow_short_prefixes_and_hashes_stay_literal():
    replace = _requirement_reference_replacer({A: "要求8", B: "要求11"})
    raw = f"| 序号 | 对应要求 | 说明 | MD5校验值 |\n| --- | --- | --- | --- |\n| 1 | {A[:8]} | {A[:8]} | {A} |\n| 2 | {B} | 普通说明 | {B} |"
    result = replace(raw)
    assert f"| 1 | 要求8 | {A[:8]} | {A} |" in result
    assert f"| 2 | 要求11 | 普通说明 | {B} |" in result
    assert result.splitlines()[:2] == raw.splitlines()[:2]
    assert "| 1 | 要求8 |" in replace(raw.replace("对应要求", "涉及要求"))
    for value in (f"SHA256：{A}", f"MD5校验值：{A}", f"https://example.invalid/{A}", f"文件 {A}.txt", "a" + A + "b", "deadbeef" * 8):
        assert replace(value) == value


def test_ambiguous_or_unknown_prefixes_are_never_guessed():
    collision = A[:8] + "d" * 24
    replace = _requirement_reference_replacer({A: "要求1", collision: "要求2"})
    assert replace("对应要求 " + A[:8]) == "对应要求 " + A[:8]
    assert replace("对应要求 " + B[:8]) == "对应要求 " + B[:8]
    assert replace("对应要求 " + B) == "对应要求 " + B
    assert replace("对应要求 " + A) == "对应要求1"
    assert replace(A[:8] + "9") == A[:8] + "9"


def test_docx_display_does_not_mutate_source_or_input_objects(tmp_path):
    original_source = "采购原文中的校验值 " + A
    requirements = [{"id": A, "text": original_source, "locator": "表格2第5行", "response": "对应要求 " + A[:8]}]
    sections = [{"title": "要求ID：" + A, "content": "正文对应要求 " + A, "requirement_ids": [A]}]
    before = copy.deepcopy((requirements, sections))
    target = tmp_path / "synthetic.docx"
    compose_docx({"name": "合成测试"}, requirements, sections, str(target))
    text = text_from_docx(target)
    assert "正文对应要求1" in text and "要求1\n" + original_source in text
    assert "表格2第5行" in text and original_source in text
    assert text.count(A) == 1  # preserved tender source, never reinterpreted
    assert (requirements, sections) == before


def test_package_keeps_global_labels_and_original_deviation_columns(tmp_path):
    requirements = [
        {"id": A, "category": "technical", "text": "技术条款", "locator": "第1段", "response": "对应要求 " + A[:8]},
        {"id": B, "category": "qualification", "text": "资格条款", "locator": "第2段", "response": "对应要求 " + B},
        {"id": C, "category": "pricing", "text": "报价条款", "locator": "第3段", "response": "对应要求 " + C[:8]},
    ]
    sections = [{"title": "技术方案", "content": "见要求ID：" + A, "requirement_ids": [A]},
                {"title": "资格响应", "content": "对应要求 " + B[:8], "requirement_ids": [B]}]
    before = copy.deepcopy((requirements, sections))
    target = tmp_path / "synthetic.zip"
    compose_bid_package({"name": "合成分册测试"}, requirements, sections, str(target))
    with ZipFile(target) as archive:
        qualification = Document(BytesIO(archive.read("01_资格文件.docx")))
        technical = Document(BytesIO(archive.read("02_商务技术文件.docx")))
        pricing = text_from_docx(BytesIO(archive.read("03_报价文件.docx")))
        qt = text_from_docx(BytesIO(archive.read("01_资格文件.docx")))
        tt = text_from_docx(BytesIO(archive.read("02_商务技术文件.docx")))
    assert "对应要求2" in qt and "要求2\n资格条款" in qt
    assert "本稿要求标识：要求1" in tt and "见要求1" in tt
    assert "要求3" in pricing and "报价条款" in pricing and "第3段" in pricing
    assert all(id not in qt + tt + pricing for id in (A, B, C))
    response_table = next(table for table in qualification.tables if table.rows[0].cells[0].text == "序号")
    assert response_table.rows[1].cells[0].text == "1"  # local form sequence remains intact
    deviation = next(table for table in technical.tables if any(cell.text == "采购文件的技术条款" for cell in table.rows[0].cells))
    assert [cell.text for cell in deviation.rows[0].cells] == ["序号", "采购文件条目号", "采购文件的技术条款", "响应文件的技术条款响应", "偏离", "说明"]
    assert deviation.rows[1].cells[0].text == "1" and deviation.rows[1].cells[1].text == "第1段"
    assert (requirements, sections) == before
