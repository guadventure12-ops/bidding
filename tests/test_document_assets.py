"""Synthetic offline raster/provenance checks; never use business data or models."""
from copy import deepcopy
import hashlib
import io
from pathlib import Path
import zipfile

import pytest
from PIL import Image

from app import document_assets as assets


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    import socket
    def blocked(*args, **kwargs):
        raise AssertionError("Document asset extraction must stay offline")
    monkeypatch.setattr(socket, "create_connection", blocked)
    monkeypatch.setattr(socket, "getaddrinfo", blocked)
    monkeypatch.setattr(socket.socket, "connect", blocked)


def raster(fmt="PNG", size=(12, 8), color=(18, 128, 64)):
    stream = io.BytesIO()
    Image.new("RGB", size, color).save(stream, format=fmt)
    return stream.getvalue()


def extract(path, tmp_path, **kwargs):
    return assets.extract_document_assets(path, document_id="source-1", output_dir=tmp_path / "assets", **kwargs)


def rewrite_zip(path, changes):
    with zipfile.ZipFile(path) as src:
        contents = {name: src.read(name) for name in src.namelist()}
    contents.update(changes)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, value in contents.items():
            # ZipInfo normally normalizes backslashes on Windows; assign the
            # literal name to construct a genuinely malicious archive fixture.
            info = zipfile.ZipInfo("temporary")
            info.filename = name
            info.orig_filename = name
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, value)


def docx_fixture(path):
    from docx import Document
    doc = Document()
    doc.add_heading("电子档案产品功能", level=1)
    doc.add_paragraph("产品原始说明 123.45。")
    pic = doc.add_picture(io.BytesIO(raster()))
    pic._inline.docPr.set("descr", "系统架构图")
    doc.add_paragraph("图1 系统架构与审计入口", style="Caption")
    doc.save(path)
    return path


def test_docx_provenance_stable_original_untouched_and_verified_word(tmp_path):
    from docx import Document
    source = docx_fixture(tmp_path / "source.docx")
    before = source.read_bytes()
    result = extract(source, tmp_path)
    assert not result["warnings"]
    assert len(result["assets"]) == 1
    image = result["assets"][0]
    assert result["source_sha256"] == hashlib.sha256(before).hexdigest()
    assert image["document_id"] == "source-1"
    assert image["source"]["part"] == "word/document.xml"
    assert image["source"]["paragraph"] == 3
    assert image["source"]["heading"] == "电子档案产品功能"
    assert image["source"]["caption"] == "系统架构图"
    assert "123.45" in image["source"]["nearby_text"]
    assert image["source"]["page"] is None
    assert (image["width"], image["height"]) == (12, 8)
    assert "approved" not in image and "status" not in image
    assert extract(source, tmp_path) == result
    assert source.read_bytes() == before
    output = Document()
    output.add_picture(str(assets.verified_asset_path(image, asset_root=tmp_path / "assets")))
    output.save(tmp_path / "export.docx")
    loaded = Document(tmp_path / "export.docx")
    assert len(loaded.inline_shapes) == 1


def test_same_image_references_have_distinct_locations_and_stable_ids(tmp_path):
    from docx import Document
    source = tmp_path / "repeated.docx"
    doc = Document()
    doc.add_picture(io.BytesIO(raster()))
    doc.add_picture(io.BytesIO(raster()))
    doc.save(source)
    result = extract(source, tmp_path)
    a, b = result["assets"]
    assert a["sha256"] == b["sha256"]
    assert a["asset_id"] != b["asset_id"]
    assert a["source"]["paragraph"] == 1 and b["source"]["paragraph"] == 2


def test_docx_deleted_picture_not_extracted_and_crop_is_explicit(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    source = tmp_path / "visible.docx"
    doc = Document()
    live = doc.add_picture(io.BytesIO(raster()))
    fill = live._inline.graphic.graphicData.pic.blipFill
    crop = OxmlElement("a:srcRect")
    crop.set("l", "25000")
    fill.append(crop)
    deleted = doc.add_picture(io.BytesIO(raster("JPEG")))
    paragraph = deleted._inline.getparent().getparent().getparent()
    run = paragraph[0]
    paragraph.remove(run)
    deletion = OxmlElement("w:del")
    deletion.append(run)
    paragraph.append(deletion)
    doc.save(source)
    result = extract(source, tmp_path)
    assert len(result["assets"]) == 1
    assert result["assets"][0]["source"]["display_transform"]["crop"] == {"l": "25000"}
    assert any(w["code"] == "display_transform_unapplied" for w in result["warnings"])


def test_source_and_document_version_change_asset_ids(tmp_path):
    source = docx_fixture(tmp_path / "version.docx")
    old = extract(source, tmp_path)["assets"][0]
    different_document = assets.extract_document_assets(source, document_id="source-2", output_dir=tmp_path / "assets")["assets"][0]
    assert old["asset_id"] != different_document["asset_id"]
    rewrite_zip(source, {"docProps/custom.xml": b"<version>2</version>"})
    new = extract(source, tmp_path)["assets"][0]
    assert old["asset_id"] != new["asset_id"]
    assert old["sha256"] == new["sha256"]


def pptx_fixture(path):
    # Minimal OOXML fixture is intentional: no PowerPoint/UI dependency needed.
    files = {
        "ppt/presentation.xml": f'<p:presentation xmlns:p="{assets.P[1:-1]}" xmlns:r="{assets.R[1:-1]}"><p:sldIdLst><p:sldId r:id="r2"/><p:sldId r:id="r1"/></p:sldIdLst></p:presentation>',
        "ppt/_rels/presentation.xml.rels": '<Relationships><Relationship Id="r1" Type="urn:test/slide" Target="slides/slide1.xml"/><Relationship Id="r2" Type="urn:test/slide" Target="slides/slide2.xml"/></Relationships>',
        "ppt/slides/slide1.xml": f'<p:sld xmlns:p="{assets.P[1:-1]}"/>',
        "ppt/slides/slide2.xml": f'''<p:sld xmlns:p="{assets.P[1:-1]}" xmlns:a="{assets.A[1:-1]}" xmlns:r="{assets.R[1:-1]}"><p:cSld><p:spTree><p:sp><p:nvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>档案功能展示</a:t></a:r></a:p></p:txBody></p:sp><p:pic><p:nvPicPr><p:cNvPr descr="档案查询截图"/></p:nvPicPr><p:blipFill><a:blip r:embed="img1"/></p:blipFill></p:pic></p:spTree></p:cSld></p:sld>''',
        "ppt/slides/_rels/slide2.xml.rels": '<Relationships><Relationship Id="img1" Type="urn:test/image" Target="../media/image1.jpeg"/></Relationships>',
        "ppt/media/image1.jpeg": raster("JPEG"),
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return path


def test_pptx_uses_presentation_order_and_preserves_slide_heading(tmp_path):
    result = extract(pptx_fixture(tmp_path / "slides.pptx"), tmp_path)
    assert not result["warnings"]
    assert len(result["assets"]) == 1
    image = result["assets"][0]
    assert image["source"]["part"] == "ppt/slides/slide2.xml"
    assert image["source"]["slide"] == 1
    assert image["source"]["page"] is None
    assert image["source"]["heading"] == "档案功能展示"
    assert image["source"]["caption"] == "档案查询截图"
    assert image["original_mime"] == "image/jpeg" and image["mime"] == "image/png"


def test_pptx_shape_fill_image_is_not_silently_lost(tmp_path):
    source = pptx_fixture(tmp_path / "fill.pptx")
    with zipfile.ZipFile(source) as archive:
        xml = archive.read("ppt/slides/slide2.xml").decode()
    xml = xml.replace("</p:spTree>", '<p:sp><p:spPr><a:blipFill><a:blip r:embed="img1"/></a:blipFill></p:spPr></p:sp></p:spTree>')
    rewrite_zip(source, {"ppt/slides/slide2.xml": xml})
    result = extract(source, tmp_path)
    assert not result["warnings"]
    assert len(result["assets"]) == 2
    assert result["assets"][1]["source"]["kind"] == "shape_or_background_fill"


def test_pdf_only_displayed_rasters_and_actual_page_locations(tmp_path):
    import fitz
    source = tmp_path / "source.pdf"
    with fitz.open() as doc:
        first = doc.new_page()
        first.insert_text((20, 20), "Architecture overview")
        first.insert_image(fitz.Rect(20, 40, 140, 120), stream=raster())
        second = doc.new_page()
        second.insert_text((20, 20), "Expense product")
        second.insert_image(fitz.Rect(20, 40, 140, 120), stream=raster("JPEG"))
        doc.save(source)
    original = source.read_bytes()
    result = extract(source, tmp_path)
    assert not result["warnings"]
    assert [a["source"]["page"] for a in result["assets"]] == [1, 2]
    assert result["assets"][0]["source"]["bbox"] == [20.0, 40.0, 140.0, 120.0]
    assert "Architecture overview" in result["assets"][0]["source"]["nearby_text"]
    assert source.read_bytes() == original


def test_pdf_soft_mask_preserved_and_vector_artwork_warned(tmp_path):
    import fitz
    rgba = Image.new("RGBA", (10, 10), (0, 100, 200, 64))
    data = io.BytesIO()
    rgba.save(data, format="PNG")
    source = tmp_path / "alpha.pdf"
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_image(fitz.Rect(10, 10, 110, 110), stream=data.getvalue())
        page.draw_rect(fitz.Rect(20, 150, 40, 180))
        doc.save(source)
    result = extract(source, tmp_path)
    assert len(result["assets"]) == 1
    with Image.open(result["assets"][0]["path"]) as img:
        assert img.getpixel((1, 1))[3] == 64
    assert any(w["code"] == "non_raster_drawing" for w in result["warnings"])


def test_markdown_inline_reference_html_and_no_code_examples(tmp_path):
    source = tmp_path / "source.md"
    (tmp_path / "image with space.webp").write_bytes(raster("WEBP"))
    source.write_text('# 查询功能\n![查询示意](<image with space.webp>)\n![复用][shot]\n[shot]: <image with space.webp>\n<img src="image with space.webp" alt="本地示意" />\n```md\n![不是正文](missing.png)\n```\n`![也不是正文](missing2.png)`', encoding="utf-8")
    result = extract(source, tmp_path)
    assert not result["warnings"]
    assert len(result["assets"]) == 3
    assert result["assets"][0]["source"]["line"] == 2
    assert result["assets"][0]["source"]["heading"] == "查询功能"
    assert all(a["original_mime"] == "image/webp" for a in result["assets"])
    assert len({a["asset_id"] for a in result["assets"]}) == 3


def test_markdown_traversal_external_and_data_uris_not_read(tmp_path):
    child = tmp_path / "source"
    child.mkdir()
    outside = tmp_path / "outside.png"
    outside.write_bytes(raster())
    source = child / "sample.md"
    source.write_text('![外部](../outside.png)\n![网络](https://example.test/image.png)\n![内联](data:image/png;base64,xxx)\n![缺失](missing.png)\n![SVG](image.svg)', encoding="utf-8")
    (child / "image.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg"/>', encoding="utf-8")
    denied = extract(source, tmp_path)
    assert not denied["assets"]
    assert len(denied["warnings"]) == 5
    allowed = extract(source, tmp_path, allowed_image_roots=[tmp_path])
    assert len(allowed["assets"]) == 1
    assert allowed["assets"][0]["source"]["image_path"] == str(outside)
    assert len(allowed["warnings"]) == 4


@pytest.mark.parametrize("broken", ["external", "svg", "missing", "bad_relationship"])
def test_docx_bad_images_do_not_become_available_assets(tmp_path, broken):
    source = docx_fixture(tmp_path / "bad.docx")
    with zipfile.ZipFile(source) as archive:
        rels = archive.read("word/_rels/document.xml.rels").decode()
    if broken == "external":
        rels = rels.replace('Target="media/image1.png"', 'Target="https://example.test/image.png" TargetMode="External"')
        changes = {"word/_rels/document.xml.rels": rels}
    elif broken == "svg":
        changes = {"word/media/image1.png": b'<svg><script>alert(1)</script></svg>'}
    elif broken == "missing":
        changes = {"word/_rels/document.xml.rels": rels.replace("media/image1.png", "media/absent.png")}
    else:
        changes = {"word/_rels/document.xml.rels": rels.replace('/relationships/image"', '/relationships/oleObject"')}
    rewrite_zip(source, changes)
    result = extract(source, tmp_path)
    assert not result["assets"]
    assert result["incomplete"] and result["warnings"][0]["code"] == "image_unavailable"


@pytest.mark.parametrize("extra", ["../escape.png", "/absolute.png", "word/../../escape.png", "word\\media\\x.png"])
def test_zip_path_validation_rejects_archive_before_extraction(tmp_path, extra):
    source = docx_fixture(tmp_path / "unsafe.docx")
    rewrite_zip(source, {extra: raster()})
    result = extract(source, tmp_path)
    assert not result["assets"]
    assert "非法路径" in result["warnings"][0]["message"]


def test_limits_xml_entities_large_pixels_and_zip_bomb(tmp_path, monkeypatch):
    source = docx_fixture(tmp_path / "bounded.docx")
    monkeypatch.setattr(assets, "MAX_IMAGE_PIXELS", 10)
    result = extract(source, tmp_path)
    assert not result["assets"] and "图像像素" in result["warnings"][0]["message"]
    monkeypatch.setattr(assets, "MAX_IMAGE_PIXELS", 25_000_000)
    rewrite_zip(source, {"custom/bomb.txt": b"0" * 2_000_000})
    result = extract(source, tmp_path)
    assert not result["assets"] and "异常压缩比" in result["warnings"][0]["message"]
    source = docx_fixture(tmp_path / "entity.docx")
    rewrite_zip(source, {"word/document.xml": b'<!DOCTYPE root [<!ENTITY x "external">]><root>&x;</root>'})
    result = extract(source, tmp_path)
    assert not result["assets"] and "DTD" in result["warnings"][0]["message"]


def test_verifier_rejects_tampering_and_paths_outside_asset_root(tmp_path):
    source = docx_fixture(tmp_path / "source.docx")
    image = extract(source, tmp_path)["assets"][0]
    changed = deepcopy(image)
    changed["width"] = 999
    with pytest.raises(ValueError, match="尺寸"):
        assets.verified_asset_path(changed, asset_root=tmp_path / "assets")
    changed = deepcopy(image)
    changed["path"] = str(tmp_path / (image["asset_id"] + ".png"))
    Path(changed["path"]).write_bytes(Path(image["path"]).read_bytes())
    with pytest.raises(ValueError, match="资源目录"):
        assets.verified_asset_path(changed, asset_root=tmp_path / "assets")
    Path(image["path"]).write_bytes(raster(color=(255, 0, 0)))
    with pytest.raises(ValueError, match="hash"):
        assets.verified_asset_path(image, asset_root=tmp_path / "assets")
    result = extract(source, tmp_path)
    assert not result["assets"] and "未覆盖" in result["warnings"][0]["message"]


def test_asset_count_and_total_byte_caps_stop_scan_without_losing_valid_previous_asset(tmp_path, monkeypatch):
    source = tmp_path / "count.md"
    (tmp_path / "image.png").write_bytes(raster())
    source.write_text('![a](image.png)\n![b](image.png)\n![c](image.png)', encoding="utf-8")
    monkeypatch.setattr(assets, "MAX_ASSETS", 1)
    result = extract(source, tmp_path)
    assert len(result["assets"]) == 1
    assert any("引用数量" in w["message"] for w in result["warnings"])
    assert "_attempts" not in result and "_output_bytes" not in result
    monkeypatch.setattr(assets, "MAX_ASSETS", 100)
    monkeypatch.setattr(assets, "MAX_OUTPUT_BYTES", 1)
    result = extract(source, tmp_path)
    assert not result["assets"]
    assert "输出总量" in result["warnings"][0]["message"]


def test_require_explicit_document_and_output_and_no_default_production_path(tmp_path):
    source = docx_fixture(tmp_path / "source.docx")
    with pytest.raises(TypeError):
        assets.extract_document_assets(source, document_id="d")
    with pytest.raises(ValueError):
        assets.extract_document_assets(source, document_id="", output_dir=tmp_path / "assets")
    with pytest.raises(ValueError):
        assets.extract_document_assets(source, document_id="d", output_dir=source)
    with pytest.raises(ValueError, match="网络路径"):
        assets.extract_document_assets("//host/share/file.docx", document_id="d", output_dir=tmp_path / "assets")


def test_source_change_during_extraction_discards_assets(tmp_path, monkeypatch):
    source = docx_fixture(tmp_path / "source.docx")
    normal = assets._docx
    def change(*args):
        normal(*args)
        with source.open("ab") as stream:
            stream.write(b"new version")
    monkeypatch.setattr(assets, "_docx", change)
    result = extract(source, tmp_path)
    assert not result["assets"]
    assert any(w["code"] == "source_changed" for w in result["warnings"])


def test_source_change_plus_partial_parse_failure_also_discards_assets(tmp_path, monkeypatch):
    source = docx_fixture(tmp_path / "partial.docx")
    normal = assets._docx
    def change(*args):
        normal(*args)
        with source.open("ab") as stream:
            stream.write(b"new version")
        raise ValueError("later part unavailable")
    monkeypatch.setattr(assets, "_docx", change)
    result = extract(source, tmp_path)
    assert not result["assets"]
    assert {w["code"] for w in result["warnings"]} == {"source_changed", "extraction_incomplete"}


def test_png_validation_does_not_recompress_and_still_rechecks_bytes(tmp_path,monkeypatch):
    source=docx_fixture(tmp_path/'verify.docx');result=extract(source,tmp_path)
    asset=result['assets'][0];from PIL import Image
    monkeypatch.setattr(Image.Image,'save',lambda *a,**k:pytest.fail('Validation must not re-encode PNGs'))
    path=assets.verified_asset_path(asset,asset_root=tmp_path/'assets')
    assets.verified_asset_path(asset,asset_root=tmp_path/'assets')
    path.write_bytes(b'changed after cached decode')
    with pytest.raises(ValueError,match='hash'):assets.verified_asset_path(asset,asset_root=tmp_path/'assets')
