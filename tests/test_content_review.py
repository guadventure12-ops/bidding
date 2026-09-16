"""Offline content cleanup/approval contracts, isolated DB and synthetic text."""
import copy
import hashlib
import json
import re

import pytest
from fastapi.testclient import TestClient

from app import content_review as cr, db, workflow, provider
from app.main import app


@pytest.fixture
def local(tmp_path, monkeypatch):
    monkeypatch.setattr(db, 'DATA', tmp_path / 'data')
    monkeypatch.setenv('MX_TESTING', '1')
    monkeypatch.setenv('LANGFUSE_ENABLED', 'false')
    monkeypatch.delenv('DEEPSEEK_API_KEY', raising=False)
    monkeypatch.setattr(provider, 'chat_json', lambda *a, **k: pytest.fail('Paid model call prohibited'))
    db.init()
    p = db.uid()
    db.insert('projects', {'id': p, 'name': '合成采购项目', 'domain': 'archive', 'company_name': '合成企业', 'created_at': db.now(), 'updated_at': db.now()})
    workflow.refresh_project_basics(p)
    return p, tmp_path


def section(p, text, title='合成章节', ids=None, reqs=None):
    id = db.uid()
    row = {'id': id, 'project_id': p, 'ordinal': 1, 'title': title, 'content': text, 'evidence_ids': ids or [], 'requirement_ids': reqs or [], 'created_at': db.now(), 'updated_at': db.now()}
    db.insert('sections', row)
    return db.one('SELECT * FROM sections WHERE id=?', (id,))


def evidence(local, text, status='approved', scope='archive'):
    _, tmp = local
    path = tmp / (db.uid() + '.txt')
    path.write_text(text, encoding='utf-8')
    doc = workflow.ingest(path)
    db.update('documents', doc['id'], {'status': status, 'scope': scope})
    chunk = db.one('SELECT * FROM chunks WHERE document_id=?', (doc['id'],))
    return doc, chunk


def test_assessments_batch_sources_and_recheck_after_revocation(local,monkeypatch):
    p,_=local;doc,chunk=evidence(local,'产品提供档案检索。')
    other_doc,other=evidence(local,'产品记录查询日志。')
    first=section(p,'产品提供档案检索。[E:'+chunk['id']+']',ids=[chunk['id']])
    second=section(p,'产品提供档案检索。[E:'+chunk['id']+']',ids=[chunk['id']])
    original=workflow._valid_evidence;calls=[]
    def tracked(ids,*args,**kwargs):
        calls.append(list(ids));return original(ids,*args,**kwargs)
    monkeypatch.setattr(workflow,'_valid_evidence',tracked)
    cr.assessments(p,replacements={second['id']:{'content':'产品记录查询日志。[E:'+other['id']+']','evidence_ids':[other['id']]}})
    assert len(calls)==1 and set(calls[0])=={chunk['id'],other['id']}
    db.update('documents',doc['id'],{'status':'pending'})
    rows,_=cr.assessments(p)
    assert all(any(c['rule_id']=='section_evidence_validity' for c in row['structural_checks']) for row in rows)


def test_project_detail_computes_the_risk_assessments_once(local,monkeypatch):
    p,_=local;section(p,'这是有实质内容的合成章节。')
    original=cr.assessments;calls=[]
    def tracked(*args,**kwargs):calls.append(args);return original(*args,**kwargs)
    monkeypatch.setattr(cr,'assessments',tracked)
    response=TestClient(app).get('/api/projects/'+p)
    assert response.status_code==200 and len(response.json()['section_risks'])==1
    assert len(calls)==1


def requirement(p, category='technical'):
    id = db.uid()
    db.insert('requirements', {'id': id, 'project_id': p, 'title': '合成要求', 'text': '采购方要求核对声明。', 'quote': '采购方要求核对声明。', 'category': category, 'created_at': db.now(), 'updated_at': db.now()})
    return id


def test_exact_cleanup_preserves_structure_numbers_fields_and_unknown_wrappers():
    wrapper = cr.WRAPPERS[1][0]
    raw = '# 标题\n\n' + wrapper + '\n\n拟用文本：我方遵守采购程序。\n' + cr.GENERIC_GAP + '\n| 参数 | 值 |\n|---|---|\n| 数量 | 10 |\n\n【待补充：报价金额】\n![图](image.png)\n普通段落中的拟用文本：应保留。\n【其他标记】'
    result = cr.cleanup_text(raw)
    assert len(result['removed']) == 3
    assert '我方遵守采购程序。' in result['content']
    assert '# 标题' in result['content'] and '| 数量 | 10 |' in result['content']
    assert '【待补充：报价金额】' in result['content'] and '![图](image.png)' in result['content']
    assert '普通段落中的拟用文本：应保留。' in result['content'] and '【其他标记】' in result['content']
    assert re.findall(r'\d+', raw) == re.findall(r'\d+', result['content'])
    assert cr.cleanup_text(result['content'])['removed'] == []


@pytest.mark.parametrize('text', ['', '# 只有标题', cr.WRAPPERS[0][0], '# 标题\n' + cr.GENERIC_GAP, '| 标题 | 说明 |\n|---|---|', '【待补充：报价】'])
def test_empty_or_only_hint_cannot_approve(local, text):
    s = section(local[0], text)
    assert not cr.approval(s)['eligible']
    with TestClient(app) as client:
        assert client.patch('/api/sections/' + s['id'], json={'status': 'approved'}).status_code == 400


def test_ordinary_pending_word_not_a_gap_and_delivery_not_content_blocker(local):
    s = section(local[0], '缺失材料待补充流程由专人登记。\n【待补充：签字盖章】')
    a = cr.approval(s)
    # R010 requires internal labels to leave the body first. The real signature
    # task remains delivery-only and does not block approval of cleaned prose.
    assert not a['eligible']
    plan = cr.preview(local[0], [s['id']], 'separate_notes')
    assert plan['assessments'][0]['eligible']
    assert any(t['blocks_delivery'] and not t['blocks_content'] for t in plan['todos'])


@pytest.mark.parametrize('placeholder', ['【待补充】', '【待确认】', '【待补充 报价金额】'])
def test_real_bracket_placeholder_not_ignored_because_colon_missing(local, placeholder):
    s = section(local[0], '按采购文件编制投标响应。\n' + placeholder)
    assert not cr.approval(s)['eligible']


def test_internal_company_material_is_reusable_when_user_recognizes_it(local):
    _, e = evidence(local, '仅供内部学习。\n系统支持档案检索。')
    s = section(local[0], '系统支持档案检索。', ids=[e['id']])
    assert cr.approval(s)['eligible']


def test_trusted_exact_reuse_has_no_extra_metadata_requirement_and_partial_evidence_not_blanket(local):
    doc, e = evidence(local, '系统支持档案检索。')
    s = section(local[0], '系统支持档案检索。[E:' + e['id'] + ']', ids=[e['id']])
    assert cr.approval(s)['eligible']
    s['content'] += '\n本系统支持100万个并发用户。'
    a = cr.approval(s)
    assert not a['eligible'] and any('100万' in r for r in a['blockers'])


def test_table_claims_need_own_support_and_markdown_reuse_keeps_meaning(local):
    _, e = evidence(local, '系统支持档案检索。')
    s = section(local[0], '- **系统支持档案检索。** [E:' + e['id'] + ']', ids=[e['id']])
    assert cr.approval(s)['eligible']
    s['content'] += '\n| 能力 | 描述 |\n|---|---|\n| 容量 | 本系统支持100万个并发用户。 |'
    assert any('100万' in reason for reason in cr.approval(s)['blockers'])


def test_project_trust_keeps_original_doc_and_other_project_unchanged(local):
    p, _ = local
    doc, e = evidence(local, '系统支持档案检索。', status='pending')
    before = db.one('SELECT * FROM documents WHERE id=?', (doc['id'],))
    p2 = db.uid()
    db.insert('projects', {'id': p2, 'name': '另一项目', 'domain': 'archive', 'created_at': db.now(), 'updated_at': db.now()})
    plan = cr.trust_preview(p)
    cr.set_trust(p, [doc['id']], plan['token'], True)
    assert len(workflow._valid_evidence([e['id']], 'archive', p)) == 1
    assert not workflow._valid_evidence([e['id']], 'archive', p2)
    assert db.one('SELECT * FROM documents WHERE id=?', (doc['id'],)) == before
    db.update('documents', doc['id'], {'updated_at': db.now()})
    assert not workflow._valid_evidence([e['id']], 'archive', p), 'Source reparsing/review changes invalidate project trust'


@pytest.mark.parametrize('body,scope', [('失败案例：系统支持档案检索。', 'archive'), ('系统支持档案检索。', 'historical'), ('其他客户专属：系统支持档案检索。', 'archive'), ('采购要求：系统支持档案检索。', 'archive'), ('AI生成：系统支持档案检索。', 'archive')])
def test_failure_other_customer_and_historical_sources_not_enterprise_facts(local, body, scope):
    _, e = evidence(local, body, scope=scope)
    assert not workflow._valid_evidence([e['id']], 'archive', local[0])


def test_cleanup_preview_no_write_apply_idempotent_and_undo(local):
    p, _ = local
    s = section(p, cr.WRAPPERS[1][0] + '\n\n拟用文本：按采购文件递交响应。')
    before = db.one('SELECT * FROM sections WHERE id=?', (s['id'],))
    snapshot_count = db.one('SELECT count(*) n FROM project_snapshots')['n']
    plan = cr.preview(p, [s['id']])
    assert plan['summary']['affected'] == 1
    assert db.one('SELECT * FROM sections WHERE id=?', (s['id'],)) == before
    assert db.one('SELECT count(*) n FROM project_snapshots')['n'] == snapshot_count
    with pytest.raises(ValueError): cr.execute(p, [s['id']], 'cleanup', plan['token'], False)
    result = cr.execute(p, [s['id']], 'cleanup', plan['token'], True)
    assert result['changed'] == 1
    plan2 = cr.preview(p, [s['id']])
    result2 = cr.execute(p, [s['id']], 'cleanup', plan2['token'], True)
    assert result2['changed'] == 0 and result2['snapshot_id'] is None
    cr.undo(result['snapshot_id'], True)
    assert db.one('SELECT * FROM sections WHERE id=?', (s['id'],)) == before


def test_batch_single_parity_refresh_and_reapproval(local):
    p, _ = local
    good = section(p, '按采购文件的格式编制响应文件。')
    bad = section(p, '投标金额为【待补充：报价金额】。')
    with TestClient(app) as client:
        single = client.patch('/api/sections/' + bad['id'], json={'status': 'approved'})
        assert single.status_code == 400
        plan = client.post(f'/api/projects/{p}/sections/batch/preview', json={'section_ids': [good['id'], bad['id']], 'action': 'approve'}).json()
        assert plan['summary']['eligible'] == 1 and plan['summary']['blocked'] == 1
        result = client.post(f'/api/projects/{p}/sections/batch/apply', json={'section_ids': plan['section_ids'], 'action': 'approve', 'token': plan['token'], 'confirmed': True}).json()
        assert result['approved'] == 1 and result['skipped'] == 1 and result['failed'] == 0
        assert db.one('SELECT status FROM sections WHERE id=?', (good['id'],))['status'] == 'approved'
        assert client.patch('/api/sections/' + good['id'], json={'content': '编辑后的实质正文。'}).status_code == 200
        assert db.one('SELECT status FROM sections WHERE id=?', (good['id'],))['status'] == 'draft'


def test_stale_preview_cross_project_and_running_jobs_rejected(local):
    p, _ = local
    s = section(p, cr.WRAPPERS[0][0] + '\n实质正文内容。')
    plan = cr.preview(p, [s['id']])
    db.update('sections', s['id'], {'content': s['content'] + '新修改'})
    with pytest.raises(ValueError, match='重新预览'): cr.execute(p, [s['id']], 'cleanup', plan['token'], True)
    with pytest.raises(ValueError): cr.preview('missing', [s['id']])
    db.insert('jobs', {'id': db.uid(), 'project_id': p, 'mode': 'generate', 'status': 'running', 'created_at': db.now()})
    with pytest.raises(ValueError, match='任务'): cr.execute(p, [s['id']], 'cleanup', cr.preview(p, [s['id']])['token'], True)


def test_undo_does_not_overwrite_later_edit(local):
    p, _ = local
    s = section(p, cr.WRAPPERS[0][0] + '\n实质正文内容。')
    plan = cr.preview(p, [s['id']])
    result = cr.execute(p, [s['id']], 'cleanup', plan['token'], True)
    db.update('sections', s['id'], {'content': '用户后续编辑，必须保留。'})
    with pytest.raises(ValueError, match='已有变化'): cr.undo(result['snapshot_id'], True)


def test_qualification_todos_dedup_not_each_proof_and_known_source_reuse(local):
    p, _ = local
    req = requirement(p, 'qualification')
    a = section(p, '我方承诺不存在采购人的附属机构情形。', reqs=[req])
    b = section(p, '我方承诺未被责令停业或破产。', reqs=[req])
    rows, todos = cr.assessments(p)
    grouped = [t for t in todos if t['kind'] == 'qualification_declaration']
    assert len(grouped) == 1 and set(grouped[0]['section_ids']) == {a['id'], b['id']}
    assert not any(t['kind'] == 'delivery' for t in todos)
    assert len([t for t in todos if t['kind'] == 'qualification_delivery']) == 1
    assert all(not row['eligible'] for row in rows)


def test_known_qualification_notes_move_to_central_task_but_real_fields_remain(local):
    p, _ = local
    req = requirement(p, 'qualification')
    body = '我方承诺未被责令停业或破产。\n' + cr.DECLARATION_NOTES[1] + '\n【待补充：统一社会信用代码】'
    s = section(p, body, reqs=[req])
    assert cr.cleanup_text(body)['content'] == body, 'No transfer without qualification context'
    plan = cr.preview(p, [s['id']])
    after = plan['changes'][0]['after']
    assert cr.DECLARATION_NOTES[1] not in after
    assert '【待补充：统一社会信用代码】' in after
    assert any(t['kind'] == 'qualification_declaration' and t['status'] == 'open' for t in plan['todos'])
    assert any(t['kind'] == 'qualification_delivery' and t['status'] == 'open' for t in plan['todos'])
    result = cr.execute(p, [s['id']], 'cleanup', plan['token'], True)
    cr.undo(result['snapshot_id'], True)
    assert db.one('SELECT content FROM sections WHERE id=?', (s['id'],))['content'] == body


def test_declaration_content_confirmation_does_not_complete_signature(local):
    p, _ = local
    req = requirement(p, 'qualification')
    s = section(p, '我方承诺未被责令停业或破产。', reqs=[req])
    _, todos = cr.assessments(p)
    item = next(t for t in todos if t['kind'] == 'qualification_declaration')
    cr.resolve_todo(p, item['id'], cr.todo_signature(item), '已核对企业实际状态及采购声明条款，认可声明内容。', True)
    a = cr.approval(db.one('SELECT * FROM sections WHERE id=?', (s['id'],)))
    assert a['eligible']
    delivery = next(t for t in a['todos'] if t['kind'] == 'qualification_delivery')
    assert delivery['status'] == 'open' and not delivery['blocks_content'] and delivery['blocks_delivery']


def test_generation_never_appends_generic_fallback_or_rewrites_body(local):
    req = {'id': 'R', 'category': 'qualification', 'text': '资格声明要求', 'quote': '资格声明要求', 'title': '资格'}
    draft = {'content': '我方承诺核对资格。', 'responses': [{'requirement_id': 'R', 'response': '我方承诺核对资格。', 'evidence_ids': []}]}
    result = workflow._qualification_templates({'requirements': [req]}, copy.deepcopy(draft))
    assert result['content'] == draft['content']
    assert result['responses'][0]['response'] == draft['responses'][0]['response']
    assert result['responses'][0]['gap'] and result['responses'][0]['gap_reason']
    assert cr.MARKER not in json.dumps(result, ensure_ascii=False)


def test_generation_exact_product_reuse_clears_false_confirmation_gap(local):
    _, e = evidence(local, '系统支持档案检索。')
    req = {'id': 'R', 'category': 'technical'}
    result = {'content': '系统支持档案检索。', 'responses': [{'requirement_id': 'R', 'response': '系统支持档案检索。', 'evidence_ids': [e['id']], 'gap': True, 'gap_reason': '再确认企业资料'}]}
    result = workflow._qualification_templates({'requirements': [req]}, result, {e['id']: e})
    assert result['responses'][0]['gap'] is False


def test_generation_preserves_specific_qualification_gaps_in_internal_tasks(local):
    p, _ = local
    rid = requirement(p, 'qualification')
    s = section(p, '我方承诺符合采购文件的资格要求。', reqs=[rid])
    req = db.one('SELECT * FROM requirements WHERE id=?', (rid,))
    result = {'content': s['content'], 'responses': [{'requirement_id': rid, 'response': s['content'], 'evidence_ids': [], 'gap': True, 'gap_reason': '缺少统一社会信用代码，需核对营业执照'}]}
    result = workflow._qualification_templates({'requirements': [req]}, result)
    cr.record_generation(p, s['id'], result)
    _, tasks = cr.assessments(p)
    assert any(t['kind'] == 'content_gap' and '统一社会信用代码' in t['message'] for t in tasks)
    assert '待企业确认的承诺模板' not in result['content']
    task = next(t for t in tasks if t['kind'] == 'content_gap')
    cr.resolve_todo(p, task['id'], cr.todo_signature(task), '已核对营业执照中的统一社会信用代码，并记录核对依据。', True)
    assert next(t for t in cr.assessments(p)[1] if t['id'] == task['id'])['status'] == 'resolved'
    db.update('sections', s['id'], {'content': s['content'] + '\n【待补充：统一社会信用代码】'})
    assert not cr.approval(db.one('SELECT * FROM sections WHERE id=?', (s['id'],)))['eligible']


def test_model_contradiction_remains_content_blocker(local):
    p, _ = local
    s = section(p, '实际正文内容。')
    db.insert('reviews', {'id': db.uid(), 'project_id': p, 'fingerprint': 'old', 'status': 'complete', 'findings': [{'section_id': s['id'], 'verdict': 'contradiction', 'reason': '明确的实质性来源冲突'}], 'created_at': db.now()})
    assert '明确的实质性来源冲突' in cr.approval(s)['blockers']


def test_export_draft_does_not_reinject_removed_wrappers(local):
    p, _ = local
    s = section(p, cr.WRAPPERS[0][0] + '\n按采购文件格式编制响应文件。')
    plan = cr.preview(p, [s['id']])
    cr.execute(p, [s['id']], 'cleanup', plan['token'], True)
    result = workflow.export_project(p, 'docx', False)
    from docx import Document
    text = '\n'.join(x.text for x in Document(result['path']).paragraphs)
    assert cr.MARKER not in text and cr.GENERIC_GAP not in text
    assert '按采购文件格式编制响应文件。' in text


def test_export_does_not_invent_proof_for_each_qualification_subclause(local):
    from app.documents import compose_docx
    from docx import Document
    p, tmp = local
    req = {'id': 'r', 'category': 'qualification', 'text': '供应商作出资格声明。', 'response': '我方承诺未被责令停业或破产。'}
    for deviation in (False, True):
        output = tmp / ('deviation.docx' if deviation else 'response.docx')
        compose_docx({'name': '合成采购', '_deviation_tables': deviation}, [req], [{'title': '资格声明', 'content': req['response']}], str(output))
        document = Document(output)
        cells = '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        assert '待补充 证明材料' not in cells and '证明材料：【待补充】' not in cells
        assert req['response'] in cells


def test_review_prompt_keeps_confirmation_out_of_each_body_paragraph():
    prompt = workflow._review_prompt(dict(name='合成', project_number='1', company_name='合成', buyer='采购方', deadline=''), [])
    assert '应在每条开头明确标明' not in prompt
    assert '不能要求每条正文追加' in prompt
