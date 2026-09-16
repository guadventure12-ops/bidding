"""Full API + worker integration on synthetic procurement; no live model or network."""
import copy
from io import BytesIO
import json
import os
import re
import socket
import uuid
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from lxml import etree

from app import db, provider, proposal_runtime, workflow
from app.main import app


SLOTS = ["评审索引表", "商务条款偏离表", "技术条款偏离表", "需求理解与解决方案", "整体设计方案",
         "系统集成方案", "功能模块设计及系统功能演示", "售后运维服务方案"]
INTERNAL = "采购人独立安排评委抽取，其流程仅由采购人负责。"
KNOWLEDGE = """# 合成档案产品资料

产品提供档案检索、授权访问、分类管理、归档记录和查询结果导出功能。

产品记录档案查询日志，支持按组织、账套和关键字检索，使用角色权限控制用户查看范围。

产品提供归档接口、档案元数据校验和服务日志查询功能。
"""


def tender_bytes():
    doc = Document()
    doc.add_paragraph("项目名称：隔离测试电子档案采购项目")
    doc.add_paragraph("采购人：合成采购单位")
    doc.add_paragraph("项目编号：MOCK-2026")
    doc.add_heading("采购需求", level=1)
    doc.add_paragraph("系统应提供档案检索和归档接口设计方案。")
    doc.add_paragraph(INTERNAL)
    table = doc.add_table(rows=5, cols=3)
    for cell, text in zip(table.rows[0].cells, ["模块", "功能名称", "功能描述"]): cell.text = text
    data = [["基础配置", "登录安全", "支持账号登录权限控制及异常日志查询。"],
            ["", "", "支持组织信息维护及员工角色分配。"],
            ["智能检索", "全文查询", "支持档案关键词检索及查询结果导出。"],
            ["", "借阅查询", "支持按档案类型和授权范围检索记录。"]]
    for row, values in zip(table.rows[1:], data):
        for cell, text in zip(row.cells, values): cell.text = text
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(1, 1).merge(table.cell(2, 1))
    table.cell(3, 0).merge(table.cell(4, 0))
    scores = doc.add_table(rows=3, cols=4)
    for cell, text in zip(scores.rows[0].cells, ["序号", "评分因素", "评审标准", "分值"]): cell.text = text
    for row, values in zip(scores.rows[1:], [["1", "需求理解与解决方案", "应提供需求理解与可落地方案，说明档案管理流程。", "10"],
                                            ["2", "功能模块设计及系统功能演示", "应按各功能模块提供操作流程和演示验证步骤。", "10"]]):
        for cell, text in zip(row.cells, values): cell.text = text
    doc.add_heading("第六章 响应文件格式", level=1)
    doc.add_paragraph("商务、技术文件目录")
    for slot in SLOTS: doc.add_paragraph(slot + "（格式自拟）")
    doc.add_paragraph('附件2：响应函')
    doc.add_paragraph('我方按照本采购项目的约定提交响应文件。')
    doc.add_paragraph('供应商： （盖章）')
    doc.add_paragraph('日期： 年 月 日')
    doc.add_paragraph("附件10：评审索引表")
    t = doc.add_table(rows=2, cols=3)
    for cell, text in zip(t.rows[0].cells, ["评审内容", "响应章节", "页码"]): cell.text = text
    doc.add_paragraph("附件11：商务条款偏离表")
    t = doc.add_table(rows=2, cols=5)
    for cell, text in zip(t.rows[0].cells, ["序号", "采购条目号", "商务条款", "响应", "说明"]): cell.text = text
    doc.add_paragraph("附件12：技术偏离表")
    t = doc.add_table(rows=2, cols=6)
    for cell, text in zip(t.rows[0].cells, ["序号", "采购条目号", "技术条款", "响应", "偏离", "说明"]): cell.text = text
    doc.add_paragraph('附件16：初次报价一览表')
    t = doc.add_table(rows=2, cols=3)
    for cell,text in zip(t.rows[0].cells,['报价项目','含税报价（元）','备注']):cell.text=text
    t.cell(1,0).text='电子档案采购项目'
    doc.add_paragraph('供应商： （盖章）')
    doc.add_paragraph('日期： 年 月 日')
    stream = BytesIO(); doc.save(stream)
    return stream.getvalue()


class MockModel:
    def __init__(self):
        self.calls = []
        self.generation_label = "初次方案"

    def __call__(self, system, prompt, cancel=None):
        if "投标目录规划助手" in system:
            task = json.loads(prompt)
            source_titles = task.get("source_required_children", [])
            structure = task.get("source_structure", [])
            titles = list(dict.fromkeys(source_titles or [row["title"] for row in structure] or [task["module"]["title"] + "编制方案"]))
            if task.get("single_form_section"):
                titles = titles[:1]
            assignments = {title: [] for title in titles}
            for requirement in task["requirements"]:
                text = requirement.get("title", "") + "\n" + requirement.get("text", "")
                owner = next((title for title in titles if title in text), titles[0])
                assignments[owner].append(requirement["id"])
            suboutline = [] if task.get("fixed_procurement_format") else [
                {"title": "方案执行流程", "children": [{"title": "操作与核对步骤"}, {"title": "异常处理方法"}]}]
            self.calls.append({"stage": "outline", "module_id": task["module"]["id"], "titles": titles,
                               "fixed_procurement_format": task.get("fixed_procurement_format"),
                               "source_required_children": source_titles,
                               "requirement_ids": [r["id"] for r in task["requirements"]]})
            return {"sections": [{"title": title, "requirement_ids": assignments[title],
                                  "suboutline": copy.deepcopy(suboutline)} for title in titles],
                    "_usage": {"prompt_tokens": 10, "completion_tokens": 20}}
        if "文档数据：\n" in prompt:
            blocks = json.loads(prompt.rsplit("文档数据：\n", 1)[1])
            rows = []
            for block in blocks:
                text = block["text"]
                is_feature = "|" in text and "支持" in text
                is_score = "|" in text and "应" in text and ("解决方案" in text or "演示" in text)
                if text.startswith("系统应提供") or is_feature or is_score or text == INTERNAL:
                    rows.append({"chunk_id": block["chunk_id"], "quote": text, "text": text,
                                 "title": text[:80], "number": "", "category": "other" if text == INTERNAL else "technical",
                                 "mandatory": False, "score": ""})
            self.calls.append({"stage": "analysis", "blocks": len(blocks), "requirements": len(rows)})
            return {"requirements": rows, "notes": [], "_usage": {"prompt_tokens": 10, "completion_tokens": 20}}
        if "\n要求数据：\n" in prompt and "\n证据数据：\n" in prompt:
            requirements = json.loads(prompt.rsplit("\n要求数据：\n", 1)[1].split("\n证据数据：\n", 1)[0])
            evidence = json.loads(prompt.rsplit("\n证据数据：\n", 1)[1])
            selected = [evidence[0]["id"]] if evidence else []
            citation = "[E:" + selected[0] + "]" if selected else ""
            body = self.generation_label + "：按授权范围执行档案检索并核对查询结果。" + citation
            marker = "完整章节任务及采购上下文（不是企业事实）：\n"
            task = json.JSONDecoder().raw_decode(prompt.split(marker, 1)[1])[0] if marker in prompt else {}
            suboutline = task.get("suboutline") or task.get("required_outline") or []
            content = "## 操作流程\n\n" + body
            if suboutline:
                parts = []
                for third in suboutline:
                    parts.extend(["### " + third["title"], body])
                    for fourth in third.get("children", []):
                        parts.extend(["#### " + fourth["title"], body])
                content = "\n\n".join(parts)
            self.calls.append({"stage": "generation", "requirement_ids": [r["id"] for r in requirements],
                               "label": self.generation_label, "evidence_count": len(evidence),
                               "suboutline": copy.deepcopy(suboutline), "section_id": task.get("section_id")})
            return {"content": content + "\n\n| 步骤 | 预期结果 |\n| --- | --- |\n| 检索 | 返回授权结果 |",
                    "responses": [{"requirement_id": r["id"], "response": body, "evidence_ids": selected, "gap": False, "gap_reason": ""} for r in requirements],
                    "_usage": {"prompt_tokens": 30, "completion_tokens": 40}}
        raise AssertionError("Unexpected model stage in isolated pipeline test: " + prompt[:120])


def execute_job(client, response):
    assert response.status_code == 200, response.text
    id = response.json()["job"]["id"]
    workflow._run(id)
    job = client.get("/api/jobs/" + id).json()["job"]
    assert job["status"] == "succeeded", json.dumps(job, ensure_ascii=False)
    return job


def create_and_analyze(client, source=None):
    response = client.post("/api/projects", json={"name": "隔离测试项目", "company_name": "合成投标企业", "domain": "archive",
                                                 "project_number": "MOCK-2026", "buyer": "合成采购单位"})
    assert response.status_code == 200, response.text
    project_id = response.json()["project"]["id"]
    response = client.post("/api/projects/" + project_id + "/documents",
                           files={"file": ("合成采购文件.docx", source if source is not None else tender_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200, response.text
    document = response.json()["document"]
    assert document["parse_status"] == "ready"
    analysis = execute_job(client, client.post(f"/api/projects/{project_id}/run", json={"mode": "analyze"}))
    return project_id, analysis


def confirm_source_outline(client, project_id):
    response = client.get(f"/api/projects/{project_id}/outline-plan")
    assert response.status_code == 200, response.text
    plan = response.json()
    assert plan["source_count"] == len(SLOTS)
    assert not plan["confirmed"]
    assert all(g["enabled"] == (g["origin"] == "tender") for g in plan["groups"])
    chosen = [{"id": g["id"], "enabled": g["origin"] == "tender"} for g in plan["groups"]]
    response = client.put(f"/api/projects/{project_id}/outline-plan", json={"revision": plan["revision"], "groups": chosen, "confirmed": True})
    assert response.status_code == 200, response.text
    saved = response.json()
    planned = execute_job(client, client.post(f"/api/projects/{project_id}/outline-plan/generate",
        json={"revision": saved["revision"], "confirmed": True}))
    assert planned["mode"] == "outline"
    return planned


def create_and_generate(client):
    project_id, analysis = create_and_analyze(client)
    confirm_source_outline(client, project_id)
    generation = execute_job(client, client.post(f"/api/projects/{project_id}/run", json={"mode": "generate"}))
    detail = client.get("/api/projects/" + project_id).json()
    return project_id, detail, analysis, generation


@pytest.fixture
def pipeline(tmp_path, monkeypatch):
    # Caller can choose a short isolated QA path on Windows. Long pytest names
    # plus request hashes otherwise exceed older Win32 path handling limits.
    base = Path(os.environ["MX_PIPELINE_TEST_ROOT"]) / uuid.uuid4().hex[:8] if os.environ.get("MX_PIPELINE_TEST_ROOT") else tmp_path / "data"
    monkeypatch.setattr(db, "DATA", base)
    monkeypatch.setenv("MX_TESTING", "1")
    monkeypatch.setenv("LANGFUSE_ENABLED", "false")
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    db.init()
    original_connect = socket.socket.connect
    def local_ipc_only(sock, address):
        # Windows asyncio implements its self-pipe via a loopback socketpair.
        # Permit that IPC, never production 8765 or an external destination.
        assert isinstance(address, tuple) and address[0] in ("127.0.0.1", "::1") and address[1] != 8765, "No external or production network in proposal integration QA"
        return original_connect(sock, address)
    def deny(*args, **kwargs): raise AssertionError("No DNS in proposal integration QA")
    monkeypatch.setattr(socket.socket, "connect", local_ipc_only)
    monkeypatch.setattr(socket, "getaddrinfo", deny)
    monkeypatch.setattr(provider, "key_configured", lambda: True)
    monkeypatch.setattr(workflow.POOL, "submit", lambda *args, **kwargs: None)
    model = MockModel(); monkeypatch.setattr(provider, "chat_json", model)
    client = TestClient(app)
    response = client.post("/api/knowledge/upload", files={"file": ("合成产品资料.md", KNOWLEDGE.encode(), "text/markdown")})
    assert response.status_code == 200, response.text
    doc_id = response.json()["document"]["id"]
    approved = client.patch("/api/knowledge/" + doc_id, json={"status": "approved", "scope": "archive"})
    assert approved.status_code == 200, approved.text
    assert approved.json()["document"]["status"] == "approved"
    return client, model, doc_id


def xml_from_docx(data):
    with zipfile.ZipFile(BytesIO(data)) as archive:
        return etree.fromstring(archive.read("word/document.xml"))


def text_of(xml):
    return "\n".join(xml.xpath("//w:t/text()", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}))


def test_new_project_api_analysis_generation_and_docx_zip_pipeline(pipeline):
    client, model, doc_id = pipeline
    pid, detail, analysis, generated = create_and_generate(client)
    assert analysis["result"]["ai_completed"] is True  # Mock provider completion, not live DeepSeek.
    assert detail["project"]["analysis_status"] == "complete"
    plan = detail["project"]["metadata"]["proposal_blueprint"]
    assert detail["project"]["metadata"]["generation_profile"] == "technical_proposal"
    assert plan["coverage"]["all_ids_preserved"]
    assert len(plan["feature_rows"]) == 4
    assert len({f["module"] for f in plan["feature_rows"]}) == 2
    assert [g["title"] for g in plan["groups"] if g["volume"] == "technical"] == SLOTS
    assert set(detail["project"]["metadata"]["proposal_source_policy"]["fact_document_ids"]) == {doc_id}
    assert not detail["response_version_issues"]
    expected_calls = sum(s["content_kind"] == "narrative" for s in plan["sections"])
    assert len([c for c in model.calls if c["stage"] == "generation"]) == expected_calls
    assert generated["result"]["generated_sections"] == len(plan["sections"])

    exported = client.post(f"/api/projects/{pid}/export", json={"format": "docx", "final": False})
    assert exported.status_code == 200, exported.text
    content = client.get(exported.json()["url"]).content
    xml = xml_from_docx(content); text = text_of(xml)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    headings = xml.xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading1"]]/w:r/w:t/text()', namespaces=ns)
    assert headings[:len(SLOTS)] == SLOTS
    assert "内部响应与交付台账" not in text and INTERNAL not in text
    assert "逐条招标要求响应表" not in text
    assert "初次方案" in text
    fields = xml.xpath('//w:fldSimple/@w:instr', namespaces=ns)
    refs = [re.search(r'PAGEREF\s+(\S+)', x).group(1) for x in fields if 'PAGEREF' in x]
    bookmark_names = set(xml.xpath('//w:bookmarkStart/@w:name', namespaces=ns))
    # The module scoring item covers both independent module chapters.
    assert len(refs) == len(set(refs)) and len(refs) >= 3 and set(refs) <= bookmark_names
    module_ids=[s['section_id'] for s in plan['sections'] if s['title'] in ('基础配置','智能检索')]
    assert {'mxs_'+sid for sid in module_ids} <= set(refs)
    assert "[[PAGE:" not in text

    archive_response = client.post(f"/api/projects/{pid}/export", json={"format": "zip", "final": False})
    assert archive_response.status_code == 200, archive_response.text
    with zipfile.ZipFile(BytesIO(client.get(archive_response.json()["url"]).content)) as archive:
        assert {"01_资格文件.docx", "02_商务技术文件.docx", "03_报价文件.docx", "内部响应与交付台账.csv", "内部来源与版本记录.json"} <= set(archive.namelist())
        assert INTERNAL in archive.read("内部响应与交付台账.csv").decode("utf-8-sig")
        technical = text_of(xml_from_docx(archive.read("02_商务技术文件.docx")))
        assert INTERNAL not in technical
        report = json.loads(archive.read("内部来源与版本记录.json"))
        assert report["response_issues"] == []


def test_single_regeneration_updates_revision_projection_only_one_section(pipeline):
    client, model, _ = pipeline
    pid, detail, _, _ = create_and_generate(client)
    plan = detail["project"]["metadata"]["proposal_blueprint"]
    spec = next(s for s in plan["sections"] if s["title"] == "智能检索")
    sid = spec["section_id"]
    before_sections = {s["id"]: copy.deepcopy(s) for s in detail["sections"]}
    before_requirements = db.all("SELECT * FROM requirements WHERE project_id=? ORDER BY id", (pid,))
    first_export = client.post(f"/api/projects/{pid}/export", json={"format": "docx", "final": False})
    old_bytes = client.get(first_export.json()["url"]).content
    model.generation_label = "单节更新方案"
    preview = client.get(f"/api/sections/{sid}/regeneration")
    assert preview.status_code == 200, preview.text
    preview = preview.json()
    call_count = len(model.calls)
    response = client.post(f"/api/sections/{sid}/regenerate", json={"revision": preview["revision"], "outline_revision": preview["outline_revision"],
                           "request_id": uuid.uuid4().hex, "instruction": "仅更新检索流程", "confirmed": True})
    single = execute_job(client, response)
    assert single["result"]["sections"] == 1
    after = client.get(f"/api/projects/{pid}").json()
    assert not after["response_version_issues"]
    assert len(after["sections"]) == len(before_sections)
    for row in after["sections"]:
        if row["id"] != sid: assert row == before_sections[row["id"]]
        else:
            assert row["title"] == before_sections[sid]["title"]
            assert "单节更新方案" in row["content"]
    assert db.all("SELECT * FROM requirements WHERE project_id=? ORDER BY id", (pid,)) == before_requirements
    selected = set(spec["requirement_ids"])
    projected = {r["id"]: r for r in after["requirements"]}
    assert all("单节更新方案" in projected[id]["response"] for id in selected)
    assert all("单节更新方案" not in r["response"] for id, r in projected.items() if id not in selected)
    assert len(model.calls) == call_count + 1
    assert set(model.calls[-1]["requirement_ids"]) == {f"R{i:02d}" for i in range(1, len(selected) + 1)}
    assert {r["id"] for r in after["requirements"] if "单节更新方案" in r["response"]} == selected
    assert client.get(first_export.json()["url"]).content == old_bytes
    result = client.post(f"/api/projects/{pid}/export", json={"format": "zip", "final": False})
    assert result.status_code == 200, result.text
    with zipfile.ZipFile(BytesIO(client.get(result.json()["url"]).content)) as archive:
        csv = archive.read("内部响应与交付台账.csv").decode("utf-8-sig")
        assert "单节更新方案" in csv
        assert "单节更新方案" in text_of(xml_from_docx(archive.read("02_商务技术文件.docx")))
        assert json.loads(archive.read("内部来源与版本记录.json"))["response_issues"] == []


def test_manual_section_edit_invalidates_stale_response_export_not_other_sections(pipeline):
    client, model, _ = pipeline
    pid, detail, _, _ = create_and_generate(client)
    spec = next(s for s in detail["project"]["metadata"]["proposal_blueprint"]["sections"] if s["title"] == "智能检索")
    sid = spec["section_id"]
    original = next(s for s in detail["sections"] if s["id"] == sid)
    response = client.patch(f"/api/sections/{sid}", json={"title": original["title"], "content": "人工新增正文：查询结果须按新流程核对。", "status": "draft"})
    assert response.status_code == 200, response.text
    refreshed = client.get(f"/api/projects/{pid}").json()
    assert {i["requirement_id"] for i in refreshed["response_version_issues"]} == set(spec["requirement_ids"])
    for requirement in refreshed["requirements"]:
        if requirement["id"] in spec["requirement_ids"]: assert requirement["response"] == ""
    export = client.post(f"/api/projects/{pid}/export", json={"format": "zip", "final": False})
    assert export.status_code == 200, export.text
    with zipfile.ZipFile(BytesIO(client.get(export.json()["url"]).content)) as archive:
        record = json.loads(archive.read("内部来源与版本记录.json"))
        assert {i["section_id"] for i in record["response_issues"]} == {sid}
        text = text_of(xml_from_docx(archive.read("02_商务技术文件.docx")))
        assert "人工新增正文" in text
