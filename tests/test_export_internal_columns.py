"""Offline protection for removing pure internal citation columns from export copies."""
from copy import deepcopy
from io import BytesIO
import json
import socket
import sqlite3
import subprocess
import zipfile

import pytest
from docx import Document
from app import documents, proposal_export, review_rules, index_materials


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    def blocked(*args, **kwargs):
        raise AssertionError("No network, database or external process in export-copy tests")
    monkeypatch.setattr(socket.socket, "connect", blocked)
    monkeypatch.setattr(socket, "getaddrinfo", blocked)
    monkeypatch.setattr(sqlite3, "connect", blocked)
    monkeypatch.setattr(subprocess, "Popen", blocked)
    # Version invalidation is tested by test_index_pages_pipeline. This suite
    # isolates the exported copy and does not query workspace sources/settings.
    monkeypatch.setattr(index_materials, "input_revision", lambda project, sections: "offline-export-copy")
    with review_rules.scope({key: True for key in review_rules.IDS}):
        yield


def table(header="证据", first="[E:known]", second=""):
    return f"| 要素 | {header} | 说明 |\n| --- | :---: | --- |\n| 权限 | {first} | 账套隔离100 |\n| 日志 | {second} | 保留原说明 |"


@pytest.mark.parametrize("header", ["证据", "来源"])
def test_only_pure_reference_columns_are_removed_and_original_values_recorded(header):
    original = "## 原标题\n\n" + table(header, "[E:known] [E:second-id_1]") + "\n\n原数字123.45。"
    result = documents._drop_internal_evidence_columns(original)
    assert result["content"] == "## 原标题\n\n| 要素 | 说明 |\n| --- | --- |\n| 权限 | 账套隔离100 |\n| 日志 | 保留原说明 |\n\n原数字123.45。"
    assert result["removed"] == [{"kind": "internal_evidence_column", "line": 3, "column": 2,
                                  "header": header, "raw_cells": [" [E:known] [E:second-id_1] ", "  "]}]
    assert documents._drop_internal_evidence_columns(result["content"]) == {"content": result["content"], "removed": []}


@pytest.mark.parametrize("header,first,second", [
    ("证据", "", ""), ("来源", "", ""), ("附件", "", ""),
    ("证明材料", "[E:known]", ""), ("报价", "", ""),
    ("资料来源", "[E:known]", ""), ("**来源**", "[E:known]", ""),
    ("来源", "[E:known]", "《产品手册》权限章节"),
    ("证据", "[E:known] 实际证明名称", ""),
    ("来源", "[E:known]；", ""), ("证据", "[E:]", ""),
    ("来源", "https://example.invalid/source", ""),
])
def test_blank_forms_substantive_cells_and_nonexact_headers_are_preserved(header, first, second):
    original = table(header, first, second)
    assert documents._drop_internal_evidence_columns(original) == {"content": original, "removed": []}


@pytest.mark.parametrize("original", [
    "正文中的证据 [E:known] 不属于表格",
    "```markdown\n" + table() + "\n```\n",
    "~~~\n" + table() + "\n~~~\n",
    table() + "\n| 不完整行 |",
    "| 证据 | 来源 |\n| --- | --- |\n| [E:known] | [E:second] |",
])
def test_non_tables_fences_ragged_rows_and_all_reference_tables_are_not_erased(original):
    assert documents._drop_internal_evidence_columns(original) == {"content": original, "removed": []}


def inputs(body, *, hide=True, profile="technical_proposal"):
    sid = "a" * 32
    spec = {"section_id": sid, "section_key": "test", "group_title": "技术方案", "title": "权限设计",
            "volume": "technical", "content_kind": "narrative"}
    project = {"id": "isolated", "name": "合成采购项目", "buyer": "合成采购方", "company_name": "合成企业",
               "project_number": "QA-ONLY", "_hide_inline_evidence": hide,
               "_omit_response_table": True, "_omit_attachments": True, "_omit_evidence_index": True,
               "metadata": {"generation_profile": profile, "proposal_blueprint":
                            {"recognized": True, "sections": [spec], "ledger": [], "score_factors": []}}}
    sections = [{"id": sid, "title": "权限设计", "outline_group_id": "g", "outline_group_title": "技术方案",
                 "content": body, "status": "approved", "requirement_ids": [], "evidence_ids": ["known"]}]
    company = {"name": "合成企业", "evidence": [{"id": "known", "document_name": "原产品手册", "locator": "权限章节"}]}
    return project, [], sections, company


@pytest.mark.parametrize("hide,profile,expected_columns", [
    (True, "technical_proposal", 2),
    (False, "technical_proposal", 3),
    (True, "legacy", 3),
    (False, "legacy", 3),
])
def test_compose_only_changes_hidden_technical_export_copy(tmp_path, hide, profile, expected_columns):
    p, reqs, sections, company = inputs(table(), hide=hide, profile=profile)
    before = deepcopy((p, reqs, sections, company))
    output = tmp_path / "technical.docx"
    result = documents.compose_docx(p, reqs, sections, str(output), company)
    doc = Document(output)
    assert len(doc.tables) == 1 and len(doc.tables[0].columns) == expected_columns
    assert any("账套隔离100" in cell.text for row in doc.tables[0].rows for cell in row.cells)
    notes = [x for x in result["source_notes"] if x["kind"] == "internal_evidence_column"]
    assert bool(notes) is (hide and profile == "technical_proposal")
    if notes:
        assert notes[0]["section_id"] == sections[0]["id"]
        assert notes[0]["raw_cells"] == [" [E:known] ", "  "]
    assert (p, reqs, sections, company) == before


def test_removed_unknown_reference_still_warns_and_keeps_internal_source_record(tmp_path):
    p, reqs, sections, company = inputs(table("来源", "[E:missing]", "[E:known]"))
    result = documents.compose_docx(p, reqs, sections, str(tmp_path / "unknown.docx"), company)
    assert any("1 个引用未找到对应企业来源" in x for x in result["warnings"])
    notes = [x for x in result["source_notes"] if x["kind"] == "internal_evidence_column"]
    assert notes[0]["raw_cells"] == [" [E:missing] ", " [E:known] "]
    assert len(Document(tmp_path / "unknown.docx").tables[0].columns) == 2


def test_column_decision_uses_original_cells_before_other_export_cleanup(tmp_path):
    p, reqs, sections, company = inputs(table("来源", "[E:known]", "来源定位：正文 / 段落 3"))
    result = documents.compose_docx(p, reqs, sections, str(tmp_path / "locator.docx"), company)
    assert len(Document(tmp_path / "locator.docx").tables[0].columns) == 3
    assert not any(x['kind'] == 'internal_evidence_column' for x in result['source_notes'])


def test_two_column_table_still_renders_as_table_after_source_column_removed(tmp_path):
    body = "| 内容 | 证据 |\n| --- | --- |\n| 不得丢失的字段123 | [E:known] |"
    p, reqs, sections, company = inputs(body)
    documents.compose_docx(p, reqs, sections, str(tmp_path / "one-column.docx"), company)
    doc = Document(tmp_path / "one-column.docx")
    assert len(doc.tables) == 1 and len(doc.tables[0].columns) == 1
    assert doc.tables[0].cell(1, 0).text == "不得丢失的字段123"


def test_zip_retains_removed_cells_and_original_evidence_in_internal_record(tmp_path):
    p, reqs, sections, company = inputs(table("来源", "[E:missing]", "[E:known]"))
    before = deepcopy((p, reqs, sections, company))
    output = tmp_path / "isolated.zip"
    result = proposal_export.write_export(p, reqs, sections, output, company, "zip")
    assert any("1 个引用未找到对应企业来源" in x for x in result["warnings"])
    with zipfile.ZipFile(output) as archive:
        internal = json.loads(archive.read("内部来源与版本记录.json"))
        notes = internal["report"]["volumes"]["technical"]["source_notes"]
        removed = next(x for x in notes if x["kind"] == "internal_evidence_column")
        assert removed["section_id"] == sections[0]["id"]
        assert removed["raw_cells"] == [" [E:missing] ", " [E:known] "]
        assert internal["evidence"] == company["evidence"]
        doc = Document(BytesIO(archive.read("02_商务技术文件.docx")))
        assert len(doc.tables[0].columns) == 2
        assert "证据" not in [c.text for c in doc.tables[0].rows[0].cells]
    assert (p, reqs, sections, company) == before
