"""Regression for real Word exports continuing list numbers across chapters."""
from docx import Document
from docx.oxml.ns import qn

from app.documents import compose_docx


def test_export_preserves_explicit_list_numbers_across_chapters(tmp_path):
    path = tmp_path / "numbers.docx"
    compose_docx({"name": "测试项目"}, [], [
        {"title": "第一章", "content": "1. 第一项\n3. 第三项\n7) 第七项"},
        {"title": "第二章", "content": "1. 重新开始\n2. 第二项\n3. 最末事项"},
    ], str(path))
    paragraphs = Document(path).paragraphs
    expected = ["1. 第一项", "3. 第三项", "7) 第七项", "1. 重新开始", "2. 第二项", "3. 最末事项"]
    matching = [p for p in paragraphs if p.text in expected]
    assert [p.text for p in matching] == expected
    assert all(p.style.name == "Normal" and not p._p.xpath("./w:pPr/w:numPr") for p in matching)


def test_only_duplicate_leading_heading_is_removed_and_quote_text_is_preserved(tmp_path):
    path = tmp_path / "headings.docx"
    compose_docx({"name": "测试项目"}, [], [
        {"title": "第一章", "content": "# **第一章**\n\n> 待企业确认，不可直接递交。\n\n## 其他条款\n内容\n## 第一章\n正文内的同名标题保留"},
    ], str(path), company={"checks": [{"code": "model_review_missing"}]})
    paragraphs = Document(path).paragraphs
    text = [p.text for p in paragraphs]
    assert text.count("第一章") == 2  # outer chapter + intentional later heading
    assert "待企业确认，不可直接递交。" in text
    assert not any(t.startswith(">") for t in text)
    assert "正文内的同名标题保留" in text
    assert any("当前版本独立 AI 复核未完成" in t for t in text)


def test_toc_is_followed_by_heading_page_break_not_empty_break_paragraph(tmp_path):
    for has_sections in (True, False):
        path = tmp_path / f"toc-{has_sections}.docx"
        compose_docx({"name": "测试项目"}, [], [{"title": "正文第一章", "content": "实际正文"}] if has_sections else [], str(path))
        paragraphs = Document(path).paragraphs
        toc = next(i for i, p in enumerate(paragraphs) if ' TOC ' in p._p.xml)
        toc_paragraph = paragraphs[toc]
        assert toc_paragraph.paragraph_format.space_before.pt == 0
        assert toc_paragraph.paragraph_format.space_after.pt == 0
        # Only the field's final paragraph mark is small. The visible entries
        # and placeholder must keep their normal text sizes after field update.
        for property_name in ("sz", "szCs"):
            marks = toc_paragraph._p.xpath(f"./w:pPr/w:rPr/w:{property_name}")
            assert len(marks) == 1 and marks[0].get(qn("w:val")) == "2"
            assert not toc_paragraph._p.xpath(f"./w:r/w:rPr/w:{property_name}")
        assert toc_paragraph.text == "在 Word 中更新域以生成目录与页码"
        assert [field.text for field in toc_paragraph._p.xpath(".//w:instrText")] == [' TOC \\o "1-3" \\h \\z \\u ']
        following = paragraphs[toc + 1]
        assert following.style.name == "Heading 1"
        assert following.paragraph_format.page_break_before is True
        assert not following._p.xpath('.//w:br[@w:type="page"]')
