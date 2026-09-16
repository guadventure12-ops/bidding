"""Iterate-6 HTTP -> worker -> projection integration, synthetic records only.

Uses the pipeline fixture's isolated database, MockModel, and DNS/network guards.
No paid model, production application port, native Word, or customer files.
"""
import copy
import json
import uuid
from io import BytesIO

from docx import Document

from app import db, provider, workflow
from test_proposal_pipeline import (SLOTS, create_and_analyze, create_and_generate,
    execute_job, pipeline, xml_from_docx)


def get_plan(client, pid):
    result = client.get(f"/api/projects/{pid}/outline-plan")
    assert result.status_code == 200, result.text
    return result.json()


def save_selection(client, pid, plan, enabled=None, reverse=False):
    chosen = [{"id": row["id"], "enabled": row["enabled"] if enabled is None else row["id"] in enabled}
              for row in plan["groups"]]
    if reverse:
        chosen.reverse()
    response = client.put(f"/api/projects/{pid}/outline-plan",
        json={"revision": plan["revision"], "confirmed": True, "groups": chosen})
    assert response.status_code == 200, response.text
    return response.json()


def plan_selected(client, pid, plan):
    return execute_job(client, client.post(f"/api/projects/{pid}/outline-plan/generate",
        json={"revision": plan["revision"], "confirmed": True}))


def get_detail(client, pid):
    result = client.get(f"/api/projects/{pid}")
    assert result.status_code == 200, result.text
    return result.json()


def business_values(rows):
    fields = ("id", "project_id", "title", "content", "status", "requirement_ids",
              "evidence_ids", "user_edited", "legacy_title", "created_at", "updated_at")
    return {row["id"]: {field: copy.deepcopy(row.get(field)) for field in fields} for row in rows}


def current_rows(pid):
    return db.all("SELECT * FROM sections WHERE project_id=? ORDER BY ordinal,id", (pid,))


def test_new_project_cannot_generate_before_saved_selection_or_pending_plan(pipeline):
    client, model, _ = pipeline
    pid, _ = create_and_analyze(client)
    plan = get_plan(client, pid)
    calls = len(model.calls)
    assert not plan["confirmed"]
    assert [g["title"] for g in plan["groups"] if g["origin"] == "tender"] == SLOTS
    assert all(not g["enabled"] for g in plan["groups"] if g["origin"] == "generic")
    blocked = client.post(f"/api/projects/{pid}/run", json={"mode": "generate"})
    assert blocked.status_code in (400, 409), blocked.text
    assert "目录" in blocked.text and len(model.calls) == calls
    blocked_plan = client.post(f"/api/projects/{pid}/outline-plan/generate",
        json={"revision": plan["revision"], "confirmed": True})
    assert blocked_plan.status_code in (400, 409), blocked_plan.text
    assert len(model.calls) == calls and current_rows(pid) == []
    saved = save_selection(client, pid, plan)
    assert saved["confirmed"] and saved["generation"]["pending_group_ids"]
    blocked = client.post(f"/api/projects/{pid}/run", json={"mode": "generate"})
    assert blocked.status_code in (400, 409) and "下级目录" in blocked.text
    assert len(model.calls) == calls and current_rows(pid) == []


def test_planning_only_selected_modules_preserves_source_children_and_does_not_duplicate(pipeline):
    client, model, _ = pipeline
    pid, _ = create_and_analyze(client)
    plan = get_plan(client, pid)
    chosen = {g["id"] for g in plan["groups"] if g["title"] in ("需求理解与解决方案", "功能模块设计及系统功能演示") and g["origin"] == "tender"}
    assert len(chosen) == 2
    saved = save_selection(client, pid, plan, chosen)
    requirements = db.all("SELECT * FROM requirements WHERE project_id=? ORDER BY id", (pid,))
    result = plan_selected(client, pid, saved)
    calls = [call for call in model.calls if call["stage"] == "outline"]
    assert len(calls) == 2 and {call["module_id"] for call in calls} == chosen
    assert {"基础配置", "智能检索"} <= {title for call in calls for title in call["source_required_children"]}
    assert not [call for call in model.calls if call["stage"] == "generation"]
    detail = get_detail(client, pid)
    rows = detail["sections"]
    assert len(rows) == len({row["id"] for row in rows})
    assert all(row["content"] == "" and row["status"] == "draft" for row in rows)
    groups = [g for g in detail["chapter_outline"]["groups"] if g["id"] in chosen]
    assert len(groups) == 2
    assert {"基础配置", "智能检索"} <= {leaf["title"] for group in groups for leaf in group["sections"]}
    for group in groups:
        for leaf in group["sections"]:
            assert leaf["children"] and all(child["children"] for child in leaf["children"])
    before = copy.deepcopy(current_rows(pid))
    stale = client.post(f"/api/projects/{pid}/outline-plan/generate", json={"revision": saved["revision"], "confirmed": True})
    assert stale.status_code in (400, 409), stale.text
    latest = get_plan(client, pid)
    repeat = client.post(f"/api/projects/{pid}/outline-plan/generate", json={"revision": latest["revision"], "confirmed": True})
    assert repeat.status_code in (400, 409), repeat.text
    workflow._run(result["id"])
    assert current_rows(pid) == before
    assert len([call for call in model.calls if call["stage"] == "outline"]) == 2
    assert db.all("SELECT * FROM requirements WHERE project_id=? ORDER BY id", (pid,)) == requirements


def test_no_source_directory_defaults_all_generic_modules_off_and_requires_manual_choice(pipeline):
    client, model, _ = pipeline
    document = Document()
    document.add_paragraph("项目名称：合成无目录项目")
    document.add_paragraph("系统应提供档案检索和归档接口设计方案。")
    stream = BytesIO(); document.save(stream)
    pid, _ = create_and_analyze(client, stream.getvalue())
    plan = get_plan(client, pid)
    assert plan["source_count"] == 0 and plan["groups"]
    assert all(g["origin"] == "generic" and not g["enabled"] for g in plan["groups"])
    assert plan["generation"]["pending_group_ids"] == []
    assert not plan["confirmed"]
    target = next(g for g in plan["groups"] if not g.get("duplicate_of"))
    saved = save_selection(client, pid, plan, {target["id"]})
    assert saved["generation"]["pending_group_ids"] == [target["id"]]
    plan_selected(client, pid, saved)
    detail = get_detail(client, pid)
    enabled = [g for g in detail["compilation_outline"]["groups"] if g["enabled"]]
    assert [g["id"] for g in enabled] == [target["id"]]
    assert all(g["source_refs"] == [] for g in enabled)
    assert len([call for call in model.calls if call["stage"] == "outline"]) == 1
    execute_job(client, client.post(f"/api/projects/{pid}/run", json={"mode": "generate"}))
    generated = get_detail(client, pid)
    chosen_ids = set(enabled[0]["section_ids"])
    assert chosen_ids and all(row["content"] for row in generated["sections"] if row["id"] in chosen_ids)
    assert {call["section_id"] for call in model.calls if call["stage"] == "generation"} == chosen_ids
    assert all(call["suboutline"] for call in model.calls if call["stage"] == "generation")


def test_selection_projection_preserves_content_records_and_can_undo_without_database_restore(pipeline):
    client, model, _ = pipeline
    pid, detail, _, _ = create_and_generate(client)
    original = current_rows(pid)
    before_business = business_values(original)
    before_requirements = db.all("SELECT * FROM requirements WHERE project_id=? ORDER BY id", (pid,))
    before_documents = db.all("SELECT * FROM documents ORDER BY id")
    export = client.post(f"/api/projects/{pid}/export", json={"format": "docx", "final": False})
    assert export.status_code == 200, export.text
    old_export = client.get(export.json()["url"]).content
    # Export has its own current-check refresh. Freeze checks after that setup
    # action so this assertion isolates the outline-selection operation.
    before_checks = db.all("SELECT * FROM checks WHERE project_id=? ORDER BY id", (pid,))
    plan = get_plan(client, pid)
    hidden_group = next(g for g in plan["groups"] if g["title"] == "系统集成方案" and g["origin"] == "tender")
    hidden_ids = set(hidden_group["section_ids"])
    enabled = {g["id"] for g in plan["groups"] if g["enabled"] and g["id"] != hidden_group["id"]}
    calls = len(model.calls)
    saved = save_selection(client, pid, plan, enabled, reverse=True)
    after = get_detail(client, pid)
    assert hidden_ids == {row["id"] for row in after["retained_sections"]}
    assert hidden_ids.isdisjoint(after["active_section_ids"])
    assert {row["id"] for row in after["sections"]} == set(after["active_section_ids"])
    assert business_values(current_rows(pid)) == before_business
    assert db.all("SELECT * FROM requirements WHERE project_id=? ORDER BY id", (pid,)) == before_requirements
    assert db.all("SELECT * FROM checks WHERE project_id=? ORDER BY id", (pid,)) == before_checks
    assert db.all("SELECT * FROM documents ORDER BY id") == before_documents
    assert client.get(export.json()["url"]).content == old_export
    assert len(model.calls) == calls
    assert get_detail(client, pid)["active_section_ids"] == after["active_section_ids"]
    for sid in hidden_ids:
        blocked = client.get(f"/api/sections/{sid}/regeneration")
        assert blocked.status_code in (400, 409) and "未编入" in blocked.text
    assert len(model.calls) == calls
    fresh = client.post(f"/api/projects/{pid}/export", json={"format": "docx", "final": False})
    assert fresh.status_code == 200, fresh.text
    fresh_xml = xml_from_docx(client.get(fresh.json()["url"]).content)
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    headings = fresh_xml.xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading1"]]/w:r/w:t/text()', namespaces=namespaces)
    assert "系统集成方案" not in headings
    undone = client.post(f"/api/outline-operations/{saved['snapshot_id']}/undo", json={"confirmed": True})
    assert undone.status_code == 200, undone.text
    assert undone.json()["status"] == "restored"
    assert current_rows(pid) == original
    assert get_detail(client, pid)["retained_sections"] == []
    assert client.get(export.json()["url"]).content == old_export
    assert len(model.calls) == calls


def test_stale_revision_and_incomplete_filter_ids_are_rejected_without_writes(pipeline):
    client, model, _ = pipeline
    pid, _ = create_and_analyze(client)
    first = get_plan(client, pid)
    saved = save_selection(client, pid, first)
    metadata = copy.deepcopy(db.one("SELECT * FROM projects WHERE id=?", (pid,))["metadata"])
    calls = len(model.calls)
    stale = client.put(f"/api/projects/{pid}/outline-plan", json={"revision": first["revision"], "confirmed": True,
        "groups": [{"id": g["id"], "enabled": False} for g in first["groups"]]})
    assert stale.status_code in (400, 409), stale.text
    incomplete = client.put(f"/api/projects/{pid}/outline-plan", json={"revision": saved["revision"], "confirmed": True,
        "groups": [{"id": g["id"], "enabled": True} for g in saved["groups"] if g["origin"] == "tender"]})
    assert incomplete.status_code in (400, 409), incomplete.text
    assert db.one("SELECT * FROM projects WHERE id=?", (pid,))["metadata"] == metadata
    assert len(model.calls) == calls and current_rows(pid) == []


def test_duplicate_model_sections_fail_after_one_repair_without_partial_chapters(pipeline, monkeypatch):
    client, model, _ = pipeline
    pid, _ = create_and_analyze(client)
    plan = get_plan(client, pid)
    target = next(g for g in plan["groups"] if g["title"] == "需求理解与解决方案" and g["origin"] == "tender")
    saved = save_selection(client, pid, plan, {target["id"]})
    calls = []
    def malformed(system, prompt, cancel=None):
        assert "投标目录规划助手" in system
        task = json.loads(prompt); calls.append(task)
        row = {"title": "重复二级主题", "requirement_ids": [r["id"] for r in task["requirements"]],
               "suboutline": [{"title": "三级主题", "children": [{"title": "四级主题"}]}]}
        return {"sections": [row, copy.deepcopy(row)]}
    monkeypatch.setattr(provider, "chat_json", malformed)
    response = client.post(f"/api/projects/{pid}/outline-plan/generate", json={"revision": saved["revision"], "confirmed": True})
    assert response.status_code == 200, response.text
    jid = response.json()["job"]["id"]; workflow._run(jid)
    result = client.get(f"/api/jobs/{jid}").json()["job"]
    assert result["status"] == "failed" and "重复" in result["error"]
    assert len(calls) == 2 and "format_error" in calls[1]
    assert current_rows(pid) == []
    assert get_plan(client, pid)["generation"]["pending_group_ids"] == [target["id"]]


def test_undo_conflict_keeps_later_body_edit_and_current_selection(pipeline):
    client, model, _ = pipeline
    pid, detail, _, _ = create_and_generate(client)
    plan = get_plan(client, pid)
    hidden = next(g for g in plan["groups"] if g["title"] == "系统集成方案" and g["origin"] == "tender")
    saved = save_selection(client, pid, plan, {g["id"] for g in plan["groups"] if g["enabled"] and g["id"] != hidden["id"]})
    active = get_detail(client, pid)
    row = next(s for s in active["sections"] if s["outline_group_title"] == "需求理解与解决方案")
    changed = row["content"] + "\n\n后续人工补充：按合成项目要求核对记录。"
    result = client.patch(f"/api/sections/{row['id']}", json={"title": row["title"], "content": changed, "status": "draft"})
    assert result.status_code == 200, result.text
    before_undo = current_rows(pid)
    calls = len(model.calls)
    undo = client.post(f"/api/outline-operations/{saved['snapshot_id']}/undo", json={"confirmed": True})
    assert undo.status_code == 200, undo.text
    assert undo.json()["status"] == "conflict" and undo.json()["restored"] == 0 and undo.json()["skipped"] > 0
    assert current_rows(pid) == before_undo
    assert not next(g for g in get_plan(client, pid)["groups"] if g["id"] == hidden["id"])["enabled"]
    assert len(model.calls) == calls


def test_disabled_sections_reject_legacy_single_bulk_and_generation_endpoints_even_with_revision_rule_off(pipeline):
    client, model, _ = pipeline
    pid, detail, _, _ = create_and_generate(client)
    config = client.get("/api/review-settings").json()
    changed = client.patch("/api/review-settings", json={"revision": config["revision"], "enabled": {"operation_revision": False}})
    assert changed.status_code == 200, changed.text
    assert changed.json()["enabled"]["operation_revision"] is False
    hidden = next(s for s in detail["sections"] if s["outline_group_title"] == "系统集成方案")
    active = next(s for s in detail["sections"] if s["outline_group_title"] == "需求理解与解决方案")
    old_previews = {}
    for row in (hidden, active):
        response = client.get(f"/api/sections/{row['id']}/regeneration")
        assert response.status_code == 200, response.text
        old_previews[row["id"]] = response.json()
    old_bulk = client.post(f"/api/projects/{pid}/sections/batch/preview", json={"section_ids": [hidden["id"]], "action": "approve"})
    assert old_bulk.status_code == 200, old_bulk.text
    plan = get_plan(client, pid)
    save_selection(client, pid, plan, {g["id"] for g in plan["groups"] if g["enabled"] and hidden["id"] not in g["section_ids"]}, reverse=True)
    before = current_rows(pid)
    calls = len(model.calls)
    direct = client.patch(f"/api/sections/{hidden['id']}", json={"title": hidden["title"], "content": hidden["content"], "status": "approved"})
    assert direct.status_code in (400, 409) and "未编入" in direct.text
    preview = client.post(f"/api/projects/{pid}/sections/batch/preview", json={"section_ids": [hidden["id"]], "action": "approve"})
    assert preview.status_code in (400, 409) and "目录" in preview.text
    apply = client.post(f"/api/projects/{pid}/sections/batch/apply", json={"section_ids": [hidden["id"]], "action": "approve", "confirmed": True, "token": old_bulk.json()["token"]})
    assert apply.status_code in (400, 409) and "目录" in apply.text
    for sid, old in old_previews.items():
        generated = client.post(f"/api/sections/{sid}/regenerate", json={"revision": old["revision"],
            "outline_revision": old["outline_revision"], "request_id": uuid.uuid4().hex, "instruction": "", "confirmed": True})
        assert generated.status_code in (400, 409), generated.text
    assert current_rows(pid) == before and len(model.calls) == calls
