"""Presentation-only cleanup and grouping, including with review rules off."""
import copy
import pytest
from docx import Document
from app import review_rules
from app.documents import compose_docx, clean_export_locators
from test_review_rules import local


def test_locator_cleaning_is_exact_idempotent_and_preserves_tables_code_and_prose():
    original=('采购条款定位：正文 / 表格 6 / 第 2 行\n'
              '金额300万元、API 3.2、2026年保持。\n'
              '系统支持采购条款定位：按业务配置执行。\n'
              '| 列名 | 信息 |\n| --- | --- |\n| 300 | 采购条款定位：正文 / 段落 58\n'
              '| 500 | 正文保持 |\n'
              '```text\n采购条款定位：正文 / 段落 99\n```\n'
              '原文位置：未知处，仍需核对实际金额300万元。\n')
    cleaned=clean_export_locators(original)
    assert len(cleaned['removed'])==2
    assert '金额300万元、API 3.2、2026年保持。' in cleaned['content']
    assert '系统支持采购条款定位：按业务配置执行。' in cleaned['content']
    assert '| 300 | \n| 500 | 正文保持 |' in cleaned['content']
    assert '```text\n采购条款定位：正文 / 段落 99\n```' in cleaned['content']
    assert '实际金额300万元。' in cleaned['content']
    assert clean_export_locators(cleaned['content'])['removed']==[]


def test_only_known_requirement_annotations_are_removed():
    id='abcdef12'+'a'*24
    text=f'### 1. 实施说明（要求ID：{id}）\n响应（对应要求 abcdef12）保持。\n标准（2026）和（要求ID：deadbeef）保留。'
    result=clean_export_locators(text,[id])
    assert len(result['removed'])==2
    assert result['content']=='### 1. 实施说明\n响应保持。\n标准（2026）和（要求ID：deadbeef）保留。'
    assert clean_export_locators('### 章节（对应要求 deadbeef）')['content']=='### 章节'


@pytest.mark.parametrize('final',[False,True])
def test_docx_has_one_family_parent_unique_h2_and_no_body_locators(tmp_path,final):
    rows=[{'id':'a','title':'资格与商务响应（1）','content':'## 资格与商务响应（1）\n### 一、供应商资格\n采购条款定位：正文 / 表格 6 / 第 2 行\n原始资格正文300万元。'},
          {'id':'b','title':'资格与商务响应（2）','content':'## 资格与商务响应（2）\n### 一、信用记录\n采购条款定位：正文 / 段落 58\n原始信用正文2026年。'},
          {'id':'c','title':'资格与商务响应（3）','content':'## 资格与商务响应（3）——授权书\n原始授权正文。\n| 字段 | 值 |\n| --- | --- |\n| API | 3.2 |'}]
    before=copy.deepcopy(rows)
    path=tmp_path/'grouped.docx'
    p={'name':'项目','buyer':'采购人','project_number':'Q-1','company_name':'企业','_omit_response_table':True,'_omit_attachments':True}
    with review_rules.scope({id:False for id in review_rules.IDS}):
        result=compose_docx(p,[],rows,str(path),final=final)
    doc=Document(path);headings=[(x.style.name,x.text) for x in doc.paragraphs if x.style.name.startswith('Heading ')]
    assert [text for style,text in headings if style=='Heading 1']==['资格与商务响应']
    assert [text for style,text in headings if style=='Heading 2']==['供应商资格','信用记录','授权书']
    assert not any('（1）' in text or '（2）' in text or '（3）' in text for _,text in headings)
    text='\n'.join(x.text for x in doc.paragraphs)
    assert '采购条款定位' not in text and '原始资格正文300万元。' in text and '原始信用正文2026年。' in text
    assert doc.tables[0].cell(1,1).text=='3.2'
    assert len(result['source_notes'])==2 and len(result['outline'])==1
    assert [part['section_id'] for part in result['outline'][0]['parts']]==['a','b','c']
    assert rows==before


def test_standalone_source_appendix_remains_separate_from_body(tmp_path):
    path=tmp_path/'sources.docx'
    compose_docx({'name':'项目'},[],[{'title':'技术方案','content':'采购条款定位：正文 / 段落 2\n原始技术正文[E:e1]。'}],str(path),
                 {'evidence':[{'id':'e1','document_name':'产品手册.md','locator':'功能章节'}]})
    doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)
    assert '采购条款定位' not in text
    assert '原始技术正文' in text and '企业资料来源索引' in text
    assert '产品手册.md' in text and '功能章节' in text


@pytest.mark.parametrize('format',['docx','zip','md'])
def test_export_api_changes_only_new_output_and_keeps_business_rows(local,format):
    from app import db
    from io import BytesIO
    from zipfile import ZipFile
    db.set_setting('review_rules',{id:False for id in review_rules.IDS})
    db.update('sections','s',{'title':'资格与商务响应（1）','content':'# 资格与商务响应（1）\n## 主体资格\n采购条款定位：正文 / 段落 2\n原有正文300万元。'})
    db.insert('sections',{'id':'s2','project_id':'p','ordinal':2,'title':'资格与商务响应（2）','content':'# 资格与商务响应（2）\n## 信用核查\n第二段API 3.2。','created_at':db.now(),'updated_at':db.now()})
    before={t:db.all('SELECT * FROM '+t) for t in ('sections','requirements','documents','chunks','reviews','jobs','settings')}
    old=db.DATA/'exports'/'previous.docx';old.write_bytes(b'Previous output remains unchanged')
    response=local.post('/api/projects/p/export',json={'format':format,'final':True})
    assert response.status_code==200,response.text
    record=db.one('SELECT * FROM exports WHERE id=?',(response.json()['id'],))
    downloaded=local.get(response.json()['url'])
    assert downloaded.status_code==200
    if format=='md':
        text=downloaded.content.decode('utf-8')
        assert text.splitlines().count('## 资格与商务响应')==1
        assert '### 主体资格' in text and '### 信用核查' in text
    else:
        payload=downloaded.content
        if format=='zip':
            with ZipFile(BytesIO(payload)) as z:payload=z.read('01_资格文件.docx')
        doc=Document(BytesIO(payload));text='\n'.join(p.text for p in doc.paragraphs)
        assert sum(p.style.name=='Heading 1' and p.text=='资格与商务响应' for p in doc.paragraphs)==1
    assert '资格与商务响应（1）' not in text and '采购条款定位' not in text
    assert '原有正文300万元。' in text and '第二段API 3.2。' in text
    assert old.read_bytes()==b'Previous output remains unchanged'
    assert before=={t:db.all('SELECT * FROM '+t) for t in before}


def test_package_keeps_family_identity_when_parts_move_to_different_volumes(tmp_path):
    from app.documents import compose_bid_package
    from io import BytesIO
    from zipfile import ZipFile
    path=tmp_path/'volumes.zip'
    sections=[{'title':'技术方案（1）','content':'## 档案功能\n技术正文。'},
              {'title':'技术方案（2）','content':'## 费用响应\n含税报价：100元'}]
    with review_rules.scope({id:False for id in review_rules.IDS}):
        compose_bid_package({'name':'项目'},[],sections,str(path),final=True)
    with ZipFile(path) as z:
        for name in ('02_商务技术文件.docx','03_报价文件.docx'):
            doc=Document(BytesIO(z.read(name)))
            headings=[p.text for p in doc.paragraphs if p.style.name=='Heading 1']
            assert '技术方案' in headings
            assert not any('技术方案（' in text for text in headings)
