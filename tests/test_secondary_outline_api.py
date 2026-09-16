"""Real FastAPI/SQLite integration for manual H2 definitions; synthetic inputs only.

Uses the pipeline's isolated database, in-process TestClient, blocked DNS and
Mock provider. It deliberately does not use the historical H1-only planner.
DOCX export pagination remains the test-only boundary stub from conftest.
"""
import copy
from io import BytesIO
import uuid

from docx import Document

from app import db
from test_proposal_pipeline import create_and_analyze, execute_job, pipeline, xml_from_docx


def source_bytes(mode="explicit"):
    doc = Document()
    doc.add_paragraph("项目名称：二级目录隔离验收")
    doc.add_paragraph("系统应提供档案检索和归档接口设计方案。")
    doc.add_heading("技术评分细则", level=1)
    table = doc.add_table(rows=3, cols=4)
    for cell, value in zip(table.rows[0].cells, ["序号", "评分因素", "评审标准", "分值"]):
        cell.text = value
    for row, values in zip(table.rows[1:], [
        ["1", "技术方案", "应提供档案检索方案，说明授权控制和操作流程。", "10"],
        ["2", "实施安排", "应提供实施计划及培训安排。", "10"],
    ]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    if mode in ("explicit", "partial"):
        doc.add_paragraph("投标文件目录")
        doc.add_paragraph("一、技术方案")
        if mode == "explicit":
            doc.add_paragraph("（一）档案检索响应")
            doc.add_paragraph("（二）归档接口响应")
        doc.add_paragraph("二、实施安排")
        if mode == "explicit":
            doc.add_paragraph("（一）实施执行计划")
        doc.add_paragraph("供应商应按规定提供响应内容。")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def get_plan(client, pid):
    response = client.get(f"/api/projects/{pid}/outline-plan")
    assert response.status_code == 200, response.text
    return response.json()


def detail(client, pid):
    response = client.get(f"/api/projects/{pid}")
    assert response.status_code == 200, response.text
    return response.json()


def payload(plan):
    return {"revision": plan["revision"], "confirmed": True, "groups": [
        {"id": g["id"], "enabled": bool(g["enabled"]), "children": [
            {"id": c["id"], "title": c["title"], "enabled": c.get("enabled", True)}
            for c in g.get("children", [])]}
        for g in plan["groups"]
    ]}


def put_plan(client, pid, body):
    response = client.put(f"/api/projects/{pid}/outline-plan", json=body)
    assert response.status_code == 200, response.text
    return response.json()


def setup_defined(client, mode="explicit"):
    pid, _ = create_and_analyze(client, source_bytes(mode))
    first = get_plan(client, pid)
    assert first["source_mode"] == ("scoring" if mode == "scoring" else mode), first
    saved = put_plan(client, pid, payload(first))
    return pid, saved


def row(sid):
    return db.one("SELECT * FROM sections WHERE id=?", (sid,))


def rows(pid):
    return db.all("SELECT * FROM sections WHERE project_id=? ORDER BY ordinal,id", (pid,))


def children(plan):
    return [c for g in plan["groups"] if g["enabled"] for c in g.get("children", []) if c.get("enabled", True)]


def approve_body(client, sid, content="保留原文123\n\n| 参数 | 值 |\n| --- | --- |\n| 档案 | 20 |"):
    config = client.patch("/api/settings", json={"section_approval_threshold": "ignore"})
    assert config.status_code == 200, config.text
    result = client.patch(f"/api/sections/{sid}", json={"content": content, "status": "approved"})
    assert result.status_code == 200, result.text
    assert result.json()["section"]["status"] == "approved"
    return result.json()["section"]


def test_full_children_put_creates_empty_sections_with_source_names_and_no_planning_call(pipeline):
    client, model, _ = pipeline
    pid, _ = create_and_analyze(client, source_bytes())
    plan = get_plan(client, pid)
    assert plan["source_mode"] == "explicit"
    assert [c["title"] for c in children(plan)] == ["档案检索响应", "归档接口响应", "实施执行计划"]
    calls = copy.deepcopy(model.calls)
    saved = put_plan(client, pid, payload(plan))
    actual = detail(client, pid)
    defined = children(saved)
    assert all(c["section_id"] for c in defined)
    assert all(c["title_locked"] for c in defined)
    assert len({c["section_id"] for c in defined}) == 3
    assert saved["generation"]["pending_group_ids"] == []
    assert all(row(c["section_id"])["content"] == "" for c in defined)
    assert all(row(c["section_id"])["status"] == "draft" for c in defined)
    assert {c["section_id"] for c in defined} <= set(actual["active_section_ids"])
    assert model.calls == calls
    original_ids = {r["id"] for r in rows(pid)}
    again = put_plan(client, pid, payload(get_plan(client, pid)))
    assert {r["id"] for r in rows(pid)} == original_ids
    assert [c["section_id"] for c in children(again)] == [c["section_id"] for c in defined]
    assert model.calls == calls


def test_explicit_title_is_locked_through_direct_section_patch_while_body_stays_editable(pipeline):
    client, model, _ = pipeline
    pid, plan = setup_defined(client)
    target = children(plan)[0]
    sid = target["section_id"]
    before = row(sid)
    calls = copy.deepcopy(model.calls)
    response = client.patch(f"/api/sections/{sid}", json={"title": "不应接受的另一个标题"})
    assert response.status_code in (400, 409), response.text
    assert "招标文件明确规定" in response.text
    assert row(sid) == before
    response = client.patch(f"/api/sections/{sid}", json={"title": target["title"], "content": "手工正文123", "status": "draft"})
    assert response.status_code == 200, response.text
    assert response.json()["section"]["content"] == "手工正文123"
    assert next(c for c in children(get_plan(client, pid)) if c["section_id"] == sid)["title"] == target["title"]
    assert model.calls == calls


def test_scoring_title_patch_keeps_section_id_syncs_plan_and_resets_approval(pipeline):
    client, model, _ = pipeline
    pid, plan = setup_defined(client, "scoring")
    target = children(plan)[0]
    assert target["origin"] == "scoring" and not target["title_locked"]
    sid = target["section_id"]
    old = approve_body(client, sid)
    calls = copy.deepcopy(model.calls)
    response = client.patch(f"/api/sections/{sid}", json={"title": "人工核对后的评分响应"})
    assert response.status_code == 200, response.text
    new = response.json()["section"]
    assert new["id"] == old["id"] and new["content"] == old["content"]
    assert new["status"] == "draft" and new["user_edited"] == 1
    current = get_plan(client, pid)
    child = next(c for c in children(current) if c["section_id"] == sid)
    assert child["title"] == new["title"]
    assert child["source_refs"] == target["source_refs"]
    raw_selection = detail(client, pid)["project"]["metadata"]["outline_selection"]
    assert next(c for g in raw_selection["groups"] for c in g.get("children", []) if c.get("section_id") == sid)["title"] == new["title"]
    assert model.calls == calls


def test_custom_child_creation_and_renaming_do_not_change_existing_ids_body_or_records(pipeline):
    client, model, _ = pipeline
    pid, plan = setup_defined(client)
    kept = children(plan)[0]["section_id"]
    approve_body(client, kept)
    existing = {r["id"]: r for r in rows(pid)}
    old_docs = db.all("SELECT * FROM documents ORDER BY id")
    old_reqs = db.all("SELECT * FROM requirements WHERE project_id=? ORDER BY id", (pid,))
    calls = copy.deepcopy(model.calls)
    data = payload(get_plan(client, pid))
    group = next(g for g in data["groups"] if g["enabled"])
    custom_id = "new-" + str(uuid.uuid4())
    group["children"].append({"id": custom_id, "title": "人工补充验证方案", "enabled": True})
    saved = put_plan(client, pid, data)
    custom = next(c for c in children(saved) if c["id"] == custom_id)
    assert custom["origin"] == "custom" and custom["source_refs"] == []
    assert custom["requirement_ids"] == [] and custom["section_id"] not in existing
    assert row(custom["section_id"])["content"] == ""
    protected_fields = ("id", "content", "status", "evidence_ids", "requirement_ids", "user_edited", "created_at", "updated_at")
    for current in rows(pid):
        if current["id"] in existing:
            assert {k: current[k] for k in protected_fields} == {k: existing[current["id"]][k] for k in protected_fields}
    assert db.all("SELECT * FROM documents ORDER BY id") == old_docs
    assert db.all("SELECT * FROM requirements WHERE project_id=? ORDER BY id", (pid,)) == old_reqs
    approve_body(client, custom["section_id"], "自定义方案正文，不影响其他章节。")
    renamed = client.patch(f"/api/sections/{custom['section_id']}", json={"title": "人工补充验证步骤"})
    assert renamed.status_code == 200, renamed.text
    assert renamed.json()["section"]["status"] == "draft"
    assert next(c for c in children(get_plan(client, pid)) if c["id"] == custom_id)["title"] == "人工补充验证步骤"
    assert model.calls == calls


def test_disable_one_child_changes_only_active_projection_and_preserves_approved_body(pipeline):
    client, model, _ = pipeline
    pid, plan = setup_defined(client)
    target = children(plan)[0]
    sid = target["section_id"]
    before = approve_body(client, sid)
    calls = copy.deepcopy(model.calls)
    data = payload(get_plan(client, pid))
    next(c for g in data["groups"] for c in g["children"] if c["id"] == target["id"])["enabled"] = False
    put_plan(client, pid, data)
    actual = detail(client, pid)
    assert sid not in actual["active_section_ids"]
    assert sid in {r["id"] for r in actual["retained_sections"]}
    assert row(sid)["content"] == before["content"] and row(sid)["status"] == "approved"
    excluded = client.get(f"/api/sections/{sid}/regeneration")
    assert excluded.status_code in (400, 409) and "未编入" in excluded.text
    assert model.calls == calls
    data = payload(get_plan(client, pid))
    next(c for g in data["groups"] for c in g["children"] if c["id"] == target["id"])["enabled"] = True
    put_plan(client, pid, data)
    assert sid in detail(client, pid)["active_section_ids"]
    assert row(sid)["content"] == before["content"] and row(sid)["status"] == "approved"


def test_undo_directory_rename_skips_conflict_from_later_section_edit(pipeline):
    client, model, _ = pipeline
    pid, plan = setup_defined(client, "scoring")
    target = children(plan)[0]
    sid = target["section_id"]
    approve_body(client, sid)
    data = payload(get_plan(client, pid))
    next(c for g in data["groups"] for c in g["children"] if c["id"] == target["id"])["title"] = "目录调整后的评分响应"
    saved = put_plan(client, pid, data)
    assert row(sid)["status"] == "draft"
    response = client.patch(f"/api/sections/{sid}", json={"content": "目录保存后的新正文456"})
    assert response.status_code == 200, response.text
    before = row(sid)
    calls = copy.deepcopy(model.calls)
    undone = client.post(f"/api/outline-operations/{saved['snapshot_id']}/undo", json={"confirmed": True})
    assert undone.status_code == 200, undone.text
    assert undone.json()["status"] == "conflict", undone.text
    assert row(sid) == before
    assert model.calls == calls


def test_stale_or_missing_child_ids_cannot_replace_current_directory(pipeline):
    client, model, _ = pipeline
    pid, saved = setup_defined(client)
    before = rows(pid)
    data = payload(saved)
    next(g for g in data["groups"] if g["enabled"])["children"].pop()
    bad = client.put(f"/api/projects/{pid}/outline-plan", json=data)
    assert bad.status_code in (400, 409) and "完整提交" in bad.text
    assert rows(pid) == before
    stale = payload(saved)
    target = children(saved)[0]
    assert client.patch(f"/api/sections/{target['section_id']}", json={"content": "并发新正文"}).status_code == 200
    bad = client.put(f"/api/projects/{pid}/outline-plan", json=stale)
    assert bad.status_code in (400, 409) and "变化" in bad.text
    assert row(target["section_id"])["content"] == "并发新正文"


def test_single_child_ai_preserves_other_chapters_and_docx_uses_saved_hierarchy(pipeline):
    client, model, _ = pipeline
    pid, plan = setup_defined(client, "scoring")
    targets = children(plan)
    assert len(targets) >= 2
    target, sibling = targets[:2]
    sid = target["section_id"]
    approve_body(client, sibling["section_id"], "保留另一章正文789")
    before = {r["id"]: r for r in rows(pid) if r["id"] != sid}
    response = client.get(f"/api/sections/{sid}/regeneration")
    assert response.status_code == 200, response.text
    generation = response.json()
    calls = len(model.calls)
    execute_job(client, client.post(f"/api/sections/{sid}/regenerate", json={
        "revision": generation["revision"], "outline_revision": generation["outline_revision"],
        "request_id": str(uuid.uuid4()), "instruction": "只编写当前目录主题", "confirmed": True,
    }))
    added_calls = model.calls[calls:]
    assert added_calls and all(c["stage"] == "generation" and c["section_id"] == sid for c in added_calls)
    assert row(sid)["content"] and row(sid)["status"] == "draft"
    for current in rows(pid):
        if current["id"] in before:
            assert current == before[current["id"]]
    after = get_plan(client, pid)
    assert [(c["id"], c["section_id"], c["title"]) for c in children(after)] == [(c["id"], c["section_id"], c["title"]) for c in targets]
    exported = client.post(f"/api/projects/{pid}/export", json={"format": "docx", "final": False})
    assert exported.status_code == 200, exported.text
    xml = xml_from_docx(client.get(exported.json()["url"]).content)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    h1 = xml.xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading1"]]/w:r/w:t/text()', namespaces=ns)
    h2 = xml.xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading2"]]/w:r/w:t/text()', namespaces=ns)
    enabled_titles = [g["title"] for g in after["groups"] if g["enabled"]]
    assert h1[:len(enabled_titles)] == enabled_titles
    assert all(c["title"] in h2 for c in targets)
    assert len(model.calls) == calls + len(added_calls)
