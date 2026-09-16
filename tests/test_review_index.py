"""Scoring-source fidelity and bounded layout requests on isolated data."""
import copy
import json
from pathlib import Path
import zipfile

from docx import Document
import pytest

from app import db, provider, proposal_export as pe, proposal_runtime as pr, review_index as ri
from test_review_rules import local
from test_proposal_export import sample


def block(row,cells,spans=None,document='tender'):
    return {'id':f'{document}-r{row}','document_id':document,'ordinal':row,'locator':f'正文 / 表格 3 / 第 {row} 行','kind':'table_row',
            'text':' | '.join(cells),'metadata':{'table':{'index':3,'row':row,'cells':cells,'spans':spans or [1]*len(cells)}}}


def source():
    return [block(1,['序号','评审标准','分值'],[1,2,1]),block(2,['技术部分（5分）'],[4]),
            block(3,['1','需求理解','第一段完整标准\n第二段含 | 公式及 a < b。','5']),
            block(4,['价格部分（35分）'],[4]),block(5,['13','报价','按报价原公式计算得分。','35'])]


def model_case():
    p,reqs,sections=sample()
    factor=p['metadata']['proposal_blueprint']['score_factors'][0]
    factor.update(text=source()[2]['metadata']['table']['cells'][2],source={'chunk_id':'tender-r3','document_id':'tender','locator':source()[2]['locator']})
    p['metadata']['proposal_blueprint']['score_factors'].append({'number':'13','title':'报价','points':35,'source':{'chunk_id':'tender-r5','document_id':'tender','locator':source()[4]['locator']}})
    return p,reqs,sections,ri.build(p,sections,source())


def test_source_table_retains_every_row_header_span_text_and_unmapped_price():
    p,reqs,sections,model=model_case()
    table=model['tables'][0]
    assert model['mode']=='source_table' and table['width']==5
    assert table['rows'][0]['values']==['序号','评审标准','','分值','页数']
    assert table['rows'][0]['spans']==[1,2,1,1]
    assert table['rows'][1]['spans']==[4,1]
    assert len(table['rows'])==5
    assert table['rows'][-1]['values']==['13','报价','按报价原公式计算得分。','35','']
    assert table['rows'][2]['values'][2]==source()[2]['metadata']['table']['cells'][2]
    markdown=ri.markdown(model)
    assert '<br>' in markdown and '&#124;' in markdown and 'a &lt; b' in markdown
    assert '备注或者说明' not in markdown
    assert any(x['row_id']==table['rows'][-1]['id'] for x in model['unmapped_rows'])


def test_fragment_rows_deduplicate_and_document_identity_is_not_mixed():
    p,_,sections,_=model_case();blocks=source()
    fragment=copy.deepcopy(blocks[2]);fragment['id']='child';blocks.insert(3,fragment)
    blocks += [block(1,['编号','评分标准','分值'],document='another'),block(2,['A','第二文档原评分',''],document='another')]
    model=ri.build(p,sections,blocks)
    assert [len(t['rows']) for t in model['tables']]==[5,2]
    assert model['tables'][1]['rows'][1]['values']==['A','第二文档原评分','','']


def test_known_incomplete_scoring_table_does_not_fall_back():
    p,_,sections,_=model_case();blocks=source()
    blocks[2]['metadata']['table']={'index':3,'row':3,'split':True}
    with pytest.raises(ValueError,match='未完整解析'):ri.build(p,sections,blocks)


def test_attachment_index_form_is_not_the_scoring_table():
    p,_,sections,_=model_case()
    form=[block(1,['条款号','评审因素','对应页码','备注或者说明'],document='form'),block(2,['1','','',''],document='form')]
    model=ri.build(p,sections,form+source())
    assert len(model['tables'])==1 and model['tables'][0]['width']==5


def test_no_table_uses_scoring_prose_without_inventing_points():
    p,_,sections,_=model_case()
    prose=[{'id':'h','kind':'heading','text':'评分细则'},
           {'id':'a','kind':'paragraph','text':'1. 产品能力（10分）：按实际演示评分。'},
           {'id':'b','kind':'paragraph','text':'2. 服务方案：根据方案合理性评分，原文未列分值。'}]
    model=ri.build(p,sections,prose);rows=model['tables'][0]['rows']
    assert model['mode']=='generic'
    assert rows[0]['values']==['序号','评审项目','评分标准','分值','页数']
    assert rows[1]['values'][3]=='10' and rows[2]['values'][3]==''
    assert rows[2]['values'][2]==prose[2]['text']


def layout(model):
    t=model['tables'][0]
    return {'tables':[{'source_table_id':t['id'],'columns':[{'key':'c0','label':'编号'},{'key':'c1','label':'评审项目'},
             {'key':'c2','label':'评分原文'},{'key':'c3','label':'分值'},{'key':'chapter','label':'响应章节'},{'key':'page','label':'页数'}],
             'row_ids':[r['id'] for r in t['rows'][1:]]}]}


def test_instruction_is_sent_once_and_model_cannot_supply_scoring_values(local,monkeypatch):
    p,_,sections,model=model_case();calls=[];answer=layout(model)
    answer['content']='伪造评分999分'
    answer['_review_index_model']={'tables':[]}
    answer['_internal_notes']=[]
    def fake(system,prompt,cancel=None):calls.append(prompt);return answer
    monkeypatch.setattr(provider,'chat_json',fake)
    result,diagnostic=ri.apply_instruction(p,{'id':sections[0]['id'],'_instruction':'把序号改为编号，新增响应章节列'},model,'j','op')
    assert len(calls)==1 and '把序号改为编号，新增响应章节列' in calls[0]
    assert result['mode']=='custom_layout'
    assert result['tables'][0]['rows'][0]['values']==['编号','评审项目','评分原文','分值','响应章节','页数']
    assert result['tables'][0]['rows'][2]['values'][2:4]==source()[2]['metadata']['table']['cells'][2:4]
    assert '999' not in ri.markdown(result) and diagnostic['_request_diagnostic']
    assert '_review_index_model' not in diagnostic and '_internal_notes' not in diagnostic


@pytest.mark.parametrize('fault',['drop_row','duplicate','unknown_row','drop_column','invent_column'])
def test_invalid_layout_is_rejected_without_auto_retry(local,monkeypatch,fault):
    p,_,sections,model=model_case();answer=layout(model);t=answer['tables'][0];calls=[]
    if fault=='drop_row':t['row_ids'].pop()
    elif fault=='duplicate':t['row_ids'].append(t['row_ids'][0])
    elif fault=='unknown_row':t['row_ids'][0]='not-source'
    elif fault=='drop_column':t['columns'].pop(2)
    else:t['columns'].insert(0,{'key':'fiction','label':'自拟得分'})
    def fake(system,prompt,cancel=None):calls.append(prompt);return answer
    monkeypatch.setattr(provider,'chat_json',fake)
    with pytest.raises(provider.ProviderError,match='原章节保持不变'):
        ri.apply_instruction(p,{'id':sections[0]['id'],'_instruction':'调整格式'},model,'bad','op-'+fault)
    assert len(calls)==1


def test_saved_source_model_exports_exact_text_spans_and_real_page_fields(local,tmp_path):
    p,reqs,sections,model=model_case()
    sections[0].update(content=ri.markdown(model),user_edited=1)
    p['metadata']=pr.record_result(p['metadata'],sections[0],{'responses':[],'_review_index_model':model})
    path=tmp_path/'index.docx';report=pe.write_export(p,reqs,sections,path,{'name':'测试企业'})
    doc=Document(path);table=doc.tables[0]
    assert len(table.columns)==5
    assert len(table.rows[0]._tr.tc_lst)==4 and len(table.rows[1]._tr.tc_lst)==2
    assert table.cell(2,2).text==source()[2]['metadata']['table']['cells'][2]
    assert table.cell(4,3).text=='35' and table.cell(4,4).text==''
    assert report['volumes']['technical']['page_reference_fields']==1
    with zipfile.ZipFile(path) as z:assert b'[[PAGE:' not in z.read('word/document.xml')
    sections[0]['content']='| 人工列名 | 人工内容 |\n| --- | --- |\n| 项目 | 用户后续修改 |'
    projected,_=pe.projected_sections(p,reqs,sections,'technical')
    assert projected[0]['content']==sections[0]['content']
    assert projected[0]['_review_index_model']['mode']=='saved_markdown'


def test_multiple_page_targets_in_one_cell_all_become_fields(local,tmp_path):
    p,reqs,sections,model=model_case()
    model['tables'][0]['rows'][2]['values'][-1]='[[PAGE:'+sections[1]['id']+']]\n[[PAGE:'+sections[0]['id']+']]'
    sections[0].update(content=ri.markdown(model),user_edited=1)
    p['metadata']=pr.record_result(p['metadata'],sections[0],{'responses':[]})  # Explicit manual page mapping.
    path=tmp_path/'multiple.docx';report=pe.write_export(p,reqs,sections,path,{'name':'测试企业'})
    assert report['volumes']['technical']['page_reference_fields']==2
    with zipfile.ZipFile(path) as z:assert b'[[PAGE:' not in z.read('word/document.xml')


def test_generic_total_points_do_not_take_a_subitem_and_scope_resets():
    p,_,sections,_=model_case()
    blocks=[{'document_id':'a','kind':'heading','text':'四、评审细则'},
            {'id':'a1','document_id':'a','kind':'paragraph','text':'1. 方案设计（10分）：证书每份最多得2分，方案内容8分。'},
            {'id':'a2','document_id':'a','kind':'paragraph','text':'2. 服务：证书每份最多得2分，合计未列明。'},
            {'id':'b1','document_id':'b','kind':'paragraph','text':'1. 报价单：报价有效期30天。'}]
    rows=ri.build(p,sections,blocks)['tables'][0]['rows']
    assert len(rows)==3
    assert rows[1]['values'][3]=='10' and rows[2]['values'][3]==''


def test_scoring_table_without_numeric_points_stays_an_original_table():
    p,_,sections,_=model_case()
    blocks=[block(1,['序号','评分内容']),block(2,['1','方案覆盖系统架构和实施计划，按完整性评价'])]
    model=ri.build(p,sections,blocks)
    assert model['mode']=='source_table'
    assert model['tables'][0]['rows'][0]['values']==['序号','评分内容','页数']


def test_duplicate_factor_numbers_match_source_not_other_document():
    p,_,sections,model=model_case();plan=p['metadata']['proposal_blueprint']
    first=plan['score_factors'][0]
    second={**first,'title':'另一评分项','source':{'chunk_id':'elsewhere','document_id':'other','locator':'表3行3'}}
    plan['score_factors'].append(second)
    plan['sections'][1]['score_factors']=[second]
    targets,_=ri._owners(first,plan,sections)
    assert targets==[]
    plan['sections'][1]['score_factors']=[first]
    targets,_=ri._owners(first,plan,sections)
    assert targets==[sections[1]['id']]


def test_manual_change_does_not_export_entities_or_restore_old_values(local,tmp_path):
    p,reqs,sections,model=model_case()
    sections[0].update(content=ri.markdown(model),user_edited=1)
    p['metadata']=pr.record_result(p['metadata'],sections[0],{'responses':[],'_review_index_model':model})
    sections[0]['content']=sections[0]['content'].replace('第一段完整标准','人工修改的第一段')
    output=tmp_path/'edited.docx';pe.write_export(p,reqs,sections,output,{'name':'测试企业'})
    table=Document(output).tables[0]
    assert table.cell(2,2).text=='人工修改的第一段\n第二段含 | 公式及 a < b。'
    assert len(table.rows[0]._tr.tc_lst)==4
    assert '第一段完整标准' not in table.cell(2,2).text
