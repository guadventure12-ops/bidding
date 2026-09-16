"""Shared persisted directory and non-destructive, reversible migration."""
import copy
from docx import Document
import pytest
from test_review_rules import local
from app import db,chapter_outline as co,export_outline
from app.documents import compose_docx


def setup_rows():
    db.update('sections','s',{'title':'资格与商务响应（1）','content':'## 资格与商务响应（1）\n### 主体资格\n原文300万元。','status':'approved','user_edited':1})
    db.insert('sections',{'id':'s2','project_id':'p','ordinal':1,'title':'资格与商务响应（2）','content':'## 资格与商务响应（2）\n### 信用核查\n原文2026年。','status':'approved','created_at':db.now(),'updated_at':db.now()})
    db.insert('projects',{'id':'other','name':'不在本次范围','created_at':db.now(),'updated_at':db.now()})
    db.insert('sections',{'id':'else','project_id':'other','ordinal':0,'title':'其他章节（1）','content':'不得改变','status':'approved','created_at':db.now(),'updated_at':db.now()})
    db.insert('reviews',{'id':'historic','project_id':'p','fingerprint':'past','status':'complete','created_at':db.now()})


def test_migration_preserves_ids_body_approval_timestamps_and_all_old_history(local):
    setup_rows()
    before={t:db.all('SELECT * FROM '+t+' ORDER BY 1') for t in ('projects','sections','requirements','reviews','exports','documents','settings')}
    result=co.migrate_project('p')
    assert result['changed']==2
    assert result['outline']['group_count']==1 and result['outline']['section_count']==2
    after=db.all('SELECT * FROM sections ORDER BY id')
    for old,new in zip(before['sections'],after):
        assert {k:v for k,v in old.items() if k not in co.FIELDS}=={k:v for k,v in new.items() if k not in co.FIELDS}
        if old['project_id']=='p':assert new['legacy_title']==old['title']
        else:assert old==new
    assert after[-2]['title']=='主体资格'
    assert all(before[t]==db.all('SELECT * FROM '+t+' ORDER BY 1') for t in before if t!='sections')
    assert co.migrate_project('p')['changed']==0
    assert len(db.all('SELECT * FROM project_snapshots'))==1
    assert len(list((db.DATA/'outline-backups').glob('*.sqlite3')))==1


def test_shared_api_and_export_titles_survive_body_and_title_edits(local,tmp_path):
    setup_rows();co.migrate_project('p')
    original=db.one('SELECT * FROM sections WHERE id="s2"')
    assert local.patch('/api/sections/s',json={'title':'新的主体主题','content':'## 模型风格改变的标题\n新正文3.2。'}).status_code==200
    data=local.get('/api/projects/p').json()
    group=data['chapter_outline']['groups'][0]
    assert group['title']=='资格与商务响应'
    assert [r['title'] for r in group['sections']]==['新的主体主题','信用核查']
    assert db.one('SELECT * FROM sections WHERE id="s2"')==original
    path=tmp_path/'shared.docx'
    compose_docx({'name':'项目','_omit_response_table':True,'_omit_attachments':True},[],data['sections'],str(path))
    doc=Document(path)
    assert [p.text for p in doc.paragraphs if p.style.name=='Heading 1']==[group['title']]
    assert [p.text for p in doc.paragraphs if p.style.name=='Heading 2']==[r['title'] for r in group['sections']]
    assert local.patch('/api/sections/'+group['id'],json={'status':'approved'}).status_code==404
    assert local.patch('/api/sections/s',json={'title':'资格与商务响应（3）'}).status_code==400
    assert local.patch('/api/sections/s',json={'title':'信用核查'}).status_code==400


def test_migration_undo_skips_later_edit_and_only_restores_directory(local):
    setup_rows();result=co.migrate_project('p')
    db.update('sections','s2',{'content':'后来修改的正文'})
    outcome={r['id']:r['status'] for r in co.restore_migration(result['snapshot_id'])}
    assert outcome=={'s':'restored','s2':'conflict'}
    assert db.one('SELECT * FROM sections WHERE id="s"')['title']=='资格与商务响应（1）'
    assert db.one('SELECT * FROM sections WHERE id="s2"')['content']=='后来修改的正文'
    assert db.one('SELECT * FROM sections WHERE id="s2"')['title']=='信用核查'


def test_old_history_snapshot_matches_directory_only_migration(local):
    setup_rows();old=db.one('SELECT * FROM sections WHERE id="s"')
    old={k:v for k,v in old.items() if k not in ('outline_group_id','outline_group_title','legacy_title')}
    co.migrate_project('p');current=db.one('SELECT * FROM sections WHERE id="s"')
    assert co.snapshot_matches(current,old)
    assert co.preserve_directory_on_restore(current,old)['title']=='主体资格'
    assert not co.snapshot_matches({**current,'content':'后来改动'},old)


def test_model_planning_creates_fixed_leaf_before_batches_and_reuses_it(local):
    from app.workflow import CATEGORY_LABELS
    db.execute('DELETE FROM sections WHERE project_id=?',('p',))
    reqs=[]
    for i in range(25):
        r={'id':f'r{i}','category':'technical','title':f'功能{i}','text':'要求正文'}
        reqs.append(r)
    plan=co.generation_plan('p',reqs,CATEGORY_LABELS)
    assert len(plan)==1 and len(plan[0]['requirements'])==25
    row=db.one('SELECT * FROM sections WHERE id=?',(plan[0]['id'],))
    assert row['title']==plan[0]['title'] and row['outline_group_id']==plan[0]['outline_group_id']
    assert row['content']=='' and row['status']=='draft'
    assert co.generation_plan('p',reqs,CATEGORY_LABELS)[0]['id']==row['id']
    assert len(db.all('SELECT * FROM sections WHERE project_id="p"'))==1
    added=co.generation_plan('p',reqs+[{'id':'new','category':'technical','title':'新增范围','text':'新要求'}],CATEGORY_LABELS)
    assert len(added)==2 and len({r['id'] for r in added})==2
    assert co.generation_plan('p',reqs+[{'id':'new','category':'technical','title':'新增范围','text':'新要求'}],CATEGORY_LABELS)==added


def test_legacy_directory_cannot_be_renamed_by_model_planning(local):
    setup_rows()
    with pytest.raises(ValueError,match='旧目录尚未同步'):
        co.generation_plan('p',[],{})
