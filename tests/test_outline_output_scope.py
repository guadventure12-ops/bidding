"""Selected directory participation reaches export/index/assets, without model calls."""
import copy
import json
import zipfile

from docx import Document
from lxml import etree as E
import pytest

from app import db, compilation_outline, index_materials, proposal_assets
from app import proposal_export as pe, proposal_runtime as pr, review_index as ri
from test_review_rules import local
from test_proposal_export import sample


def select(project, sections, enabled):
    project['metadata']['outline_selection']={
        'confirmed':True, 'library_snapshot':{'revision':'frozen-v1'},
        'groups':[{'id':'selection-'+s['id'], 'enabled':s['id'] in enabled,
                   'section_ids':[s['id']]} for s in sections], 'fixed_groups':[]}


def xml_text(document):
    return '\n'.join([p.text for p in document.paragraphs]+
                     [c.text for t in document.tables for r in t.rows for c in r.cells])


def test_legacy_export_remains_complete_without_a_selection(local):
    p,reqs,sections=sample()
    assert [s['id'] for s in pe.projected_sections(p,reqs,sections,'technical')[0]]==[s['id'] for s in sections[:2]]
    assert ri.build(p,sections,blocks=[])['tables'][0]['rows'][1]['values'][-1]=='[[PAGE:'+sections[1]['id']+']]'


def test_selected_docx_omits_retained_section_and_old_page_binding_without_mutating_source(local,tmp_path):
    p,reqs,sections=sample();select(p,sections,{sections[0]['id']})
    before=copy.deepcopy((p,reqs,sections))
    path=tmp_path/'selected.docx'
    report=pe.write_export(p,reqs,sections,path,{'name':'测试企业'})
    assert report['volumes']['technical']['section_ids']==[sections[0]['id']]
    assert report['volumes']['technical']['page_reference_fields']==0
    doc=Document(path)
    assert [x.text for x in doc.paragraphs if x.style.name=='Heading 2']==['索引']
    assert sections[1]['content'] not in xml_text(doc)
    assert '需求理解' in xml_text(doc)  # The procurement scoring row remains.
    with zipfile.ZipFile(path) as z:
        root=E.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    assert [x.get('{'+ns['w']+'}name') for x in root.findall('.//w:bookmarkStart',ns)
            if x.get('{'+ns['w']+'}name','').startswith('mxs_')]==['mxs_'+sections[0]['id']]
    assert (p,reqs,sections)==before


def test_response_validation_still_receives_full_durable_sections(local,monkeypatch):
    p,reqs,sections=sample();select(p,sections,{sections[0]['id']})
    original=pr.project_responses;seen=[]
    def checked(project,rows,requirements):
        seen.append([s['id'] for s in rows])
        return original(project,rows,requirements)
    monkeypatch.setattr(pr,'project_responses',checked)
    pe.projected_sections(p,reqs,sections,'technical')
    assert seen==[[s['id'] for s in sections]]
    assert len(p['metadata']['proposal_blueprint']['sections'])==5


def test_package_keeps_full_internal_ledger_and_blueprint(local,tmp_path):
    p,reqs,sections=sample();select(p,sections,{s['id'] for s in sections if s!=sections[1]})
    path=tmp_path/'selected.zip';pe.write_export(p,reqs,sections,path,{'name':'测试企业'},format='zip')
    with zipfile.ZipFile(path) as z:
        ledger=z.read('内部响应与交付台账.csv').decode('utf-8-sig')
        metadata=json.loads(z.read('内部来源与版本记录.json'))
    assert all(r['id'] in ledger for r in reqs)
    assert metadata['blueprint']==p['metadata']['proposal_blueprint']
    assert '本次响应' in ledger


def test_selection_order_controls_rendering_without_reordering_records(local):
    p,reqs,sections=sample();select(p,sections,{s['id'] for s in sections})
    p['metadata']['outline_selection']['groups'].reverse()
    assert [s['id'] for s in pe.projected_sections(p,reqs,sections,'technical')[0]]==[sections[1]['id'],sections[0]['id']]
    assert sections[0]['title']=='索引'


def test_disabled_scoring_owner_remains_unmapped_not_deleted(local):
    p,reqs,sections=sample();select(p,sections,{sections[0]['id']})
    model=ri.build(p,sections,blocks=[])
    row=model['tables'][0]['rows'][1]
    assert row['values']==['1','需求理解','','5','']
    assert row['chapter']==''


def test_unmatched_custom_index_row_cannot_keep_disabled_page_link(local):
    p,reqs,sections=sample();select(p,sections,{sections[0]['id']})
    model={'tables':[{'rows':[{'kind':'data','values':['保留自定义描述','[[PAGE:'+sections[1]['id']+']]']}]}]}
    ri.refresh_model_pages(p,model,sections,set())
    assert model['tables'][0]['rows'][0]['values']==['保留自定义描述','']


@pytest.mark.parametrize('with_keys',[True,False])
def test_custom_index_derived_chapter_column_does_not_name_disabled_owner(local,with_keys):
    p,reqs,sections=sample();select(p,sections,{sections[0]['id']})
    model={'mode':'custom_layout','tables':[{'header_index':0,'rows':[
        {'kind':'header','values':['评审项目','所属章节' if with_keys else '响应章节','页数']},
        {'kind':'data','factor_number':'1','factor_title':'需求理解','values':['需求理解',sections[1]['outline_group_title'],'[[PAGE:'+sections[1]['id']+']]']}]}]}
    if with_keys:model['tables'][0]['column_keys']=['c0','chapter','page']
    ri.refresh_model_pages(p,model,sections,set())
    assert model['tables'][0]['rows'][1]['values']==['需求理解','','']


def test_unknown_page_target_still_fails_instead_of_being_silently_erased(local,tmp_path):
    p,reqs,sections=sample();select(p,sections,{sections[0]['id']})
    sections[0]['content']='| 页数 |\n| --- |\n| [[PAGE:'+('f'*32)+']] |'
    sections[0]['user_edited']=True
    with pytest.raises(ValueError,match='章节不存在'):
        pe.write_export(p,reqs,sections,tmp_path/'unknown.docx',{'name':'测试企业'})


def test_material_targets_exclude_disabled_sections(local,monkeypatch):
    p,reqs,sections=sample();select(p,sections,{sections[0]['id']})
    monkeypatch.setattr(db,'all',lambda *a,**k:sections)
    assert index_materials.scoped_sections(p)==[]
    p['metadata']['outline_selection']['groups'][1]['enabled']=True
    assert [s['id'] for s,_ in index_materials.scoped_sections(p)]==[sections[1]['id']]


def test_page_revision_tracks_frozen_selection_but_not_shared_library_edits(local,monkeypatch):
    p,reqs,sections=sample();select(p,sections,{sections[0]['id']})
    from app import workflow
    monkeypatch.setattr(workflow,'knowledge_fingerprint',lambda:'source-v1')
    monkeypatch.setattr(workflow,'fingerprint',lambda id:'tender-v1')
    settings={'outline_library':{'revision':'shared-v1'},'company_name':'测试企业'}
    monkeypatch.setattr(db,'get_settings',lambda:settings)
    first=index_materials.input_revision(p,sections)
    settings['outline_library']={'revision':'shared-v2','modules':[{'title':'新共享模块'}]}
    assert index_materials.input_revision(p,sections)==first
    p['metadata']['outline_selection']['groups'][1]['enabled']=True
    assert index_materials.input_revision(p,sections)!=first
    second=index_materials.input_revision(p,sections)
    p['metadata']['outline_selection']['library_snapshot']['revision']='another-frozen-v2'
    assert index_materials.input_revision(p,sections)!=second
    third=index_materials.input_revision(p,sections)
    settings['company_name']='另一企业'
    assert index_materials.input_revision(p,sections)!=third


def test_images_are_reallocated_only_to_enabled_sections_in_selected_order(local,monkeypatch):
    p,reqs,sections=sample()
    a=sections[1];b={**a,'id':'b'*32};sections.append(b)
    spec=p['metadata']['proposal_blueprint']['sections'][1]
    spec.update(title='智能检索',group_title='功能设计')
    p['metadata']['proposal_blueprint']['sections'].append({**spec,'section_id':b['id']})
    asset={'asset_id':'img','sha256':'hash','title':'智能检索界面','reviewed_tags':['智能检索']}
    monkeypatch.setattr(proposal_assets,'load_manifest',lambda p:[asset])
    select(p,sections,{b['id']})
    assert proposal_assets.allocate(p)=={b['id']:[asset]}
    p['_proposal_asset_allocation']={a['id']:[asset],b['id']:[asset]}
    assert proposal_assets.for_section(p,a)==[]
    assert proposal_assets.for_section(p,b)==[asset]
    select(p,sections,{a['id'],b['id']})
    p['metadata']['outline_selection']['groups'].reverse()
    assert proposal_assets.allocate(p)=={b['id']:[asset]}
