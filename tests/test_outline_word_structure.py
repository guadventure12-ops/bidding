"""Actual OOXML heading order equals the editor view, including fenced code."""
from docx import Document
from app import chapter_outline,documents


def test_word_body_and_editor_share_four_level_structure_and_self_wrapper():
    section={'id':'s','title':'接口方案','outline_group_title':'系统集成',
             'content':'# 系统集成\n## 接口方案\n### 1. 业务流程 ###\n正文数字123。\n#### 1.1. 异常处理\n保持原文。\n````text\n### 示例代码\n```\n#### 仍是代码\n````\n'}
    before=section.copy()
    doc=Document()
    documents._write_markdown(doc,chapter_outline.structure_body(section),base_level=3)
    actual=[{'level':int(p.style.name.split()[-1]),'title':p.text} for p in doc.paragraphs if p.style.name.startswith('Heading ')]
    assert actual==chapter_outline.body_headings(section)==[{'level':3,'title':'业务流程'},{'level':4,'title':'异常处理'}]
    assert '正文数字123。' in [p.text for p in doc.paragraphs]
    assert section==before


def test_new_directory_generation_result_checks_order_even_without_review_gate(monkeypatch):
    import pytest
    from app import workflow,provider,review_rules
    spec={'title':'接口方案','group_title':'系统集成','content_kind':'narrative',
          'suboutline':[{'title':'流程','children':[{'title':'校验'}]}]}
    monkeypatch.setattr(review_rules,'active',lambda key:False)
    section={'requirements':[],'_proposal_spec':spec}
    with pytest.raises(provider.ProviderError,match='顺序'):
        workflow._validate_generation_result(section,{}, {'content':'### 其他主题\n实质文字','responses':[]})
    workflow._validate_generation_result(section,{}, {'content':'### 流程\n#### 校验\n实质文字','responses':[]})
