"""The price-form label follows verified tender text, not an internal project name."""
from copy import deepcopy
from io import BytesIO
import json
import zipfile

from docx import Document
import pytest

from app.documents import _quotation_item_name, compose_bid_package


def requirement(name='电子会计档案项目', **changes):
    return {'id': 'price-label', 'category': 'pricing', 'verified': 1,
            'chunk_id': 'source-chunk', 'locator': '正文 / 表格 14 / 第 2 行',
            'title': '报价项目填写', 'text': f'初次报价一览表报价项目应填写“{name}”。',
            'quote': f'{name} |  |', 'response': '待企业核对报价', **changes}


def test_price_form_uses_verified_literal_instead_of_internal_project_name(tmp_path):
    source = requirement()
    original = deepcopy(source)
    path = tmp_path / 'bid.zip'
    compose_bid_package({'name': '示例电子会计档案项目 · 验收'}, [source], [], str(path))
    with zipfile.ZipFile(path) as archive:
        doc = Document(BytesIO(archive.read('03_报价文件.docx')))
        manifest = json.loads(archive.read('导出记录.json'))
    assert doc.tables[0].cell(1, 0).text == '电子会计档案项目'
    assert doc.tables[0].cell(1, 1).text == ''
    assert any('尚无企业确认的金额' in warning for warning in manifest['warnings'])
    assert manifest['quotation_item_source']['sources'][0]['quote'] == source['quote']
    assert not manifest['quotation_confirmed']
    assert source == original


def test_another_product_and_explicit_source_statement_are_not_hardcoded():
    source = requirement('集团费控软件', quote='报价项目应填写“集团费控软件”。', locator='正文 / 段落 85')
    source['text'] = '错误概括应填写其他值；本字段只用于识别，不得作为取值。'
    assert _quotation_item_name([source])['name'] == '集团费控软件'
    assert _quotation_item_name([requirement('财务共享软件', quote='| 财务共享软件 | | |')])['name'] == '财务共享软件'


@pytest.mark.parametrize('changes', [
    {'verified': 0}, {'chunk_id': ''}, {'locator': ''}, {'category': 'technical'},
    {'quote': '', 'response': '报价项目应填写“模型写入值”。'},
    {'quote': '报价金额500000元，具体项目自行填写。'},
    {'quote': '示例：报价项目应填写“示例名称”。'},
    {'quote': '报价项目 |  |'},
    {'quote': '档案系统 | 500000 |'},
    {'quote': '档案系统 |  |\n费控系统 | |'},
    {'quote': '档案系统 |  |', 'title': '其他事项', 'text': '其他事项'},
    {'quote': '档案系统 |  |', 'locator': '正文 / 段落 85'},
    {'quote': '报价项目不应填写“内部名称”。'},
    {'quote': '【待填写 项目名】 |  |'},
])
def test_unverified_response_amount_example_and_header_do_not_supply_value(changes):
    assert _quotation_item_name([requirement(**changes)])['status'] == 'not_found'


def test_conflicting_verified_names_are_not_silently_selected(tmp_path):
    sources = [requirement('档案系统'), requirement('费控系统', id='second')]
    result = _quotation_item_name(sources)
    assert result['status'] == 'conflict' and '待填写' in result['name']
    assert len(result['candidates']) == 2
    path = tmp_path / 'conflict.zip'
    compose_bid_package({'name': '不能掩盖冲突的界面名'}, sources, [], str(path))
    with zipfile.ZipFile(path) as archive:
        doc = Document(BytesIO(archive.read('03_报价文件.docx')))
        manifest = json.loads(archive.read('导出记录.json'))
    assert doc.tables[0].cell(1, 0).text == '待企业确认'
    assert manifest['quotation_item_source'] == result
    assert any('名称' in item and '冲突' in item for item in manifest['warnings'])


def test_duplicate_source_values_remain_unambiguous_and_missing_source_warns(tmp_path):
    result = _quotation_item_name([requirement(), requirement(id='same-value-other-source')])
    assert result['status'] == 'source_verified' and len(result['sources']) == 2
    path = tmp_path / 'fallback.zip'
    compose_bid_package({'name': '原项目名称'}, [], [], str(path))
    with zipfile.ZipFile(path) as archive:
        doc = Document(BytesIO(archive.read('03_报价文件.docx')))
        manifest = json.loads(archive.read('导出记录.json'))
    assert doc.tables[0].cell(1, 0).text == '原项目名称'
    assert manifest['quotation_item_source']['status'] == 'not_found'
    assert any('暂用项目名称' in item for item in manifest['warnings'])


def test_source_locator_alias_preserves_export_provenance():
    source = requirement(locator='', source_locator='正文 / 表格 14 / 第 2 行')
    result = _quotation_item_name([source])
    assert result['name'] == '电子会计档案项目'
    assert result['sources'][0]['locator'] == source['source_locator']


def test_conflicting_literals_in_one_source_are_not_silently_selected():
    source = requirement(quote='报价项目应填写“档案系统”。报价项目须填写“费控系统”。')
    result = _quotation_item_name([source])
    assert result['status'] == 'conflict'
    assert {candidate['name'] for candidate in result['candidates']} == {'档案系统', '费控系统'}
