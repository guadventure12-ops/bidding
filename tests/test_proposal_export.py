import json
import zipfile
from docx import Document
from lxml import etree as E
from test_review_rules import local
from app import proposal_export as pe, proposal_runtime as pr


def sample():
    sections=[];specs=[];reqs=[];ledger=[]
    definitions=[('评审索引表','索引','technical','form'),('需求理解与解决方案','目标与实施','technical','narrative'),
                 ('资格文件','声明与附件','qualification','form'),('报价文件','报价一览','pricing','form'),
                 ('内部响应与交付台账','内部采购程序','internal','internal')]
    for index,(group,title,volume,kind) in enumerate(definitions,1):
        sid=f'{index:032x}';rid='r'+str(index);key='leaf'+str(index)
        sections.append({'id':sid,'title':title,'outline_group_id':'g'+str(index),'outline_group_title':group,'legacy_title':'',
          'content':'仅内部采购程序唯一码' if volume=='internal' else '这是'+title+'的正文。','requirement_ids':[rid],'evidence_ids':[]})
        specs.append({'section_id':sid,'section_key':key,'group_title':group,'title':title,'volume':volume,'content_kind':kind,
          'score_factors':[{'number':'1'}] if kind=='narrative' else []})
        reqs.append({'id':rid,'title':title,'text':title+'采购要求','category':'technical','number':str(index),
          'response':'旧响应不得自动当新结果','status':'drafted','locator':'位置'+str(index),'evidence_ids':[]})
        ledger.append({'requirement_id':rid,'canonical_id':rid,'owner_section_key':key,'disposition':'internal' if volume=='internal' else 'narrative'})
    p={'id':'p','name':'来源项目','company_name':'测试企业','project_number':'TEST-01','buyer':'测试采购方',
       'metadata':{'generation_profile':pr.PROFILE,'proposal_blueprint':{'recognized':True,'sections':specs,'ledger':ledger,
        'score_factors':[{'number':'1','title':'需求理解','points':5}]}}}
    from app import review_index
    model=review_index.build(p,sections,blocks=[])
    sections[0]['content']=review_index.markdown(model)
    for s in sections:p['metadata']=pr.record_result(p['metadata'],s,{'responses':[{'requirement_id':s['requirement_ids'][0],'response':'本次响应','evidence_ids':[],'gap':False}],**({'_review_index_model':model} if s is sections[0] else {})})
    return p,reqs,sections


def test_technical_docx_uses_only_technical_outline_and_real_page_fields(local,tmp_path):
    p,reqs,sections=sample();path=tmp_path/'technical.docx'
    result=pe.write_export(p,reqs,sections,path,{'name':'测试企业'})
    assert not any('未自动回填' in message for message in result['warnings'])
    doc=Document(path)
    assert [x.text for x in doc.paragraphs if x.style.name=='Heading 1']==['评审索引表','需求理解与解决方案']
    assert [x.text for x in doc.paragraphs if x.style.name=='Heading 2']==['索引','目标与实施']
    with zipfile.ZipFile(path) as z:
        xml=z.read('word/document.xml');root=E.fromstring(xml)
    assert b'[[PAGE:' not in xml and '仅内部采购程序唯一码'.encode() not in xml
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    assert len(root.findall('.//w:fldSimple',ns))==1
    assert result['volumes']['technical']['page_reference_fields']==1
    assert result['volumes']['technical']['page_numbers_refreshed'] is True
    assert result['volumes']['technical']['section_pages']


def test_package_keeps_all_requirements_in_separate_ledger_not_three_books(local,tmp_path):
    p,reqs,sections=sample();path=tmp_path/'package.zip'
    result=pe.write_export(p,reqs,sections,path,{'name':'测试企业'},'zip')
    with zipfile.ZipFile(path) as z:
        assert set(z.namelist())=={'01_资格文件.docx','02_商务技术文件.docx','03_报价文件.docx','内部响应与交付台账.csv','内部来源与版本记录.json'}
        ledger=z.read('内部响应与交付台账.csv').decode('utf-8-sig')
        assert 'r5' in ledger and 'r1' in ledger
        for name in ('01_资格文件.docx','02_商务技术文件.docx','03_报价文件.docx'):
            from io import BytesIO
            with zipfile.ZipFile(BytesIO(z.read(name))) as part:assert '仅内部采购程序唯一码'.encode() not in part.read('word/document.xml')
    assert result['internal_ledger_requirements']==5


def test_reference_target_cannot_point_at_a_missing_section(local,tmp_path):
    p,reqs,sections=sample();path=tmp_path/'bad.docx'
    doc=Document();doc.add_paragraph('[[PAGE:'+('f'*32)+']]');doc.save(path)
    import pytest
    with pytest.raises(ValueError,match='章节不存在'):pe.bind_page_references(path,sections)


def test_modified_body_never_exports_the_older_response(local,tmp_path):
    p,reqs,sections=sample();sections[1]['content']='后续修改，响应未核对'
    _,issues=pe.projected_sections(p,reqs,sections,'technical')
    assert any(i['section_id']==sections[1]['id'] for i in issues)


def test_deviation_form_is_preserved_and_never_filled_from_unconfirmed_responses(local,tmp_path):
    p,reqs,sections=sample()
    p['metadata']['proposal_blueprint']['sections'][0]['group_title']='技术条款偏离表'
    sections[0]['content']='| 序号 | 偏离 |\n| --- | --- |\n| | |'
    projected,issues=pe.projected_sections(p,reqs,sections,'technical')
    assert projected[0]['content']==sections[0]['content']
    assert any(i['kind']=='deviation_decision_missing' for i in issues)
    sections[0].update(user_edited=1,content='| 序号 | 偏离 |\n| --- | --- |\n| 1 | 提供离线导入，其他条款待核对 |')
    projected,_=pe.projected_sections(p,reqs,sections,'technical')
    assert projected[0]['content']==sections[0]['content']


def test_manual_index_is_not_overwritten_and_default_uses_source_columns(local,tmp_path):
    p,reqs,sections=sample()
    assert pe.index_content(p).startswith('| 序号 | 评审项目 | 评分标准 | 分值 | 页数 |')
    sections[0].update(user_edited=1,content='| 条款号 | 对应页码 |\n| --- | --- |\n| 人工条款 | [[PAGE:'+sections[1]['id']+']] |')
    projected,_=pe.projected_sections(p,reqs,sections,'technical')
    assert projected[0]['content']==sections[0]['content']


def test_internal_csv_escapes_spreadsheet_formula_without_changing_source(local,tmp_path):
    p,reqs,sections=sample(); reqs[-1]['text']='=HYPERLINK("https://invalid.example")'
    output=tmp_path/'package.zip';pe.write_export(p,reqs,sections,output,{'name':'测试企业'},'zip')
    import csv,io
    with zipfile.ZipFile(output) as z:rows=list(csv.reader(io.StringIO(z.read('内部响应与交付台账.csv').decode('utf-8-sig'))))
    assert rows[-1][3].startswith("'=")
    assert reqs[-1]['text'].startswith('=')


def test_only_short_qualification_forms_keep_rows_together():
    from app.documents import _write_markdown
    short='| 姓名 |  |\n| --- | --- |\n| 职务 |  |\n| 地址 |  |'
    for enabled in (False,True):
        doc=Document();_write_markdown(doc,short,keep_short_forms=enabled)
        assert bool(doc.tables[0].rows[1].cells[0].paragraphs[0].paragraph_format.keep_with_next) is enabled
        assert not doc.tables[0].rows[-1].cells[0].paragraphs[0].paragraph_format.keep_with_next
    doc=Document();_write_markdown(doc,short+'\n'+'\n'.join('| 条件 | 很长的资格条件 |' for _ in range(9)),keep_short_forms=True)
    assert not doc.tables[0].rows[1].cells[0].paragraphs[0].paragraph_format.keep_with_next
