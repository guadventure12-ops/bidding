"""Synthetic document checks for per-operation configurable review rules."""
import copy
import io
import json
import zipfile
from unittest.mock import patch

import pytest
from docx import Document

from app import documents as docs, extra_documents, review_rules


def document_text(path):
    doc = Document(path)
    return '\n'.join([p.text for p in doc.paragraphs] +
                     [cell.text for t in doc.tables for row in t.rows for cell in row.cells])


def quote(total='100', **extra):
    return {'confirmed': True, 'total_including_tax': total, **extra}


def check_quote(data, disabled=None):
    with review_rules.scope(disabled or {}):
        return docs._confirmed_quotation({'quotation': data}, {}, [])


def test_default_helpers_do_not_read_database(tmp_path):
    with patch.object(review_rules.db, 'get_settings', side_effect=AssertionError('database read')):
        assert docs._confirmed_quotation({'quotation': quote()}, {}, [])['total_including_tax'] == '100'
        with pytest.raises(ValueError, match='正式导出被阻止'):
            docs.compose_docx({}, [], [], str(tmp_path / 'invalid.docx'), final=True)


def test_basic_fields_switch_only_affects_its_gate(tmp_path):
    target = tmp_path / 'unfilled.docx'
    with pytest.raises(ValueError, match='项目名称'):
        docs.compose_docx({}, [], [], str(target), final=True)
    with review_rules.scope({'export_basic_fields': False}):
        docs.compose_docx({}, [], [], str(target), final=True)
    assert target.is_file()
    with pytest.raises(ValueError, match='项目名称'):
        docs.compose_docx({}, [], [], str(tmp_path / 'blocked-again.docx'), final=True)


@pytest.mark.parametrize('value', ['-1', '12.345', '10000000000000000'])
def test_decimal_gate_disabled_keeps_exact_value_without_guessing_uppercase(value):
    data = quote(value)
    before = copy.deepcopy(data)
    with pytest.raises(ValueError):
        check_quote(data)
    parsed = check_quote(data, {'quote_decimal': False})
    assert parsed['total_including_tax'] == value
    assert parsed['total_uppercase'] == ''
    assert data == before


@pytest.mark.parametrize('value', ['NaN', 'Infinity', 'not money'])
def test_disabled_review_still_requires_executable_numeric_input(value):
    with pytest.raises(ValueError):
        check_quote(quote(value), {'quote_decimal': False, 'quote_uppercase': False,
                                 'quote_arithmetic': False, 'package_quotation': False})


def test_arithmetic_disabled_does_not_fix_user_supplied_amounts():
    items = [{'name': '测试许可', 'quantity': '2', 'unit_price': '100', 'subtotal': '300'}]
    data = quote('350', items=items)
    with pytest.raises(ValueError, match='小计不一致'):
        check_quote(data)
    parsed = check_quote(data, {'quote_arithmetic': False})
    assert parsed['items'] == items
    assert parsed['total_including_tax'] == '350'
    with pytest.raises(ValueError, match='非负'):
        check_quote(quote('-350', items=items), {'quote_arithmetic': False})


def test_uppercase_off_preserves_supplied_uppercase():
    data = quote('100', total_uppercase='贰佰元整')
    with pytest.raises(ValueError, match='大写与含税总金额不一致'):
        check_quote(data)
    assert check_quote(data, {'quote_uppercase': False})['total_uppercase'] == '贰佰元整'
    assert check_quote(quote('100'), {'quote_uppercase': False})['total_uppercase'] == '壹佰元整'
    assert check_quote(quote('100', total_uppercase='人民币 贰佰元整'),
                       {'quote_uppercase': False})['total_uppercase'] == '人民币 贰佰元整'


def test_package_gate_off_exports_blank_price_without_assuming_budget(tmp_path):
    target = tmp_path / 'package.zip'
    project = {'name': '合成项目', 'project_number': 'SYN-1', 'buyer': '合成采购人', 'budget': '888888'}
    company = {'name': '合成企业'}
    with pytest.raises(ValueError, match='缺少企业确认的报价'):
        docs.compose_bid_package(project, [], [], str(target), company, final=True)
    with review_rules.scope({'package_quotation': False}):
        docs.compose_bid_package(project, [], [], str(target), company, final=True)
    with zipfile.ZipFile(target) as archive:
        manifest = json.loads(archive.read('导出记录.json'))
        assert manifest['quotation_confirmed'] is False
        assert manifest['disabled_rules'] == ['package_quotation']
        assert '部分检查已停用' in archive.read('内部审阅清单.md').decode('utf-8')
        pricing = document_text(io.BytesIO(archive.read('03_报价文件.docx')))
        assert '888888' not in pricing
        assert '________________' in pricing
        assert '人民币大写：' in pricing


def test_missing_confirmed_total_can_stay_blank_when_package_check_off():
    with pytest.raises(ValueError, match='缺少有效'):
        check_quote({'confirmed': True})
    assert check_quote({'confirmed': True}, {'package_quotation': False}) is None


def test_all_document_rules_off_exports_without_manufacturing_values(tmp_path):
    target = tmp_path / 'all-off.zip'
    flags = {rule: False for rule in ('export_basic_fields', 'package_quotation',
             'quote_decimal', 'quote_uppercase', 'quote_arithmetic', 'body_separation')}
    data = quote('-100.123', total_uppercase='原有大写值',
                 items=[{'name': '合成条目', 'quantity': '2', 'unit_price': '-1', 'subtotal': '-8'}])
    project = {'quotation': data}
    before = copy.deepcopy(project)
    with review_rules.scope(flags):
        docs.compose_bid_package(project, [], [], str(target), final=True)
    with zipfile.ZipFile(target) as archive:
        pricing = document_text(io.BytesIO(archive.read('03_报价文件.docx')))
        assert '-100.123' in pricing and '原有大写值' in pricing
        assert '-8' in pricing
    assert project == before


def test_package_gate_off_keeps_existing_unfilled_content(tmp_path):
    project = {'name': '合成项目', 'project_number': 'SYN-1', 'buyer': '合成采购人',
               'quotation': quote('100', items=[{'name': '许可', 'quantity': '1',
                                               'unit_price': '100', 'subtotal': '100'}])}
    sections = [{'title': '报价条款', 'content': '待填写：项目折扣方案'}]
    with pytest.raises(ValueError, match='报价文件仍含待填写内容'):
        docs.compose_bid_package(project, [], sections, str(tmp_path / 'blocked.zip'),
                                 {'name': '合成企业'}, final=True)
    with review_rules.scope({'package_quotation': False, 'body_separation': False}):
        docs.compose_bid_package(project, [], sections, str(tmp_path / 'allowed.zip'),
                                 {'name': '合成企业'}, final=True)
    with zipfile.ZipFile(tmp_path / 'allowed.zip') as archive:
        assert '待填写：项目折扣方案' in document_text(io.BytesIO(archive.read('03_报价文件.docx')))


def test_body_separation_off_keeps_both_sections_and_responses(tmp_path):
    sections = [{'title': '实施方案', 'content': '提供档案检索。\n\n【待补充：实施负责人姓名】'}]
    requirements = [{'text': '提供服务', 'response': '安排现场交付。\n\n【待补充：上线日期】'}]
    originals = copy.deepcopy([sections, requirements])
    off = tmp_path / 'off.docx'
    on = tmp_path / 'on.docx'
    with review_rules.scope({'body_separation': False}):
        result = docs.compose_docx({'name': '测试'}, requirements, sections, str(off))
    assert result['internal_notes'] == []
    assert '【待补充：实施负责人姓名】' in document_text(off)
    assert '【待补充：上线日期】' in document_text(off)
    result = docs.compose_docx({'name': '测试'}, requirements, sections, str(on))
    assert result['internal_notes']
    assert '【待补充：实施负责人姓名】' not in document_text(on)
    assert [sections, requirements] == originals


def test_file_and_text_limits_follow_scope(tmp_path, monkeypatch):
    source = tmp_path / 'small.md'
    source.write_text('合成内容超过测试阈值', encoding='utf-8')
    monkeypatch.setattr(docs, 'MAX_FILE_BYTES', 1)
    monkeypatch.setattr(docs, 'MAX_TEXT_CHARS', 2)
    with pytest.raises(ValueError, match='512 MB'):
        docs.parse_document(str(source))
    with review_rules.scope({'file_parse_safety': False}):
        result = docs.parse_document(str(source))
    assert result['blocks'][0]['text'] == '合成内容超过测试阈值'
    assert result['truncated'] is False
    monkeypatch.setattr(docs, 'MAX_FILE_BYTES', 100)
    assert docs.parse_document(str(source))['truncated'] is True


def test_xml_and_zip_limits_disabled_but_entities_still_rejected(monkeypatch):
    monkeypatch.setattr(docs, 'MAX_XML_BYTES', 4)
    monkeypatch.setattr(docs, 'MAX_ZIP_BYTES', 4)
    with pytest.raises(ValueError, match='XML 超过'):
        docs._safe_xml(b'<root>synthetic</root>')
    data = io.BytesIO()
    with zipfile.ZipFile(data, 'w') as archive:
        archive.writestr('content.txt', 'synthetic')
    with zipfile.ZipFile(data) as archive:
        with pytest.raises(ValueError, match='解压规模'):
            docs._check_zip(archive)
        with review_rules.scope({'file_parse_safety': False}):
            docs._check_zip(archive)
            assert docs._safe_xml(b'<root>synthetic</root>').text == 'synthetic'
            with pytest.raises(ValueError, match='XML 实体声明'):
                docs._safe_xml(b'<!DOCTYPE root [<!ENTITY x "value">]><root>&x;</root>')


def test_pdf_page_limit_can_be_disabled_without_new_ocr_calls(tmp_path, monkeypatch):
    from pypdf import PdfWriter
    source = tmp_path / 'synthetic.pdf'
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    pdf.write(str(source))
    monkeypatch.setattr(docs, 'MAX_PDF_PAGES', 0)
    with pytest.raises(ValueError, match='PDF 超过'):
        docs.parse_document(str(source))
    with review_rules.scope({'file_parse_safety': False}), patch.object(docs, '_ocr_bytes', side_effect=AssertionError('unexpected OCR')):
        assert docs.parse_document(str(source))['page_count'] == 1


def test_extra_attachment_custom_limit_obeys_same_scope():
    class TinySyntheticLargeBytes(bytes):
        def __len__(self):
            return 32 * 1024 * 1024 + 1
    class Source:
        def read_bytes(self):
            return TinySyntheticLargeBytes(b'synthetic,data')
    with pytest.raises(ValueError, match='32 MB'):
        extra_documents._decode(Source(), {})
    with review_rules.scope({'file_parse_safety': False}):
        assert extra_documents._decode(Source(), {}) == 'synthetic,data'


def test_unreadable_document_is_an_execution_error_even_when_gate_off(tmp_path):
    damaged = tmp_path / 'damaged.docx'
    damaged.write_bytes(b'not a valid archive')
    with review_rules.scope({'file_parse_safety': False}):
        with pytest.raises(ValueError, match='损坏'):
            docs.parse_document(str(damaged))
        with pytest.raises(FileNotFoundError):
            docs.parse_document(str(tmp_path / 'does-not-exist.md'))
