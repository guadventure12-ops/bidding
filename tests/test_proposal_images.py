"""Selected local assets become actual Word drawings, with no model/network I/O."""
from copy import deepcopy
from pathlib import Path

import pytest
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from app import review_rules
from app.document_assets import extract_document_assets
from app.documents import _write_markdown, compose_docx


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    import socket
    def blocked(*args, **kwargs):
        raise AssertionError("Word asset tests must never use network")
    monkeypatch.setattr(socket.socket, "connect", blocked)
    monkeypatch.setattr(socket, "getaddrinfo", blocked)
    with review_rules.scope({key: True for key in review_rules.IDS}):
        yield


@pytest.fixture
def image_asset(tmp_path):
    image = tmp_path / "product.png"
    Image.new("RGB", (1200, 500), (40, 120, 140)).save(image)
    source = tmp_path / "product.md"
    source.write_text('# 产品查询功能\n![原始查询图](product.png)\n', encoding="utf-8")
    result = extract_document_assets(source, document_id="doc-product", output_dir=tmp_path / "assets")
    assert not result["warnings"]
    return result["assets"][0]


def project():
    return {"name": "合成电子档案项目", "buyer": "合成采购人", "project_number": "QA-IMAGES",
            "_omit_response_table": True, "_omit_attachments": True}


def section(content):
    return {"id": "section-leaf", "title": "档案查询", "outline_group_id": "group-tech",
            "outline_group_title": "技术方案", "legacy_title": "技术响应（1）", "ordinal": 1,
            "content": content, "status": "approved", "requirement_ids": [], "evidence_ids": []}


def company(asset, tmp_path):
    return {"name": "合成投标企业", "assets": [asset], "asset_root": str(tmp_path / "assets")}


def marker(asset, caption="图1 档案查询示意"):
    return f"![{caption}](asset:{asset['asset_id']})"


def test_selected_manifest_becomes_inline_picture_caption_and_source_trace(tmp_path, image_asset):
    body = '查询正文123.45。\n\n' + marker(image_asset) + '\n\n后续正文2026年。'
    rows = [section(body)]
    enterprise = company(image_asset, tmp_path)
    before = deepcopy((rows, enterprise))
    target = tmp_path / "with-image.docx"
    result = compose_docx(project(), [], rows, str(target), enterprise)
    doc = Document(target)
    assert len(doc.inline_shapes) == 1
    picture = doc.inline_shapes[0]
    page = doc.sections[-1]
    assert picture.width <= page.page_width - page.left_margin - page.right_margin
    assert picture.height <= page.page_height - page.top_margin - page.bottom_margin - Pt(45)
    assert abs(picture.width / picture.height - image_asset["width"] / image_asset["height"]) < 0.00001
    assert picture._inline.docPr.get("descr") == "图1 档案查询示意"
    assert picture._inline.docPr.get("title") == "图1 档案查询示意"
    paragraph_index = next(i for i, p in enumerate(doc.paragraphs) if p._p.xpath('.//wp:inline'))
    picture_paragraph, caption = doc.paragraphs[paragraph_index:paragraph_index + 2]
    assert picture_paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert picture_paragraph.paragraph_format.first_line_indent.pt == 0
    assert picture_paragraph.paragraph_format.keep_with_next is True
    assert picture_paragraph.paragraph_format.keep_together is True
    assert caption.style.name == "Caption" and caption.text == "图1 档案查询示意"
    assert caption.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert caption.paragraph_format.first_line_indent.pt == 0
    assert caption._p.xpath('./w:pPr/w:ind')[0].get(qn('w:firstLineChars')) == "0"
    assert caption.paragraph_format.keep_together is True
    assert caption.paragraph_format.keep_with_next is False
    assert caption.runs[0].font.italic is False and caption.runs[0].font.size.pt == 12
    for text in ('查询正文123.45。', '后续正文2026年。'):
        p = next(p for p in doc.paragraphs if p.text == text)
        assert p.style.name == "Normal" and p.paragraph_format.first_line_indent.pt == 24
    assert result["used_assets"] == [{
        "asset_id": image_asset["asset_id"], "sha256": image_asset["sha256"],
        "document_id": "doc-product", "source_sha256": image_asset["source_sha256"],
        "source": image_asset["source"], "caption": "图1 档案查询示意", "section_id": "section-leaf"}]
    assert (rows, enterprise) == before


def test_tall_picture_fits_page_with_caption_and_stays_proportional(tmp_path):
    png = tmp_path / "tall.png"
    Image.new("RGB", (120, 1800), "white").save(png)
    md = tmp_path / "tall.md"
    md.write_text('![功能清单](tall.png)', encoding="utf-8")
    asset = extract_document_assets(md, document_id="tall", output_dir=tmp_path / "assets")["assets"][0]
    target = tmp_path / "tall.docx"
    compose_docx(project(), [], [section(marker(asset))], str(target), company(asset, tmp_path))
    doc = Document(target)
    picture = doc.inline_shapes[0]
    page = doc.sections[-1]
    assert picture.height <= page.page_height - page.top_margin - page.bottom_margin - Pt(45)
    assert abs(picture.width / picture.height - 120 / 1800) < 0.00001


def test_caption_heading_symbols_are_plain_caption_not_directory(tmp_path, image_asset):
    caption = "## 1. 档案查询与API 3.2"
    target = tmp_path / "caption.docx"
    compose_docx(project(), [], [section(marker(image_asset, caption))], str(target), company(image_asset, tmp_path))
    doc = Document(target)
    headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading ")]
    assert headings == [("Heading 1", "技术方案"), ("Heading 2", "档案查询")]
    captions = [p for p in doc.paragraphs if p.style.name == "Caption"]
    assert len(captions) == 1 and captions[0].text == caption
    assert not captions[0]._p.xpath('./w:pPr/w:numPr')


def test_image_does_not_change_other_prose_table_or_persisted_outline(tmp_path, image_asset):
    body = '## 查询配置\n原始正文50000条，版本3.2。\n\n| 功能 | 值 |\n| --- | --- |\n| 数量 | 300 |\n\n末尾段落。'
    plain = tmp_path / "plain.docx"
    with_image = tmp_path / "pictured.docx"
    before = compose_docx(project(), [], [section(body)], str(plain), {"name": "合成投标企业"})
    after = compose_docx(project(), [], [section(body + '\n' + marker(image_asset))], str(with_image), company(image_asset, tmp_path))
    a, b = Document(plain), Document(with_image)
    assert before["outline"] == after["outline"]
    assert [(p.text, p.style.name) for p in a.paragraphs] == [
        (p.text, p.style.name) for p in b.paragraphs
        if p.style.name != "Caption" and not p._p.xpath('.//wp:inline')]
    assert [t._tbl.xml for t in a.tables] == [t._tbl.xml for t in b.tables]
    assert before["used_assets"] == []


@pytest.mark.parametrize("change", ["unknown", "path", "hash", "crop", "rotation", "missing_root"])
def test_unselected_or_changed_or_unreproduced_image_rejected_without_overwriting_export(tmp_path, image_asset, change):
    selected = deepcopy(image_asset)
    content = marker(image_asset)
    enterprise = company(selected, tmp_path)
    if change == "unknown":
        content = '![未知图片](asset:asset-' + 'f' * 64 + ')'
    elif change == "path":
        selected["path"] = str(tmp_path / "product.png")
    elif change == "hash":
        Path(selected["path"]).write_bytes(b"changed contents")
    elif change == "crop":
        selected["source"]["display_transform"] = {"crop": {"l": "25000"}}
    elif change == "rotation":
        selected["source"]["display_transform"] = {"rotation_flip": {"rot": "5400000"}}
    else:
        enterprise.pop("asset_root")
    target = tmp_path / "preserve.docx"
    target.write_bytes(b"prior export must remain")
    with pytest.raises(ValueError):
        compose_docx(project(), [], [section(content)], str(target), enterprise)
    assert target.read_bytes() == b"prior export must remain"


@pytest.mark.parametrize("content", [
    '![图片](C:/private/image.png)', '![图片](https://example.test/image.png)',
    '![图片](../../image.png)', '![图片](data:image/png;base64,abcd)',
    '前缀 ![图片](asset:asset-{id})', '![图片](asset:asset-{id}) 后缀',
    '| 图片 | 值 |\n| --- | --- |\n| ![图片](asset:asset-{id}) | 1 |',
    '![图片](asset:asset-wrong-id)',
])
def test_only_standalone_selected_asset_syntax_can_embed(tmp_path, image_asset, content):
    target = tmp_path / "invalid.docx"
    content = content.replace("{id}", image_asset["asset_id"].removeprefix("asset-"))
    with pytest.raises(ValueError):
        compose_docx(project(), [], [section(content)], str(target), company(image_asset, tmp_path))
    assert not target.exists()


def test_no_manifest_does_not_accept_new_asset_marker_but_keeps_old_text_links(tmp_path):
    with pytest.raises(ValueError, match="获准范围"):
        compose_docx(project(), [], [section('![未选择](asset:asset-' + 'a' * 64 + ')')], str(tmp_path / "unknown.docx"))
    target = tmp_path / "legacy.docx"
    result = compose_docx(project(), [], [section('![旧文本](local.png)\n[外部文档](https://example.test/help)')], str(target), {"assets": []})
    doc = Document(target)
    assert not doc.inline_shapes
    assert "旧文本（local.png）" in [p.text for p in doc.paragraphs]
    assert "外部文档（https://example.test/help）" in [p.text for p in doc.paragraphs]
    assert result["used_assets"] == []


def test_repeated_asset_uses_traced_per_section_and_unused_manifest_not_loaded(tmp_path, image_asset):
    extra = deepcopy(image_asset)
    extra["asset_id"] = "asset-" + "b" * 64
    extra["path"] = "C:/not-read/image.png"
    rows = [section(marker(image_asset)), {**section(marker(image_asset, "图2 重复引用")), "id": "second", "title": "查询示例", "ordinal": 2}]
    enterprise = company(image_asset, tmp_path)
    enterprise["assets"].append(extra)
    target = tmp_path / "reused.docx"
    result = compose_docx(project(), [], rows, str(target), enterprise)
    assert len(Document(target).inline_shapes) == 2
    assert [u["section_id"] for u in result["used_assets"]] == ["section-leaf", "second"]
    assert all(u["asset_id"] == image_asset["asset_id"] for u in result["used_assets"])


def test_code_examples_do_not_embed_assets_and_direct_writer_returns_trace(tmp_path, image_asset):
    code = '```markdown\n' + marker(image_asset) + '\n```\n`' + marker(image_asset) + '`'
    doc = Document()
    args = {"assets": {image_asset["asset_id"]: image_asset}, "asset_root": tmp_path / "assets"}
    assert _write_markdown(doc, code, **args) == []
    assert not doc.inline_shapes
    result = _write_markdown(doc, marker(image_asset), **args)
    assert len(doc.inline_shapes) == 1
    assert result[0]["asset_id"] == image_asset["asset_id"]


def test_conflicting_manifest_ids_are_not_resolved_by_last_entry(tmp_path, image_asset):
    changed = deepcopy(image_asset)
    changed["sha256"] = "0" * 64
    enterprise = company(image_asset, tmp_path)
    enterprise["assets"].append(changed)
    with pytest.raises(ValueError, match="冲突"):
        compose_docx(project(), [], [section(marker(image_asset))], str(tmp_path / "ambiguous.docx"), enterprise)
