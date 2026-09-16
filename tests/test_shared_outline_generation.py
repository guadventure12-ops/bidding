"""Canonical logical leaves remain stable across model transport batches.

Every model call below is an isolated stub. No provider or network is used.
"""
import copy
import json
import socket

import pytest

from app import db, workflow as w, section_generation as sg, provider, model_jobs, review_rules as rr, chapter_outline


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(db, 'DATA', tmp_path)
    monkeypatch.setenv('MX_TESTING', '1')
    monkeypatch.setenv('LANGFUSE_ENABLED', 'false')
    monkeypatch.setattr(socket.socket, 'connect', lambda *a, **k: (_ for _ in ()).throw(AssertionError('network prohibited')))
    monkeypatch.setattr(provider, 'chat_json', lambda *a, **k: (_ for _ in ()).throw(AssertionError('paid model prohibited')))
    monkeypatch.setattr(provider, 'key_configured', lambda: True)
    monkeypatch.setattr(w, 'search_evidence', lambda *a, **k: [])
    monkeypatch.setattr(w, '_requirement_source_context', lambda *a: {})
    db.init()
    for id in ('p', 'other'):
        db.insert('projects', {'id':id, 'name':'合成项目', 'company_name':'测试企业', 'domain':'archive', 'created_at':db.now(), 'updated_at':db.now()})
    db.insert('documents', {'id':'knowledge', 'name':'合成产品资料', 'path':'synthetic.md', 'sha256':'1', 'status':'approved', 'scope':'archive', 'parse_status':'ready', 'created_at':db.now(), 'updated_at':db.now()})
    requirements=[]
    for index in range(25):
        row={'id':f'r{index:02d}','project_id':'p','category':'technical','title':f'功能要求{index}',
             'text':f'应提供档案检索功能{index}','quote':f'应提供档案检索功能{index}','response':f'保留的逐条响应{index}',
             'created_at':db.now(),'updated_at':db.now()}
        db.insert('requirements',row);requirements.append(row['id'])
    for id,pid,rids,ordinal,topic in [('selected','p',requirements,0,'归档与检索'),('untouched','p',[],1,'借阅与权限'),('foreign','other',[],0,'其他项目主题')]:
        db.insert('sections', {'id':id,'project_id':pid,'ordinal':ordinal,'title':topic,'outline_group_id':pid+'-technical',
            'outline_group_title':'技术与功能方案','legacy_title':'技术与功能方案（1）','content':'已认可正文，数值123。\n\n|字段|值|\n|---|---|\n|保留|123|',
            'status':'approved','user_edited':1,'requirement_ids':rids,'created_at':db.now(),'updated_at':db.now()})
    db.update('projects','p',{'analysis_status':'complete','analysis_fingerprint':w.fingerprint('p')})
    db.insert('reviews',{'id':'historical','project_id':'p','fingerprint':'old','status':'passed','created_at':db.now()})
    db.insert('exports',{'id':'old-export','project_id':'p','name':'old.docx','path':'untouched.docx','format':'docx','created_at':db.now()})
    calls=[]
    def model(job_id, mode, key, system, prompt, inputs, cancel=None):
        calls.append({'mode':mode,'key':key,'prompt':prompt,'inputs':copy.deepcopy(inputs)})
        reqs=inputs['requirements']
        marker=','.join(row['id'] for row in reqs)
        return {'content':'# '+inputs['section_title']+'\n\n### 实际功能说明\n\n本次正文覆盖 '+marker+'，保留数值123与表格。\n\n|字段|值|\n|---|---|\n|数值|123|',
                'responses':[{'requirement_id':row['id'],'response':'实际功能响应'+row['id'],'evidence_ids':[],'gap':False} for row in reqs],
                '_usage':{'prompt_tokens':11,'completion_tokens':7,'prompt_tokens_details':{'cached_tokens':3}}}
    monkeypatch.setattr(model_jobs,'chat_json',model)
    return calls,model


def single_job(request_id='request-1'):
    preview=sg.preview('selected')
    payload=sg.admit('selected',preview['revision'],'按既有目录完善正文',request_id,preview['outline_revision'])
    job={'id':request_id,'project_id':'p','mode':'generate','payload':payload,'checkpoint':{},'created_at':db.now()}
    db.insert('jobs',job)
    return job,preview


def all_job(id='whole-1'):
    job={'id':id,'project_id':'p','mode':'generate','payload':{},'checkpoint':{},'created_at':db.now()}
    db.insert('jobs',job)
    return job


def test_single_leaf_three_batches_preserves_other_rows_and_history(isolated):
    calls,_=isolated
    before={table:db.all('SELECT * FROM '+table+' ORDER BY id') for table in ('sections','requirements','reviews','exports','documents')}
    before_outline=chapter_outline.view('p')
    job,preview=single_job()
    assert preview['batch_count']==3 and '分3批' in preview['notice']
    assert preview['outline_group_title']=='技术与功能方案'
    result=sg.run(job['id'],job)
    assert result['sections']==1 and len(calls)==3
    assert [len(call['inputs']['requirements']) for call in calls]==[12,12,1]
    assert len({call['key'] for call in calls})==3
    after=db.one('SELECT * FROM sections WHERE id=?',('selected',))
    original=next(row for row in before['sections'] if row['id']=='selected')
    assert w._outline_identity(after)==w._outline_identity(original)
    assert after['status']=='draft' and after['user_edited']==1
    assert '# 归档与检索' not in after['content']
    assert after['content'].count('### 实际功能说明')==3
    assert after['content'].count('|数值|123|')==3
    assert 'r00' in after['content'] and 'r24' in after['content']
    for row in before['sections']:
        if row['id']!='selected':assert db.one('SELECT * FROM sections WHERE id=?',(row['id'],))==row
    for table in ('requirements','reviews','exports','documents'):
        assert db.all('SELECT * FROM '+table+' ORDER BY id')==before[table]
    snapshot=db.one('SELECT * FROM project_snapshots WHERE id=?',(result['snapshot_id'],))
    assert snapshot['payload']['before']==original
    after_outline=chapter_outline.view('p')
    # Only body/status may differ. Directory order and titles are frozen.
    def directory(view):
        return [(group['id'],group['title'],[(part['id'],part['title']) for part in group['sections']]) for group in view['groups']]
    assert directory(before_outline)==directory(after_outline)


def test_later_batch_failure_preserves_whole_leaf_and_no_snapshot(isolated,monkeypatch):
    calls,model=isolated
    before={table:db.all('SELECT * FROM '+table+' ORDER BY id') for table in ('sections','requirements','reviews','exports','project_snapshots')}
    def failing(*a,**k):
        if len(calls)==1:raise provider.ProviderError('second batch transport failure')
        return model(*a,**k)
    monkeypatch.setattr(model_jobs,'chat_json',failing)
    job,_=single_job()
    with pytest.raises(provider.ProviderError,match='second batch'):sg.run(job['id'],job)
    for table,rows in before.items():assert db.all('SELECT * FROM '+table+' ORDER BY id')==rows


def test_all_approved_or_manual_leaves_are_preserved_even_with_rules_off(isolated):
    calls,_=isolated
    before=db.all('SELECT * FROM sections ORDER BY id')
    db.set_setting('review_rules',{rule:False for rule in rr.IDS})
    result=w.run_generate('whole-preserve',all_job('whole-preserve'))
    assert result['generated_sections']==0 and result['preserved_sections']==2
    assert not calls
    assert db.all('SELECT * FROM sections ORDER BY id')==before


def test_repeated_whole_generation_does_not_create_batch_chapters(isolated):
    calls,_=isolated
    # One established logical chapter owns 25 requirements; request batches must
    # never become extra chapters. New-project confirmation is tested by the API suite.
    db.execute('DELETE FROM sections WHERE id=?',('untouched',))
    db.update('sections','selected',{'content':'','status':'draft','user_edited':0})
    result=w.run_generate('whole-first',all_job('whole-first'))
    first=db.all('SELECT * FROM sections WHERE project_id=?',('p',))
    assert result['generated_sections']==1 and len(first)==1 and len(calls)==3
    assert len(first[0]['requirement_ids'])==25
    assert '（' not in first[0]['title']
    second=w.run_generate('whole-second',all_job('whole-second'))
    again=db.all('SELECT * FROM sections WHERE project_id=?',('p',))
    assert second['generated_sections']==1 and len(again)==1
    assert w._outline_identity(first[0])==w._outline_identity(again[0])
    assert db.one('SELECT * FROM sections WHERE id=?',('foreign',))['status']=='approved'


def test_new_explicit_operation_has_distinct_batch_cache_identity(isolated):
    calls,_=isolated
    job,_=single_job('operation-1');sg.run(job['id'],job)
    job,_=single_job('operation-2');sg.run(job['id'],job)
    assert len(calls)==6
    assert {call['key'] for call in calls[:3]}.isdisjoint(call['key'] for call in calls[3:])
    assert calls[0]['inputs']['outline']['id']=='selected'
    assert calls[0]['inputs']['outline']['_operation_key']=='operation-1'
    # The citation JSON remains parseable after operation metadata is inserted.
    assert json.loads(calls[0]['prompt'].split('\n证据数据：\n')[-1])==[]


def test_outline_change_is_blocked_even_when_revision_rule_disabled(isolated,monkeypatch):
    calls,model=isolated
    db.set_setting('review_rules',{'operation_revision':False})
    preview=sg.preview('selected')
    db.update('sections','selected',{'title':'新目录主题'})
    with pytest.raises(ValueError,match='目录'):
        sg.admit('selected',preview['revision'],'','new',preview['outline_revision'])
    with pytest.raises(ValueError,match='刷新'):
        sg.admit('selected',preview['revision'],'','new')
    job,_=single_job()
    def changed(*a,**k):
        result=model(*a,**k)
        if len(calls)==1:db.update('sections','selected',{'title':'调用期间新目录'})
        return result
    monkeypatch.setattr(model_jobs,'chat_json',changed)
    with pytest.raises(ValueError,match='目录'):sg.run(job['id'],job)
    assert db.one('SELECT * FROM sections WHERE id=?',('selected',))['title']=='调用期间新目录'
    assert not db.all('SELECT * FROM project_snapshots')


def test_batch_notes_and_responses_never_escape_selected_requirements(isolated,monkeypatch):
    def generated(j,key,unit,evidence,prompt,**kwargs):
        rid=unit['requirements'][0]['id']
        return {'content':'保留有效功能正文。','responses':[{'requirement_id':rid,'response':'已有响应','evidence_ids':[]},
            {'requirement_id':'foreign','response':'不得写入外部要求','evidence_ids':[]}],
            '_internal_notes':[{'requirement_ids':['foreign'],'raw':'外部','message':'外部'},
                               {'requirement_ids':[rid,'foreign'],'raw':'本节','message':'本节'}]}
    monkeypatch.setattr(w,'_generate_with_repairs',generated)
    section,p,reqs,_=sg.inputs('selected')
    with rr.scope({'generation_coverage':False}):
        result,_=w.generate_section_batches('test',p,{**section,'requirements':reqs},'one')
    assert all(row['requirement_id']!='foreign' for row in result['responses'])
    assert all('foreign' not in note['requirement_ids'] for note in result['_internal_notes'])
    assert len(result['_internal_notes'])==3


def test_usage_only_sums_actual_fields_and_headers_are_narrow(isolated):
    section,p,reqs,_=sg.inputs('selected')
    result,_=w.generate_section_batches('usage',p,{**section,'requirements':reqs},'one')
    assert result['_usage']=={'prompt_tokens':33,'completion_tokens':21,'prompt_tokens_details':{'cached_tokens':9}}
    assert w._strip_generation_heading('# 技术与功能方案\n\n## 归档与检索\n\n正文123\n# 归档与检索',section)=='正文123\n# 归档与检索'
    assert w._strip_generation_heading('# 其他业务标题\n数字123',section)=='# 其他业务标题\n数字123'
    assert w._strip_generation_heading('    indented code 123',section)=='    indented code 123'
    with pytest.raises(ValueError,match='8000'):w.generation_batches([{'id':'long','text':'字'*8001}])
    assert len(w.generation_batches([{'id':'a','text':'字'*4001},{'id':'b','text':'字'*4000}]))==2


def test_regeneration_export_uses_the_same_fixed_hierarchy(isolated,tmp_path):
    from docx import Document
    from app.documents import compose_docx
    before=chapter_outline.view('p')
    job,_=single_job();sg.run(job['id'],job)
    rows=db.all('SELECT * FROM sections WHERE project_id=? ORDER BY ordinal,id',('p',))
    path=tmp_path/'regenerated-shared-outline.docx'
    with rr.scope({rule:False for rule in rr.IDS}):
        compose_docx({'name':'合成项目','_omit_response_table':True,'_omit_attachments':True},[],rows,str(path))
    exported=Document(path)
    h1=[p.text for p in exported.paragraphs if p.style.name=='Heading 1']
    h2=[p.text for p in exported.paragraphs if p.style.name=='Heading 2']
    assert h1==[group['title'] for group in before['groups']]
    assert h2==[s['title'] for group in before['groups'] for s in group['sections']]
    assert not any('（1）' in text for text in h1+h2)
    assert len(rows)==2 and rows[0]['id']=='selected'
    assert exported.tables[0].cell(1,1).text=='123'


def test_whole_leaf_failure_never_partially_replaces_an_existing_draft(isolated,monkeypatch):
    calls,model=isolated
    db.update('sections','selected',{'status':'draft','user_edited':0})
    before={table:db.all('SELECT * FROM '+table+' ORDER BY id') for table in ('sections','requirements','reviews','exports')}
    def fail_later(*a,**k):
        if len(calls)==1:raise provider.ProviderError('later batch failure')
        return model(*a,**k)
    monkeypatch.setattr(model_jobs,'chat_json',fail_later)
    job=all_job()
    with pytest.raises(provider.ProviderError,match='later batch'):w.run_generate(job['id'],job)
    for table,rows in before.items():assert db.all('SELECT * FROM '+table+' ORDER BY id')==rows


def test_distinct_logical_leaves_may_run_two_at_once_but_keep_identity(isolated,monkeypatch):
    import threading
    import time
    calls,model=isolated
    rids=[row['id'] for row in db.all('SELECT id FROM requirements WHERE project_id=? ORDER BY id',('p',))]
    db.update('sections','selected',{'status':'draft','user_edited':0,'requirement_ids':rids[:12]})
    db.update('sections','untouched',{'status':'draft','user_edited':0,'requirement_ids':rids[12:]})
    barrier=threading.Barrier(2,timeout=5)
    lock=threading.Lock()
    active=maximum=count=0
    def concurrent(*args,**kwargs):
        nonlocal active,maximum,count
        with lock:
            active+=1;maximum=max(maximum,active);count+=1;index=count
        if index<=2:barrier.wait()
        time.sleep(.02)
        result=model(*args,**kwargs)
        with lock:active-=1
        return result
    monkeypatch.setattr(model_jobs,'chat_json',concurrent)
    before={row['id']:w._outline_identity(row) for row in db.all('SELECT * FROM sections WHERE project_id=?',('p',))}
    job=all_job()
    result=w.run_generate(job['id'],job)
    assert maximum==2 and count==3 and result['generated_sections']==2
    assert {row['id']:w._outline_identity(row) for row in db.all('SELECT * FROM sections WHERE project_id=?',('p',))}==before
    assert len(db.one('SELECT * FROM jobs WHERE id=?',(job['id'],))['checkpoint']['sections'])==2


def test_sibling_leaf_commits_once_failed_leaf_retries_all_batches(isolated,monkeypatch):
    import threading
    calls,model=isolated
    rids=[row['id'] for row in db.all('SELECT id FROM requirements WHERE project_id=? ORDER BY id',('p',))]
    db.update('sections','selected',{'status':'draft','user_edited':0,'requirement_ids':rids[:12]})
    db.update('sections','untouched',{'status':'draft','user_edited':0,'requirement_ids':rids[12:]})
    original=db.one('SELECT * FROM sections WHERE id=?',('selected',))
    barrier=threading.Barrier(2,timeout=5)
    seen=set();lock=threading.Lock()
    def failing(*args,**kwargs):
        unit_id=args[5]['outline']['id']
        with lock:first=unit_id not in seen;seen.add(unit_id)
        if first:barrier.wait()
        if unit_id=='selected':raise provider.ProviderError('selected transport failure')
        return model(*args,**kwargs)
    monkeypatch.setattr(model_jobs,'chat_json',failing)
    job=all_job()
    with pytest.raises(provider.ProviderError,match='selected transport'):w.run_generate(job['id'],job)
    assert db.one('SELECT * FROM sections WHERE id=?',('selected',))==original
    completed=db.one('SELECT * FROM sections WHERE id=?',('untouched',))
    checkpoint=db.one('SELECT * FROM jobs WHERE id=?',(job['id'],))['checkpoint']
    assert set(checkpoint['sections'])=={'untouched'}
    monkeypatch.setattr(model_jobs,'chat_json',model)
    retry=all_job('retry');retry['checkpoint']=checkpoint
    db.update('jobs',retry['id'],{'checkpoint':checkpoint})
    previous_calls=len(calls)
    result=w.run_generate(retry['id'],retry)
    assert len(calls)==previous_calls+1
    assert result['generated_sections']==1 and result['resumed_sections']==1
    assert db.one('SELECT * FROM sections WHERE id=?',('untouched',))==completed


def test_pre_migration_generation_history_restores_body_without_old_directory(isolated):
    current=db.one('SELECT * FROM sections WHERE id=?',('selected',))
    old_after={key:value for key,value in current.items() if key not in ('outline_group_id','outline_group_title','legacy_title')}
    old_after['title']=current['legacy_title']
    old_before={**old_after,'content':'历史模型前原正文300，表格保留。'}
    db.insert('project_snapshots',{'id':'pre-migration','project_id':'p','label':'section_regeneration:old',
        'created_at':db.now(),'payload':{'kind':'section_regeneration','section_id':'selected',
        'before':old_before,'after':old_after,'before_todos':[],'after_todos':[]}})
    sg.restore('pre-migration')
    restored=db.one('SELECT * FROM sections WHERE id=?',('selected',))
    assert restored['content']==old_before['content'] and restored['status']=='draft'
    assert w._outline_identity(restored)==w._outline_identity(current)
    assert restored['title']=='归档与检索'
    assert db.one('SELECT * FROM project_snapshots WHERE id=?',('pre-migration',))['payload']['after']==old_after
